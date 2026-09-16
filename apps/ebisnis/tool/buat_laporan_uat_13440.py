"""Generate per-variant UAT reports from verified local-run evidence, never mocks."""
import json
import re
import sys
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'output' / 'uat-1.34.40-20260917'
NAMES = {'albahjah': 'Al-Bahjah POS', 'nahl': 'TokoQu Al-Bahjah An Nahl'}
STEPS = {
    'rekap-2': 'Buka Rincian Produk, pilih Rekap per produk dan filter Voucher Santri.',
    'rekap-3': 'Ubah filter menjadi Voucher Pejuang tanpa mengubah produk atau periode.',
    'rekap-6': 'Ubah filter menjadi Tunai/Transfer/QRIS. Gabungkan tiga transaksi reguler.',
    'rincian-reguler': 'Dengan filter reguler tetap aktif, beralih ke mode Rincian.',
    'hpp-awal': 'Buka Data Produk dan tunggu katalog fixture tersimpan di SQLite.',
    'hpp-edit': 'Buka produk UAT MINUMAN, temukan Harga Beli, isi 6000.',
    'hpp-sesudah-simpan': 'Tekan Simpan, tunggu antrean dikirim ke fixture, lalu segarkan daftar.',
    'hpp-restart-offline': 'Tutup koneksi database, buka ulang layar dan SQLite; fixture mengembalikan HTTP 503.',
    'pelunasan-kalender': 'Buka Entri Pelunasan Piutang lalu buka pemilih tanggal.',
    'pelunasan-entri': 'Pilih tanggal dua hari sebelumnya, isi Rp15.000 dan keterangan, lalu simpan.',
    'histori-online': 'Buka Histori Pembayaran untuk pelanggan UAT dan periode bulan berjalan.',
    'histori-offline': 'Saat histori tersimpan lokal, simulasikan HTTP 503 lalu tekan Muat Ulang.',
}

def table(doc, headings, rows, widths):
    t = doc.add_table(rows=1, cols=len(headings))
    t.autofit = False
    for c, title, width in zip(t.rows[0].cells, headings, widths):
        c.width = Inches(width)
        c.text = title
        for r in c.paragraphs[0].runs:
            r.bold = True
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EEF2F5'); c._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for c, value, width in zip(cells, row, widths):
            c.width = Inches(width); c.text = str(value)
    for row in t.rows:
        for c in row.cells:
            pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'CCD3DA'); borders.append(e)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), '85'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            pr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                for r in p.runs: r.font.size = Pt(10)
    return t

def build(variant):
    evidence = json.loads((OUT / variant / 'skenario.json').read_text(encoding='utf-8'))
    assert len(evidence) == 12 and all(r['status'] == 'LULUS' for r in evidence)
    logs = [('Regresi modul dan kontrak varian', f'regresi-{variant}.log', 105), ('Database bersama', 'core-db.log', 14), ('Isolasi pembaruan', 'core-update.log', 15)]
    for _, name, count in logs:
        content = (OUT / 'logs' / name).read_text(encoding='utf-8-sig')
        assert re.search(rf'\+{count}: All tests passed!', content), name
    integration = (OUT / 'logs' / f'layar-{variant}.log').read_text(encoding='utf-8-sig')
    assert 'All tests passed!' in integration
    assert 'No issues found!' in (OUT / 'logs' / 'analyzer.log').read_text(encoding='utf-8-sig')
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    sec.top_margin = sec.bottom_margin = Inches(.55)
    sec.left_margin = sec.right_margin = Inches(.65)
    sec.header_distance = sec.footer_distance = Inches(.22)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        s = doc.styles[name]; s.font.name = 'Calibri'; s.font.color.rgb = RGBColor(0,0,0)
        s.font.underline = False
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(26)
    doc.styles['Heading 1'].font.size = Pt(17)
    footer = sec.footer.paragraphs[0]
    footer.add_run(f'{NAMES[variant]} | UAT lokal 1.34.40 build 203 | 17 September 2026     ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for r in footer.runs: r.font.size = Pt(9)
    doc.add_heading('Laporan UAT pascarilis', 0)
    doc.add_paragraph(f'{NAMES[variant]}\nVersi 1.34.40 build 203', style='Subtitle')
    revision = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()
    doc.add_paragraph(f'Revisi sumber {revision[:12]} | Eksekusi 17 September 2026')
    doc.add_heading('Keputusan pengujian', 1)
    doc.add_paragraph('LULUS pada lingkup UAT lokal yang disetujui. Seluruh skenario otomatis dan bukti layar yang tercantum di laporan ini lulus. Ini bukan persetujuan operasional dari petugas toko atau hasil pengujian server produksi.')
    table(doc, ['Kelompok pengujian', 'Hasil', 'Bukti'], [
        (logs[0][0], '105/105 lulus', logs[0][1]),
        (logs[1][0], '14/14 lulus', logs[1][1]),
        (logs[2][0], '15/15 lulus', logs[2][1]),
        ('Alur layar Windows', '1 alur, 12 checkpoint lulus', f'layar-{variant}.log dan 12 screenshot'),
        ('Analisis statis perubahan', '14 berkas, tanpa temuan', 'analyzer.log'),
    ], [3.4, 2.3, 4.0])
    doc.add_heading('Lingkup dan lingkungan', 1)
    doc.add_paragraph('Tiga perubahan setelah rilis sebelumnya: rekap produk per kelompok pembayaran; persistensi HPP; tanggal pelunasan dan histori pembayaran piutang. Layar produksi dijalankan oleh Flutter integration_test pada Windows x64, dengan HTTP loopback dan SQLite bernamespace khusus UAT. Seluruh pelanggan, produk, nominal dan transaksi adalah data uji. Tidak ada kredensial atau mutasi server toko.')
    doc.add_paragraph('Target distribusi sesuai persetujuan: APK release dengan debug signing dan installer Windows unsigned, tanpa ZIP auto-update. Verifikasi paket dicatat terpisah pada catatan rilis. Pengujian Android fisik, printer toko, saldo produksi, dan koneksi cabang tidak termasuk cakupan ini. Endpoint fixture di luar skenario memberi data kosong; kartu ringkasan global bukan objek uji. Data dan preferensi pengguna tidak dihapus.')
    doc.add_page_break()
    doc.add_heading('Matriks penerimaan', 1)
    table(doc, ['Modul', 'Kriteria yang diuji', 'Keputusan'], [
        ('Rekap pembayaran', 'Voucher Santri dan Pejuang terpisah; Tunai/Transfer/QRIS digabung. Filter diterapkan sebelum rekap, paginasi dan penyediaan data ekspor.', 'LULUS'),
        ('Kasus tepi laporan', 'Pembayaran campuran lintas kelompok tidak dihitung ganda. Metode kosong/tidak dikenal tidak ditebak menjadi tunai. Respons gagal/terpotong tidak menghapus cache sah.', 'LULUS'),
        ('HPP produk', 'Nilai 5520 diubah 6000, tersimpan dan tetap terbaca setelah tutup/buka database serta kegagalan jaringan. Nilai nol asli berbeda dengan detail belum dimuat.', 'LULUS'),
        ('Migrasi dan antrean', 'Migrasi SQLite 20 ke 21; detail pending/gagal tidak tertimpa refresh. Simpan lokal dan outbox atomik, bertahan setelah restart.', 'LULUS'),
        ('Tanggal pelunasan', 'Tanggal lampau dipertahankan dalam payload antrean. Nominal tidak valid ditolak. Outbox tercatat sebelum permintaan jaringan.', 'LULUS'),
        ('Histori pembayaran', 'Hanya baris pembayaran, bukan penambahan piutang. Filter pelanggan/periode, urutan tanggal, jumlah, pagination, cache offline dan batas respons diuji.', 'LULUS'),
        ('Pemisahan varian', 'Tag, nama aset, identitas aplikasi/installer dan namespace terpisah. Updater Nahl lama menolak Al-Bahjah, termasuk nama aset yang mirip.', 'LULUS'),
    ], [1.6, 6.9, 1.2])
    doc.add_heading('Catatan interpretasi dan pemulihan', 1)
    doc.add_paragraph('HPP kosong pada laporan pengguna dapat berasal dari hilangnya detail pada cache lokal; pengujian ini tidak membuktikan semua HPP server sebelumnya utuh dan tidak mengarang pemulihan harga. Cache lama yang belum lengkap ditandai Belum dimuat dan pengeditan dicegah sampai detail tersedia. Histori menampilkan pembayaran terkonfirmasi; antrean belum sinkron diperiksa melalui menu Sinkronisasi.')
    doc.add_paragraph('Sebelum memasang, cadangkan data lokal dan pastikan antrean tersimpan. Jika modul inti gagal, hentikan distribusi varian terkait dan kumpulkan log tanpa menghapus cache/outbox. Jangan menurunkan versi langsung pada database v21; gunakan perbaikan maju atau pemulihan cadangan pra-upgrade yang telah direkonsiliasi. Perubahan ini tidak menyebarkan patch backend produksi.')
    for i, item in enumerate(evidence, 1):
        doc.add_page_break()
        doc.add_heading(f'{i:02d} {item["judul"]}', 1)
        doc.add_paragraph(STEPS.get(item['id'], item['judul']))
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(str(OUT / variant / item['screenshot']), width=Inches(8.05))
        p = doc.add_paragraph(f'LULUS — {item["hasil"]}')
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs: r.font.size = Pt(10)
    destination = OUT / variant / f'UAT-{variant}-1.34.40-build203.docx'
    doc.save(destination)
    print(destination)

if __name__ == '__main__':
    for item in sys.argv[1:] or ['albahjah', 'nahl']:
        build(item)
