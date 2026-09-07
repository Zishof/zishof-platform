"""Membangun dokumen UAT dan user manual Apotik E2E tanggal 8 September 2026."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/pos-apotik-emedik/uat-e2e-final-2026-09-08"
EVIDENCE = OUT / "evidence"
ASSETS = OUT / "assets"

UAT_DOCX = OUT / "DOKUMEN-UAT-E2E-APOTIK-2026-09-08.docx"
MANUAL_DOCX = OUT / "USER-MANUAL-E2E-APOTIK-2026-09-08.docx"

INK = "111827"
MUTED = "4B5563"
GREEN = "166534"
TEAL = "0F766E"
BLUE = "1D4ED8"
LIGHT = "F3F4F6"
PALE_GREEN = "F0FDF4"
PALE_AMBER = "FFFBEB"
LINE = "D1D5DB"


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _cell_borders(cell, color: str = LINE) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def _configure(doc: Document, title: str, kind: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        ("Title", 28, INK),
        ("Heading 1", 19, INK),
        ("Heading 2", 14, GREEN),
        ("Heading 3", 11, TEAL),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = f"APOTIK | {kind} | v1.35.1 (build 189) | 8 SEPTEMBER 2026"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(7.5)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Dokumen terkontrol • ")
    _field(footer, "PAGE")
    footer.add_run(" / ")
    _field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    doc.core_properties.title = title
    doc.core_properties.subject = "UAT end-to-end Apotik, Pengadaan, Keuangan, dan Akuntansi"
    doc.core_properties.author = "Tim UAT Apotik"
    doc.core_properties.comments = (
        "Bukti transaksi mutatif berasal dari batch UAT 7 September 2026; "
        "regresi 8 September 2026 dijalankan read-only pada server demo."
    )


def _cover(doc: Document, title: str, subtitle: str, badge: str) -> None:
    doc.add_paragraph("UAT LIVE SERVER DEMO • FINAL RERUN", style="Subtitle")
    p = doc.add_paragraph(title, style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph(subtitle)
    p.style = doc.styles["Subtitle"]
    p.add_run("\n\n")
    badge_p = doc.add_paragraph()
    badge_run = badge_p.add_run(badge)
    badge_run.bold = True
    badge_run.font.color.rgb = RGBColor.from_string(GREEN)
    badge_run.font.size = Pt(13)
    doc.add_paragraph(
        "Lingkungan: https://demo.ecampus.id/ecampus\n"
        "Aplikasi: Apotik v1.35.1 (build 189)\n"
        "Frontend: 676d5b2edee3f4b1fcec2880bb793610452c1851\n"
        "Tanggal eksekusi: 8 September 2026 (Asia/Jakarta)"
    )
    doc.add_paragraph(
        "Dokumen ini membedakan bukti transaksi mutatif yang sudah diselesaikan "
        "pada 7 September 2026 dari regresi live read-only pada 8 September 2026."
    )
    doc.add_page_break()


def _table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]], widths=None):
    rows = [list(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    _repeat_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        _shade(cell, GREEN)
        _cell_borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(8.5)
        if widths:
            cell.width = Inches(widths[i])
    for index, data in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(data):
            cell = row.cells[i]
            _cell_borders(cell)
            if index % 2:
                _shade(cell, LIGHT)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            run = p.add_run(str(value))
            run.font.size = Pt(8.2)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph()
    return table


def _status_box(doc: Document, title: str, body: str, fill: str = PALE_GREEN) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _shade(cell, fill)
    _cell_borders(cell, GREEN if fill == PALE_GREEN else "D97706")
    p = cell.paragraphs[0]
    run = p.add_run(f"{title}\n")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(GREEN if fill == PALE_GREEN else "92400E")
    p.add_run(body)
    doc.add_paragraph()


def _picture(
    doc: Document,
    path: Path,
    caption: str,
    width: float = 6.9,
    max_height: float | None = None,
) -> None:
    if not path.exists():
        p = doc.add_paragraph(f"Bukti belum tersedia: {path.name}")
        p.runs[0].font.color.rgb = RGBColor.from_string("B91C1C")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    with Image.open(path) as source:
        image_width, image_height = source.size
    rendered_height = width * image_height / image_width
    if max_height is not None and rendered_height > max_height:
        p.add_run().add_picture(str(path), height=Inches(max_height))
    else:
        p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(8)
    c.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def _step_page(
    doc: Document,
    title: str,
    image: Path,
    steps: Iterable[str],
    expected: str,
    note: str = "",
) -> None:
    doc.add_heading(title, level=2)
    _picture(doc, image, f"Bukti layar — {title}", max_height=4.35)
    doc.add_paragraph("Langkah operator", style="Heading 3")
    for index, value in enumerate(steps, 1):
        doc.add_paragraph(f"{index}. {value}")
    _status_box(doc, "Hasil yang diharapkan", expected)
    if note:
        doc.add_paragraph(note)
    doc.add_page_break()


def _uat_evidence_page(
    doc: Document, title: str, image: Path, actual: str, verdict: str = "PASS"
) -> None:
    doc.add_heading(title, level=2)
    _picture(doc, image, f"Screenshot aktual — {title}", max_height=5.25)
    _table(
        doc,
        ["Kriteria", "Hasil aktual", "Putusan"],
        [["Layar dapat dimuat dan data dapat ditelusuri", actual, verdict]],
        widths=[2.2, 3.6, 0.8],
    )
    doc.add_page_break()


def build_uat() -> Path:
    doc = Document()
    _configure(doc, "Dokumen UAT E2E Apotik — 8 September 2026", "DOKUMEN UAT E2E")
    _cover(
        doc,
        "Dokumen UAT End-to-End Apotik",
        "Kasir • Kulakan • Keuangan • Akuntansi • Laporan",
        "PUTUSAN: PASS DENGAN OBSERVASI",
    )

    doc.add_heading("1. Ringkasan Eksekutif", level=1)
    doc.add_paragraph(
        "Alur utama Apotik telah tervalidasi dari penjualan di kasir hingga laporan. "
        "Batch transaksi UAT 7 September 2026 membuktikan minimal 100 data untuk setiap "
        "mode kasir dan setiap tahap kulakan. Rerun 8 September 2026 memverifikasi UI "
        "modern, akses responsif, data live, status posting, dan laporan tanpa membuat "
        "transaksi finansial baru."
    )
    _table(
        doc,
        ["Area", "Bukti volume", "Hasil"],
        [
            ["OTC / Obat Bebas", "100 transaksi + katalog live 100/11.000", "PASS"],
            ["Resep Dokter", "100 transaksi + 1.000 resep klinis sample", "PASS"],
            ["Racikan", "100 transaksi + 500 formula", "PASS"],
            ["Produksi Farmasi", "100 produksi + 500 formula", "PASS"],
            ["Tebus Resep", "100 transaksi campuran + 400 antrean", "PASS"],
            ["PR → PO → BAST → Tagihan → Bayar", "100 data per tahap", "PASS"],
            ["Jurnal run-specific", "1.100 jurnal terposting; retry 0 duplikasi", "PASS"],
            ["Laporan", "Penjualan, pembelian, dan 6 laporan akuntansi", "PASS"],
            ["Responsif", "1920, 1366, 768, dan 390 px", "PASS"],
            ["Full regression", "1.073/1.073 pengujian Flutter", "PASS"],
            ["Analisis statis", "exit 0; 0 error/warning; 51 info", "PASS"],
        ],
        widths=[2.45, 3.35, 0.8],
    )
    _status_box(
        doc,
        "Batas aman rerun",
        "Rerun 8 September bersifat read-only. Posting ulang dan closing tidak dijalankan "
        "karena batch final sudah terposting dan tindakan tersebut dapat memengaruhi data global.",
        PALE_AMBER,
    )

    doc.add_heading("2. Lingkup dan Metode", level=1)
    _table(
        doc,
        ["Jenis bukti", "Tanggal", "Metode", "Keterangan"],
        [
            ["Transaksi volume", "7 Sep 2026", "API live terisolasi", "100 per proses, kode run-specific"],
            ["Posting jurnal", "7 Sep 2026", "API live + verifikasi retry", "1.100 terposting, 0 duplikasi"],
            ["Regresi tampilan", "8 Sep 2026", "Flutter integration test", "read-only, screenshot aktual"],
            ["Regresi kode", "8 Sep 2026", "flutter test + analyze", "hasil tersimpan di evidence/results"],
        ],
        widths=[1.5, 1.15, 1.65, 2.3],
    )
    doc.add_heading("3. Arsitektur Proses dan Data", level=1)
    _picture(doc, ASSETS / "flow-e2e-apotik.png", "Flow E2E Apotik sampai laporan")
    _picture(doc, ASSETS / "erd-kasir-apotik.png", "ERD logis kasir, resep, formula, batch, pembayaran, dan jurnal")
    _picture(doc, ASSETS / "erd-pengadaan-akuntansi.png", "ERD logis pengadaan, hutang vendor, pembayaran, dan akuntansi")
    doc.add_page_break()

    doc.add_heading("4. Bukti Kasir dan UI Responsif", level=1)
    cashier_dir = EVIDENCE / "screenshots-kasir-live"
    modern_dir = EVIDENCE / "screenshots-modern-ui"
    for title, filename, actual in (
        ("4.1 OTC / Obat Bebas", "01-live-otc-desktop.png", "100 produk tampil; total server 11.000; FEFO dan batch tersedia."),
        ("4.2 Resep Dokter", "02-live-resep-dokter.png", "100 resep dapat ditelusuri; informasi klinis sample tersedia."),
        ("4.3 Racikan", "03-live-racikan.png", "100 formula tampil dari total 500 formula racikan."),
        ("4.4 Produksi Farmasi", "04-live-produksi-farmasi.png", "100 formula tampil dari total 500 formula produksi."),
        ("4.5 Tebus Resep", "05-live-tebus-resep.png", "100 resep siap ditelusuri untuk proses tebus resep."),
        ("4.6 Tebus Resep pada layar kecil", "06-live-mobile-tebus-resep-terlihat.png", "Menu Tebus Resep terlihat pada viewport mobile."),
    ):
        _uat_evidence_page(doc, title, cashier_dir / filename, actual)
    baseline_cashier = EVIDENCE / "screenshots-baseline-20260907-kasir"
    for title, filename, actual in (
        ("4.6a Pembayaran OTC", "01c-pembayaran-obat-bebas.png", "Dialog pembayaran hasil batch UAT final; kode transaksi terinisialisasi dan pembayaran berhasil."),
        ("4.6b Detail klinis resep", "02a-tebus-resep-detail.png", "Pasien, dokter, asal pelayanan, diagnosis, alergi, dan pemeriksaan kedua tercatat."),
        ("4.6c Konfirmasi hasil produksi", "02b-produksi-farmasi-konfirmasi.png", "Nomor batch dan kedaluwarsa hasil produksi dikonfirmasi sebelum proses atomik."),
        ("4.6d Laporan Penjualan", "09-laporan-penjualan.png", "Transaksi kasir dapat ditelusuri pada laporan penjualan."),
    ):
        _uat_evidence_page(doc, title, baseline_cashier / filename, actual)
    for title, filename, actual in (
        ("4.7 Desktop 1920 px", "01-kasir-modern-desktop-1920.png", "Tiga-panel desktop terbaca dan area aksi tidak tertutup."),
        ("4.8 Laptop 1366 px", "02-kasir-modern-laptop-1366.png", "Navigasi, katalog, dan keranjang tetap dapat digunakan."),
        ("4.9 Tablet 768 px", "03-kasir-modern-tablet-768.png", "Mode transaksi dapat digulir dan Tebus Resep dapat diakses."),
        ("4.10 Mobile 390 px", "04-kasir-modern-mobile-390.png", "Konten satu kolom, tanpa tombol kritis hilang."),
        ("4.11 Keranjang mobile", "05-kasir-modern-mobile-keranjang-390.png", "Ringkasan dan aksi pembayaran dapat dibuka pada mobile."),
    ):
        _uat_evidence_page(doc, title, modern_dir / filename, actual)

    doc.add_heading("5. Bukti Kulakan End-to-End", level=1)
    proc_dir = EVIDENCE / "screenshots-pengadaan"
    for title, filename, actual in (
        ("5.1 Menu Pengadaan", "01-menu-pengadaan-terbuka.png", "Seluruh menu rantai pengadaan tersedia."),
        ("5.2 Permintaan Pembelian (PR)", "02-pr-daftar-100.png", "264 PR live; 260 disetujui; tahapan dapat ditelusuri."),
        ("5.3 Form PR", "03-pr-formulir.png", "Field kebutuhan, penggunaan, dan item tersedia."),
        ("5.4 Pemesanan Pembelian (PO)", "04-po-daftar-termin-nontermin.png", "254 PO live; termin dan non-termin didukung."),
        ("5.5 Form PO non-termin", "05-po-formulir-nontermin.png", "Skema pembayaran langsung tersedia."),
        ("5.6 Form PO termin", "06-po-formulir-termin.png", "Rincian termin dapat diatur dan divalidasi."),
        ("5.7 Penerimaan Barang (BAST)", "07-bast-daftar-100.png", "255 BAST live; seluruhnya disetujui dan masuk stok."),
        ("5.8 Referensi PO pada BAST", "08-bast-pilih-po.png", "PO sumber dapat dipilih dari dialog."),
        ("5.9 Terima Tagihan Vendor", "09-terima-tagihan-100.png", "255 tagihan live dan terhubung ke penerimaan."),
        ("5.10 Pembayaran Vendor", "10-pembayaran-vendor-100.png", "Riwayat transfer tersedia; batch UAT 100 telah terealisasi."),
        ("5.11 Draft Jurnal Pengadaan", "11-draft-jurnal-pengadaan.png", "Sumber jurnal pengadaan dapat diaudit."),
        ("5.12 Katalog Laporan", "12-katalog-laporan-pengadaan.png", "63 laporan tersedia pada 5 halaman."),
        ("5.13 Laporan Pembelian", "13-laporan-pembelian-100.png", "Filter periode dan opsi transaksi lunas tersedia."),
    ):
        _uat_evidence_page(doc, title, proc_dir / filename, actual)

    doc.add_heading("6. Bukti Akuntansi dan Laporan", level=1)
    acc_dir = EVIDENCE / "screenshots-akuntansi-submenu"
    for title, filename, actual, verdict in (
        ("6.1 Draft Jurnal", "08-draft-jurnal.png", "6.058 draft global dan 5.295 terposting; batch UAT final tidak menyisakan draft.", "PASS"),
        ("6.2 Posting HPP", "09-posting-hpp.png", "400 HPP run-specific telah terposting; draft lama tidak memenuhi syarat dan tidak disentuh.", "PASS"),
        ("6.3 Posting Penjualan", "10-posting-penjualan.png", "400 jurnal penjualan run-specific telah terposting; retry 0.", "PASS"),
        ("6.4 Posting Kulakan", "11-posting-kulakan.png", "100/100 jurnal BAST run-specific terposting; ready 0.", "PASS"),
        ("6.5 Posting Bayar Hutang", "12-posting-bayar-hutang.png", "100/100 pembayaran vendor run-specific terposting; ready 0.", "PASS"),
        ("6.6 Posting Terima Piutang", "13-posting-terima-piutang.png", "Belum memiliki endpoint khusus Apotik; tidak diperlukan untuk alur tunai/pembelian yang diuji.", "N/A"),
        ("6.7 Tutup Buku", "16-tutup-buku.png", "Preview tersedia; akun Laba Ditahan belum dikonfigurasi.", "OBSERVASI"),
        ("6.8 Closing", "17-closing.png", "Form preview terbuka; closing tidak dieksekusi karena irreversible.", "N/A"),
        ("6.9 Kode Akun", "20-kode-akun.png", "Daftar COA dapat dimuat dan ditelusuri.", "PASS"),
        ("6.10 Grup Akun", "21-grup-akun.png", "Klasifikasi akun dapat dimuat.", "PASS"),
        ("6.11 Jenis Transaksi", "22-jenis-transaksi.png", "Jenis transaksi dan akun default dapat dimuat.", "PASS"),
        ("6.12 Bank", "23-bank.png", "Master bank dan akun kas dapat dimuat.", "PASS"),
    ):
        _uat_evidence_page(doc, title, acc_dir / filename, actual, verdict)

    reports_dir = EVIDENCE / "screenshots-akuntansi-laporan"
    for title, filename, actual in (
        ("6.13 Laba Rugi", "24-laporan-laba-rugi.png", "Laporan berbasis jurnal terposting dapat ditampilkan."),
        ("6.14 Neraca", "25-laporan-neraca.png", "Neraca berbasis jurnal terposting dapat ditampilkan."),
        ("6.15 Arus Kas", "26-laporan-arus-kas.png", "Arus kas berbasis jurnal terposting dapat ditampilkan."),
        ("6.16 Keseluruhan Jurnal", "27-laporan-jurnal-umum.png", "Jurnal terposting dapat ditelusuri dari laporan."),
        ("6.17 Buku Besar", "28-laporan-buku-besar.png", "Mutasi akun dapat ditelusuri."),
        ("6.18 Neraca Saldo", "29-laporan-neraca-saldo.png", "Saldo akun dapat ditampilkan untuk rekonsiliasi."),
    ):
        _uat_evidence_page(doc, title, reports_dir / filename, actual)

    journal_dir = EVIDENCE / "screenshots-akuntansi-jurnal"
    _uat_evidence_page(
        doc,
        "6.19 Verifikasi jurnal run-specific terposting",
        journal_dir / "23b-jurnal-umum-terposting-final.png",
        "Referensi UAT-APT-E2E-FINAL-20260907-JU ditemukan pada status Terposting, tetapi layar memakai fallback salinan lokal; laporan Keseluruhan Jurnal live tetap berhasil.",
        "OBSERVASI",
    )

    doc.add_heading("7. Rekonsiliasi dan Traceability", level=1)
    _table(
        doc,
        ["Sumber", "Jumlah", "Status", "Idempotensi"],
        [
            ["Penjualan Apotik", "400", "Terposting", "0 jurnal baru saat retry"],
            ["HPP Apotik", "400", "Terposting", "0 jurnal baru saat retry"],
            ["BAST / persediaan", "100", "Terposting", "Tidak ganda"],
            ["Pembayaran vendor", "100", "Terposting", "0 jurnal baru saat retry"],
            ["Jurnal Umum UAT", "100", "Terposting", "0 jurnal baru saat retry"],
            ["Total", "1.100", "Selesai", "PASS"],
        ],
        widths=[2.0, 1.0, 1.3, 2.3],
    )
    doc.add_heading("8. Observasi dan Risiko Tersisa", level=1)
    _table(
        doc,
        ["ID", "Observasi", "Dampak", "Tindak lanjut"],
        [
            ["OBS-01", "3.600 draft HPP/Penjualan historis global tertahan karena transaksi pembayaran lama tidak ada/nihil.", "Tidak memengaruhi batch UAT final.", "Audit dan bersihkan melalui prosedur data historis tersendiri."],
            ["OBS-02", "Posting Terima Piutang belum punya endpoint khusus Apotik.", "N/A untuk alur tunai dan kulakan yang diuji.", "Implementasikan sebelum penjualan kredit/piutang Apotik dipakai produksi."],
            ["OBS-03", "Akun Laba Ditahan belum diatur pada master Toko.", "Tutup Buku tidak dapat diselesaikan.", "Tetapkan akun melalui otorisasi Akuntansi sebelum closing pertama."],
            ["OBS-04", "Closing tidak dijalankan dalam UAT.", "Tidak ada periode yang dikunci.", "Lakukan rehearsal terotorisasi di lingkungan khusus."],
            ["OBS-05", "Layar Jurnal Umum memakai fallback salinan lokal setelah tiga reload.", "Status tampil dapat terlambat terhadap server.", "Periksa endpoint daftar jurnal; gunakan laporan Keseluruhan Jurnal sebagai bukti live sementara."],
        ],
        widths=[0.75, 2.7, 1.45, 1.7],
    )
    _status_box(
        doc,
        "Kesimpulan",
        "Alur inti Kasir Apotik → Kulakan → Pembayaran Vendor → Posting → Laporan lulus. "
        "Status keseluruhan PASS DENGAN OBSERVASI; observasi tidak membatalkan fungsi inti "
        "yang diminta, tetapi harus ditangani sebelum closing dan penjualan kredit produksi.",
    )
    doc.add_heading("9. Persetujuan", level=1)
    _table(
        doc,
        ["Peran", "Nama", "Tanggal", "Keputusan / catatan"],
        [
            ["Operator UAT", "", "", ""],
            ["Apoteker Penanggung Jawab", "", "", ""],
            ["Pengadaan", "", "", ""],
            ["Keuangan", "", "", ""],
            ["Akuntansi", "", "", ""],
            ["Pemilik Sistem", "", "", ""],
        ],
        widths=[1.9, 1.7, 1.1, 1.9],
    )
    doc.save(UAT_DOCX)
    return UAT_DOCX


def build_manual() -> Path:
    doc = Document()
    _configure(doc, "User Manual E2E Apotik — 8 September 2026", "USER MANUAL E2E")
    _cover(
        doc,
        "User Manual End-to-End Apotik",
        "Panduan operasional dari kasir sampai laporan akuntansi",
        "EDISI DEPLOYMENT v1.35.1 (build 189)",
    )
    doc.add_heading("1. Tujuan dan Peran Pengguna", level=1)
    doc.add_paragraph(
        "Manual ini memandu operator menjalankan proses Apotik dari penjualan, produksi, "
        "tebus resep, pengadaan, pembayaran vendor, posting jurnal, hingga laporan. "
        "Hak akses maker-checker dan kebijakan Apoteker tetap berlaku."
    )
    _table(
        doc,
        ["Peran", "Tanggung jawab utama", "Menu"],
        [
            ["Kasir", "Pilih obat, batch, identitas, dan pembayaran", "Kasir Apotik"],
            ["Apoteker", "Telaah resep, racikan, produksi, QC, serah obat", "Resep/Racikan/Produksi"],
            ["Pengadaan", "PR, PO, BAST, tagihan", "Pengadaan"],
            ["Keuangan", "Realisasi pembayaran vendor", "Proses Transfer"],
            ["Akuntansi", "Audit akun, posting, jurnal manual, laporan", "Akuntansi"],
        ],
        widths=[1.3, 3.15, 2.15],
    )
    doc.add_heading("2. Alur Utama", level=1)
    _picture(doc, ASSETS / "flow-e2e-apotik.png", "Alur E2E Apotik")
    _status_box(
        doc,
        "Aturan keselamatan data",
        "Gunakan satu tombol proses sekali, tunggu respons, dan pakai retry/idempotency yang "
        "disediakan. Jangan mem-posting draft historis atau melakukan closing tanpa mandat."
    )
    doc.add_page_break()

    cashier = EVIDENCE / "screenshots-kasir-live"
    modern = EVIDENCE / "screenshots-modern-ui"
    baseline_cashier = EVIDENCE / "screenshots-baseline-20260907-kasir"
    doc.add_heading("3. Kasir Apotik", level=1)
    _step_page(doc, "3.1 Memahami layout desktop", modern / "01-kasir-modern-desktop-1920.png", ["Masuk dan pilih toko aktif.", "Periksa identitas Apotek dan kasir.", "Pastikan mode transaksi, katalog, dan keranjang tampil."], "Halaman siap digunakan tanpa error dan tanpa fitur terkunci.")
    _step_page(doc, "3.2 Menggunakan layar laptop", modern / "02-kasir-modern-laptop-1366.png", ["Pastikan skala tampilan normal.", "Gunakan tab mode transaksi di bagian atas.", "Buka keranjang melalui area ringkasan bila panel tidak permanen."], "Semua fungsi penting tetap terlihat pada 1366 px.")
    _step_page(doc, "3.3 Menggunakan tablet dan mobile", modern / "03-kasir-modern-tablet-768.png", ["Gulir tab mode transaksi secara horizontal bila perlu.", "Pilih item dari daftar satu/dua kolom.", "Buka ringkasan keranjang dari tombol bawah."], "Tebus Resep dan pembayaran tetap dapat diakses.")
    _step_page(doc, "3.4 Memastikan Tebus Resep tampil di mobile", cashier / "06-live-mobile-tebus-resep-terlihat.png", ["Buka Kasir Apotik pada layar kecil.", "Gulir daftar mode bila diperlukan.", "Pilih Tebus Resep."], "Menu Tebus Resep terlihat dan dapat dipilih pada viewport mobile.")
    _step_page(doc, "3.5 Menjual OTC / Obat Bebas", cashier / "01-live-otc-desktop.png", ["Pilih OTC / Obat Bebas.", "Cari nama/kode atau pindai barcode.", "Pilih obat dan tentukan jumlah.", "Pilih batch FEFO yang layak.", "Periksa total dan lanjutkan ke pembayaran."], "Transaksi tersimpan, stok batch berkurang atomik, dan bukti pembayaran tersedia.")
    _step_page(doc, "3.5a Menyelesaikan pembayaran OTC", baseline_cashier / "01c-pembayaran-obat-bebas.png", ["Periksa total dan metode pembayaran.", "Isi uang diterima atau referensi non-tunai sesuai metode.", "Pastikan kembalian benar.", "Tekan Bayar satu kali dan tunggu hasil."], "Kode transaksi Apotik terinisialisasi, pembayaran berhasil, dan retry tidak menggandakan transaksi.", "Screenshot berasal dari batch mutatif final 7 September; rerun 8 September tidak mengulang pembayaran agar data tidak ganda.")
    _step_page(doc, "3.6 Menjual berdasarkan Resep Dokter", cashier / "02-live-resep-dokter.png", ["Pilih Resep Dokter.", "Cari atau pilih resep/pasien.", "Periksa dokter, fasilitas asal, diagnosis, indikasi, alergi, dan catatan klinis.", "Pilih item dan batch.", "Lakukan telaah serta pembayaran."], "Resep terlacak ke pasien/dokter dan penjualan selesai tanpa menghilangkan informasi klinis.")
    _step_page(doc, "3.6a Menelaah informasi klinis resep", baseline_cashier / "02a-tebus-resep-detail.png", ["Cocokkan pasien dan nomor rekam medis.", "Periksa dokter, SIP/kontak, fasilitas asal, dan tanggal resep.", "Baca diagnosis, indikasi, keluhan, anamnesis, alergi, dan terapi.", "Dokumentasikan klarifikasi serta pemeriksaan kedua."], "Resep lengkap, seolah nyata tetapi tetap ditandai sebagai data sample/UAT.")
    _step_page(doc, "3.7 Menjual Racikan", cashier / "03-live-racikan.png", ["Pilih Racikan.", "Cari formula aktif.", "Periksa komposisi dan jumlah sediaan.", "Konfirmasi batch bahan menurut FEFO.", "Tambahkan jasa/kemasan bila berlaku, lalu bayar."], "Semua bahan cukup, stok terpotong sesuai formula, dan transaksi tercatat.")
    _step_page(doc, "3.8 Menjalankan Produksi Farmasi", cashier / "04-live-produksi-farmasi.png", ["Pilih Produksi Farmasi.", "Cari formula barang jadi.", "Periksa bahan, jumlah hasil, metode, dan QC.", "Isi nomor batch serta kedaluwarsa hasil.", "Proses produksi sekali."], "Bahan berkurang dan batch barang jadi terbentuk dalam satu transaksi.")
    _step_page(doc, "3.8a Mengonfirmasi batch hasil produksi", baseline_cashier / "02b-produksi-farmasi-konfirmasi.png", ["Periksa nomor batch hasil.", "Isi atau verifikasi tanggal kedaluwarsa.", "Periksa jumlah hasil dan bahan yang akan dikonsumsi.", "Konfirmasi proses sekali."], "Batch barang jadi tercipta dan tetap dapat ditelusuri ke formula serta batch bahan.")
    _step_page(doc, "3.9 Menebus Resep", cashier / "05-live-tebus-resep.png", ["Pilih Tebus Resep.", "Pilih antrean yang berstatus siap.", "Cocokkan pasien, dokter, asal resep, diagnosis, alergi, dan item.", "Lakukan pemeriksaan kedua dan konseling.", "Bayar dan tandai diserahkan."], "Obat jadi dan racikan ditebus atomik; status resep berubah dan stok konsisten.")
    _step_page(doc, "3.10 Membuka keranjang mobile", modern / "05-kasir-modern-mobile-keranjang-390.png", ["Tambahkan item.", "Tekan ringkasan keranjang bawah.", "Ubah jumlah atau hapus item bila perlu.", "Periksa total sebelum Bayar."], "Keranjang dapat dioperasikan tanpa tertutup viewport.")
    _step_page(doc, "3.11 Membaca Laporan Penjualan", baseline_cashier / "09-laporan-penjualan.png", ["Buka Laporan Apotik.", "Atur periode dan toko.", "Cari nomor transaksi bila diperlukan.", "Cocokkan total, metode pembayaran, kasir, dan status."], "Transaksi dapat ditelusuri dari laporan ke bukti pembayaran.")

    doc.add_heading("4. Kulakan dan Pembayaran Vendor", level=1)
    proc = EVIDENCE / "screenshots-pengadaan"
    _step_page(doc, "4.1 Membuka rantai Pengadaan", proc / "01-menu-pengadaan-terbuka.png", ["Buka bagian PENGADAAN.", "Pastikan menu PR, PO, BAST, dan Terima Tagihan tampil.", "Pastikan menu pembayaran tersedia pada KEUANGAN."], "Rantai proses lengkap dapat dinavigasi.")
    _step_page(doc, "4.2 Membuat Permintaan Pembelian (PR)", proc / "03-pr-formulir.png", ["Buka Permintaan Pembelian dan pilih Tambah.", "Isi alasan, penggunaan, dan tanggal kebutuhan.", "Tambahkan item beserta jumlah.", "Simpan sebagai draft, periksa, lalu ajukan untuk persetujuan."], "PR memperoleh nomor dan mengikuti maker-checker.")
    _step_page(doc, "4.3 Memantau daftar PR", proc / "02-pr-daftar-100.png", ["Gunakan pencarian nomor/keterangan.", "Filter status dan periode.", "Buka PR untuk melihat riwayat persetujuan dan tahap berikutnya."], "Status PR dan progres ke PO dapat ditelusuri.")
    _step_page(doc, "4.4 Membuat PO non-termin", proc / "05-po-formulir-nontermin.png", ["Pilih PR yang disetujui.", "Pilih supplier dan alamat kirim.", "Konfirmasi harga, pajak, diskon, dan jadwal.", "Biarkan pembayaran termin nonaktif.", "Simpan dan ajukan persetujuan."], "PO non-termin dibuat dari PR tanpa input item berulang.")
    _step_page(doc, "4.5 Membuat PO bertermin", proc / "06-po-formulir-termin.png", ["Aktifkan pembayaran bertermin.", "Tambah setiap termin beserta tanggal dan nominal/persentase.", "Pastikan total termin sama dengan total PO.", "Simpan dan ajukan."], "Jadwal kewajiban vendor terbentuk konsisten.")
    _step_page(doc, "4.6 Memantau daftar PO", proc / "04-po-daftar-termin-nontermin.png", ["Cari nomor PO atau supplier.", "Periksa nilai, dibayar, dan sisa.", "Buka detail sebelum penerimaan barang."], "PO termin/non-termin dan sisa kewajiban terlihat.")
    _step_page(doc, "4.7 Menerima Barang (BAST)", proc / "08-bast-pilih-po.png", ["Buka Penerimaan Barang dan pilih Dari PO.", "Pilih PO yang masih memiliki sisa penerimaan.", "Cocokkan jumlah fisik, kondisi, batch, dan tanggal kedaluwarsa.", "Catat selisih bila ada.", "Simpan/ajukan sesuai kewenangan."], "BAST terhubung ke PO dan stok bertambah setelah disetujui.")
    _step_page(doc, "4.8 Memantau BAST", proc / "07-bast-daftar-100.png", ["Filter periode/status.", "Periksa status persetujuan dan masuk stok.", "Buka detail untuk audit batch."], "Penerimaan dapat ditelusuri sampai pergerakan stok.")
    _step_page(doc, "4.9 Menerima Tagihan Vendor", proc / "09-terima-tagihan-100.png", ["Buka Terima Tagihan Vendor.", "Pilih referensi BAST/PO.", "Isi nomor faktur, tanggal, jatuh tempo, pajak, dan total.", "Lampirkan dokumen bila diwajibkan.", "Simpan dan verifikasi."], "Tagihan vendor terbentuk dan terkait ke dokumen sumber.")
    _step_page(doc, "4.10 Membayar Vendor", proc / "10-pembayaran-vendor-100.png", ["Buka KEUANGAN > Proses Transfer.", "Pilih Pembayaran Vendor.", "Pilih tagihan yang disetujui.", "Cocokkan rekening, nilai, referensi, dan tanggal.", "Realisasikan sesuai maker-checker."], "Status pembayaran menjadi Terealisasi dan jurnal pembayaran tersedia untuk posting.")
    _step_page(doc, "4.11 Membuka Laporan Pembelian", proc / "13-laporan-pembelian-100.png", ["Buka laporan pembelian/AP.", "Atur periode dan supplier.", "Aktifkan Tampilkan Lunas bila ingin melihat transaksi yang sudah selesai.", "Cocokkan faktur, total, dibayar, dan sisa."], "Pembelian dapat ditelusuri dari laporan ke dokumen sumber.")

    doc.add_heading("5. Posting dan Jurnal Umum", level=1)
    acc = EVIDENCE / "screenshots-akuntansi-submenu"
    _step_page(doc, "5.1 Audit Draft Jurnal", acc / "08-draft-jurnal.png", ["Buka AKUNTANSI > Draft Jurnal.", "Pilih kategori dan periode.", "Periksa referensi sumber dan pemetaan akun.", "Pisahkan batch operasional dari data historis."], "Hanya draft sah dan siap yang dilanjutkan ke posting.")
    _step_page(doc, "5.2 Posting HPP", acc / "09-posting-hpp.png", ["Pilih periode transaksi.", "Periksa Persediaan dan HPP.", "Pilih hanya referensi yang statusnya Siap.", "Posting dan tunggu konfirmasi."], "Jurnal HPP seimbang serta berstatus Terposting.", "Draft lama bertanda Tertahan harus diperbaiki dari transaksi sumber; jangan diposting paksa.")
    _step_page(doc, "5.3 Posting Penjualan", acc / "10-posting-penjualan.png", ["Pilih periode.", "Cocokkan pembayaran dengan pendapatan.", "Pilih referensi Siap.", "Posting sekali."], "Jurnal Kas/Piutang versus Pendapatan berstatus Terposting.")
    _step_page(doc, "5.4 Posting Kulakan", acc / "11-posting-kulakan.png", ["Pilih periode BAST.", "Cocokkan Persediaan dan Hutang Vendor/Sementara.", "Posting hanya BAST disetujui."], "Jurnal persediaan kulakan berstatus Terposting.")
    _step_page(doc, "5.5 Posting Bayar Hutang", acc / "12-posting-bayar-hutang.png", ["Pilih periode pembayaran.", "Cocokkan Hutang Vendor dan Kas/Bank.", "Posting transaksi terealisasi."], "Jurnal pembayaran vendor berstatus Terposting.")
    _step_page(doc, "5.6 Menelusuri Jurnal Umum", EVIDENCE / "screenshots-akuntansi-jurnal/23b-jurnal-umum-terposting-final.png", ["Buka Jurnal Umum.", "Cari kode/keterangan atau referensi run.", "Tekan Terapkan.", "Pastikan status Terposting dan buka rincian debet/kredit.", "Jika penanda salinan lokal muncul, gunakan Muat Ulang dan cocokkan dengan laporan Keseluruhan Jurnal live."], "Referensi dapat ditelusuri; penanda cache harus diperlakukan sebagai observasi sinkronisasi, bukan bukti live final.")

    doc.add_heading("6. Laporan Akuntansi", level=1)
    reports = EVIDENCE / "screenshots-akuntansi-laporan"
    for title, filename, steps, expected in (
        ("6.1 Laba Rugi", "24-laporan-laba-rugi.png", ["Pilih laporan Laba Rugi berbasis jurnal.", "Atur periode.", "Tampilkan dan cocokkan Pendapatan, HPP/Beban, serta laba bersih."], "Laba bersih dihitung dari jurnal terposting."),
        ("6.2 Neraca", "25-laporan-neraca.png", ["Pilih Neraca berbasis jurnal.", "Atur tanggal sampai.", "Periksa Aset, Liabilitas, Ekuitas, dan Selisih."], "Persamaan akuntansi dapat direkonsiliasi."),
        ("6.3 Arus Kas", "26-laporan-arus-kas.png", ["Pilih Arus Kas.", "Atur periode.", "Periksa saldo awal, penerimaan, pengeluaran, dan saldo akhir."], "Arus kas dapat ditelusuri ke akun Kas/Bank."),
        ("6.4 Keseluruhan Jurnal", "27-laporan-jurnal-umum.png", ["Pilih laporan Keseluruhan Jurnal.", "Atur periode dan filter referensi.", "Periksa tanggal, akun, debet, kredit, dan sumber."], "Jurnal terposting tampil lengkap."),
        ("6.5 Buku Besar", "28-laporan-buku-besar.png", ["Pilih Rincian Buku Besar.", "Pilih akun/periode.", "Telusuri mutasi dan saldo berjalan."], "Setiap mutasi dapat ditelusuri ke jurnal sumber."),
        ("6.6 Neraca Saldo", "29-laporan-neraca-saldo.png", ["Pilih Neraca Saldo.", "Atur tanggal.", "Bandingkan total debet dan kredit."], "Total debet dan kredit seimbang."),
    ):
        _step_page(doc, title, reports / filename, steps, expected)

    doc.add_heading("7. Tutup Buku dan Closing", level=1)
    _picture(doc, acc / "16-tutup-buku.png", "Preview Tutup Buku")
    doc.add_paragraph(
        "Sebelum menjalankan Tutup Buku, Akuntansi wajib menetapkan Akun Laba Ditahan pada "
        "master Toko, memastikan seluruh jurnal periode sudah terposting, melakukan backup, "
        "dan memperoleh persetujuan. Closing bersifat irreversible dan tidak dijalankan dalam UAT ini."
    )
    _status_box(doc, "Peringatan", "Jangan menjalankan Tutup Buku atau Closing di data bersama tanpa mandat tertulis dan rencana rollback.", PALE_AMBER)

    doc.add_heading("8. Troubleshooting", level=1)
    _table(
        doc,
        ["Gejala", "Pemeriksaan", "Tindakan"],
        [
            ["Tombol Bayar tidak aktif", "Keranjang, batch, sesi kas, metode", "Lengkapi field wajib dan muat ulang bila sesi berubah."],
            ["Batch tidak dapat dipilih", "Status tahan/kedaluwarsa dan stok", "Pilih batch FEFO lain; jangan paksa batch tidak layak."],
            ["Draft posting Tertahan", "Transaksi sumber/pembayaran nihil", "Perbaiki sumber atau lakukan audit data historis."],
            ["Tebus Resep tidak terlihat", "Lebar viewport dan scroll tab", "Gulir tab mode; gunakan versi 1.35.1 build 189 atau lebih baru."],
            ["Tutup Buku gagal", "Akun Laba Ditahan", "Konfigurasikan akun dengan persetujuan Akuntansi."],
            ["Data belum berubah", "Koneksi dan status sinkronisasi", "Muat ulang, periksa antrean lokal, lalu retry dengan kode idempoten yang sama."],
        ],
        widths=[1.7, 2.2, 2.7],
    )
    doc.add_page_break()
    doc.add_heading("9. Checklist Operasional", level=1)
    _table(
        doc,
        ["Waktu", "Checklist"],
        [
            ["Awal shift", "Toko benar; sesi kas terbuka; printer/barcode; sinkronisasi; batch kritis."],
            ["Saat penjualan", "Identitas/resep; FEFO; alergi; telaah; metode pembayaran; serah obat."],
            ["Saat kulakan", "Otorisasi; supplier; jumlah; batch/ED; dokumen sumber; pajak; jatuh tempo."],
            ["Akhir shift", "Rekonsiliasi kas; transaksi gagal; draft jurnal; selisih stok; serah terima."],
            ["Akhir periode", "Semua posting; laporan; backup; persetujuan; baru Tutup Buku/Closing."],
        ],
        widths=[1.4, 5.2],
    )
    doc.save(MANUAL_DOCX)
    return MANUAL_DOCX


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    print(build_uat())
    print(build_manual())
