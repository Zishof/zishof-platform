from __future__ import annotations

import datetime as dt
import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "apotik-uiux-v2" / "uat-2026-09-08"
SHOTS = OUT / "screenshots"
ASSETS = OUT / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)

TEAL = "0D756D"
GREEN = "154F3B"
PURPLE = "6A4BB2"
INK = "172033"
MUTED = "5D687C"
PALE = "EEF7F5"
PALE_PURPLE = "F1EDFB"
PALE_RED = "FFF0F0"
WHITE = "FFFFFF"
LINE = "D7DEE8"


def font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        Path(r"C:\Windows\Fonts\aptos.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    if bold:
        candidates = [
            Path(r"C:\Windows\Fonts\aptos-display-bold.ttf"),
            Path(r"C:\Windows\Fonts\arialbd.ttf"),
        ] + candidates
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def annotate(
    source: Path,
    target: Path,
    label: str,
    boxes: list[tuple[int, int, int, int]],
    crop: tuple[int, int, int, int] | None = None,
) -> None:
    image = Image.open(source).convert("RGB")
    if crop:
        image = image.crop(crop)
        boxes = [
            (x1 - crop[0], y1 - crop[1], x2 - crop[0], y2 - crop[1])
            for x1, y1, x2, y2 in boxes
        ]
    draw = ImageDraw.Draw(image)
    width = max(4, round(image.width / 320))
    for i, box in enumerate(boxes, start=1):
        draw.rounded_rectangle(box, radius=12, outline="#E23B3B", width=width)
        cx = box[0] + 18
        cy = box[1] + 18
        draw.ellipse((cx - 15, cy - 15, cx + 15, cy + 15), fill="#E23B3B")
        draw.text((cx, cy), str(i), fill="white", font=font(19, True), anchor="mm")
    bar_h = max(48, round(image.height * 0.075))
    canvas = Image.new("RGB", (image.width, image.height + bar_h), "white")
    canvas.paste(image, (0, 0))
    bar = ImageDraw.Draw(canvas)
    bar.rectangle((0, image.height, image.width, image.height + bar_h), fill="#154F3B")
    bar.text(
        (20, image.height + bar_h // 2),
        label,
        fill="white",
        font=font(max(16, round(bar_h * 0.34)), True),
        anchor="lm",
    )
    canvas.save(target, optimize=True)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str) -> None:
    draw.line((start, end), fill=color, width=5)
    ex, ey = end
    sx, sy = start
    dx, dy = ex - sx, ey - sy
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    head = 16
    wing = 8
    p1 = (ex - ux * head + px * wing, ey - uy * head + py * wing)
    p2 = (ex - ux * head - px * wing, ey - uy * head - py * wing)
    draw.polygon([end, p1, p2], fill=color)


def rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    subtitle: str,
    fill: str,
    outline: str,
) -> None:
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = box
    draw.text((x1 + 18, y1 + 18), title, fill="#172033", font=font(22, True))
    draw.multiline_text(
        (x1 + 18, y1 + 54),
        subtitle,
        fill="#425069",
        font=font(15),
        spacing=5,
    )


def build_diagrams() -> None:
    flow = Image.new("RGB", (1500, 760), "white")
    d = ImageDraw.Draw(flow)
    rounded_box(d, (40, 80, 310, 230), "Kasir Apotik", "Pilih mode, cari item,\nbaca penanda keselamatan", "#EEF7F5", "#0D756D")
    rounded_box(d, (400, 80, 670, 230), "Keranjang", "Batch FEFO, qty, harga,\nidentitas pembeli", "#F1EDFB", "#6A4BB2")
    rounded_box(d, (760, 80, 1030, 230), "Pembayaran", "Validasi server, metode bayar,\nkode idempoten", "#EEF7F5", "#0D756D")
    rounded_box(d, (1120, 80, 1460, 230), "Hasil transaksi", "Nota, mutasi stok, status resep,\ncatatan audit", "#F1EDFB", "#6A4BB2")
    for x in [(310, 155, 400, 155), (670, 155, 760, 155), (1030, 155, 1120, 155)]:
        arrow(d, (x[0], x[1]), (x[2], x[3]), "#0D756D")
    rounded_box(d, (220, 390, 570, 560), "Baca lokal dulu", "Katalog, formula, dan antrean resep\ndapat tampil dari cache dengan penanda usia data", "#F8FAFC", "#9AA6B7")
    rounded_box(d, (645, 390, 995, 560), "Mutasi tetap aman", "Pembayaran dan produksi tidak mengaku\nberhasil sebelum server mengonfirmasi", "#FFF0F0", "#E23B3B")
    rounded_box(d, (1070, 390, 1420, 560), "Sinkronisasi", "Data server menyegarkan cache dan\nperubahan petugas lain diberi penanda", "#F8FAFC", "#9AA6B7")
    arrow(d, (535, 230), (395, 390), "#6A4BB2")
    arrow(d, (895, 230), (820, 390), "#6A4BB2")
    arrow(d, (1290, 230), (1245, 390), "#6A4BB2")
    d.text((40, 655), "Alur operasional POS Apotik dan batas integritas local-first", fill="#172033", font=font(28, True))
    flow.save(ASSETS / "diagram-alur-pos-apotik.png", optimize=True)

    erd = Image.new("RGB", (1500, 860), "white")
    d = ImageDraw.Draw(erd)
    entities = {
        "Item obat": (60, 80, 360, 250, "id, kode, nama\nbentuk, kekuatan, harga"),
        "Batch": (600, 80, 900, 250, "item_id, lot, kedaluwarsa\nstok, status"),
        "Formula": (1080, 80, 1410, 250, "id, jenis\nkomposisi, hasil"),
        "Resep": (60, 520, 360, 710, "id, pasien, dokter\ndiagnosis, status"),
        "Transaksi": (600, 500, 900, 730, "id, kode idempoten\nmetode, total, status"),
        "Rincian transaksi": (1060, 500, 1420, 730, "transaksi_id, item_id\nbatch_id, qty, harga"),
    }
    for title, (x1, y1, x2, y2, body) in entities.items():
        rounded_box(d, (x1, y1, x2, y2), title, body, "#F8FAFC", "#0D756D")
    arrow(d, (360, 165), (600, 165), "#0D756D")
    d.text((455, 130), "1 : N", fill="#425069", font=font(18, True))
    arrow(d, (900, 165), (1080, 165), "#6A4BB2")
    d.text((965, 130), "bahan/hasil", fill="#425069", font=font(17, True))
    arrow(d, (360, 610), (600, 610), "#0D756D")
    d.text((450, 575), "0..1 : N", fill="#425069", font=font(18, True))
    arrow(d, (900, 615), (1060, 615), "#0D756D")
    d.text((940, 580), "1 : N", fill="#425069", font=font(18, True))
    arrow(d, (1215, 250), (1215, 500), "#6A4BB2")
    d.text((1230, 350), "racikan / produksi", fill="#425069", font=font(17, True))
    arrow(d, (750, 250), (750, 500), "#0D756D")
    d.text((765, 350), "alokasi FEFO", fill="#425069", font=font(17, True))
    d.text((60, 790), "Model data konseptual untuk panduan pengguna; relasi fisik tetap mengikuti skema backend.", fill="#5D687C", font=font(19))
    erd.save(ASSETS / "diagram-data-konseptual-apotik.png", optimize=True)


def build_annotated_assets() -> None:
    annotate(
        SHOTS / "02-kasir-modern-laptop-1366.png",
        ASSETS / "manual-01-mode-dan-tebus.png",
        "1  Pilih mode transaksi    2  Tebus resep tetap terlihat",
        [(12, 25, 596, 80), (1060, 23, 1205, 82)],
        crop=(0, 0, 1220, 205),
    )
    annotate(
        SHOTS / "02-kasir-modern-laptop-1366.png",
        ASSETS / "manual-02-cari-dan-filter.png",
        "1  Cari nama, kode, atau barcode    2  Gunakan filter cepat",
        [(12, 80, 995, 133), (12, 136, 648, 178)],
        crop=(0, 66, 1010, 220),
    )
    annotate(
        SHOTS / "01-kasir-modern-desktop-1920.png",
        ASSETS / "manual-03-kartu-obat.png",
        "1  Identitas obat    2  Harga dan stok    3  Penanda keselamatan    4  Tambahkan",
        [(272, 142, 684, 318), (282, 160, 502, 244), (282, 247, 612, 308), (621, 160, 675, 218)],
        crop=(250, 125, 1120, 490),
    )
    source = Image.open(SHOTS / "01-kasir-modern-desktop-1920.png").convert("RGB")
    top = source.crop((1510, 0, 1920, 345))
    bottom = source.crop((1510, 845, 1920, 1080))
    separator = 44
    bar_h = 50
    cart = Image.new(
        "RGB", (410, top.height + separator + bottom.height + bar_h), "white"
    )
    cart.paste(top, (0, 0))
    cart.paste(bottom, (0, top.height + separator))
    draw = ImageDraw.Draw(cart)
    draw.rectangle((0, top.height, 410, top.height + separator), fill="#EEF2F7")
    draw.text(
        (205, top.height + separator // 2),
        "⋮  bagian daftar item  ⋮",
        fill="#5D687C",
        font=font(17, True),
        anchor="mm",
    )
    draw.rounded_rectangle(
        (5, 5, 405, top.height + separator + bottom.height - 5),
        radius=12,
        outline="#E23B3B",
        width=4,
    )
    draw.ellipse((13, 13, 45, 45), fill="#E23B3B")
    draw.text((29, 29), "1", fill="white", font=font(19, True), anchor="mm")
    draw.rectangle(
        (0, top.height + separator + bottom.height, 410, cart.height), fill="#154F3B"
    )
    draw.text(
        (16, cart.height - bar_h // 2),
        "Panel keranjang desktop",
        fill="white",
        font=font(18, True),
        anchor="lm",
    )
    cart.save(ASSETS / "manual-04-keranjang-desktop.png", optimize=True)
    annotate(
        SHOTS / "04-kasir-modern-mobile-390.png",
        ASSETS / "manual-05-mobile.png",
        "Ponsel: mode dan Tebus Resep berada di bagian atas katalog",
        [(10, 24, 379, 136), (12, 142, 380, 190)],
    )
    annotate(
        SHOTS / "05-kasir-modern-mobile-keranjang-390.png",
        ASSETS / "manual-06-mobile-cart.png",
        "Ringkasan keranjang melekat di bawah setelah obat dipilih",
        [(0, 777, 390, 843)],
    )


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: int = 5) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)


def set_run_font(run, name: str = "Aptos", size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr)
    run._r.append(fld_char2)


def setup_document(title: str, subject: str) -> Document:
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.subject = subject
    doc.core_properties.author = "Tim Implementasi Apotik"
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [
        ("Title", 30, GREEN),
        ("Heading 1", 18, GREEN),
        ("Heading 2", 13, TEAL),
        ("Heading 3", 10.5, PURPLE),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    footer = section.footer
    table = footer.add_table(rows=1, cols=3, width=Inches(6.9))
    table.columns[0].width = Inches(2.4)
    table.columns[1].width = Inches(2.1)
    table.columns[2].width = Inches(2.4)
    table.cell(0, 0).text = "Apotik UI/UX V2"
    table.cell(0, 1).text = "8 September 2026"
    p = table.cell(0, 2).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("Halaman ")
    page_field(p)
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=8)
            run.font.color.rgb = RGBColor.from_string(MUTED)
    return doc


def cover(doc: Document, title: str, subtitle: str, status: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(32)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("APOTIK")
    set_run_font(run, "Aptos Display", 15, True)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p = doc.add_paragraph()
    run = p.add_run(subtitle)
    set_run_font(run, size=14)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_paragraph()
    t = doc.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    t.columns[0].width = Cm(4.2)
    t.columns[1].width = Cm(12.6)
    for i, (k, v) in enumerate([
        ("Dokumen", title),
        ("Versi", "UI/UX V2 · 8 September 2026"),
        ("Lingkungan", "Build Windows Apotik dan pengujian integrasi lokal"),
        ("Status", status),
    ]):
        t.cell(i, 0).text = k
        t.cell(i, 1).text = v
        set_cell_shading(t.cell(i, 0), PALE)
        for cell in t.rows[i].cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9.5, bold=(cell == t.cell(i, 0)))
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("Dokumen ini mencatat hasil yang dapat direproduksi dari source code dan build lokal. UAT live pada server dilakukan setelah frontend versi ini dideploy.")
    set_run_font(run, size=10)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    p.paragraph_format.space_before = Pt(12)
    doc.add_page_break()


def add_summary_box(doc: Document, heading: str, text: str, fill: str = PALE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, TEAL, 8)
    p = cell.paragraphs[0]
    r = p.add_run(heading + "\n")
    set_run_font(r, size=11, bold=True)
    r.font.color.rgb = RGBColor.from_string(GREEN)
    r = p.add_run(text)
    set_run_font(r, size=9.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        set_cell_shading(cell, GREEN)
        set_cell_border(cell, WHITE, 5)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=8.2, bold=True)
            run.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cells[col].text = str(value)
            set_cell_shading(cells[col], "FFFFFF" if row_index % 2 else "F6F8FB")
            set_cell_border(cells[col])
            cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[col].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=7.8)
                    if value == "PASS":
                        run.font.color.rgb = RGBColor.from_string(TEAL)
                        run.bold = True
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, image: Path, caption: str, width_inches: float = 6.7) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image), width=Inches(width_inches))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(6)
    r = c.add_run(caption)
    set_run_font(r, size=8.2)
    r.italic = True
    r.font.color.rgb = RGBColor.from_string(MUTED)


def build_uat() -> Path:
    doc = setup_document(
        "Dokumen UST dan UAT Redesign Apotik",
        "Hasil pengujian UI UX V2 Kasir Apotik",
    )
    cover(
        doc,
        "Dokumen UST dan UAT\nRedesign Kasir Apotik",
        "Verifikasi tampilan modern, responsif, keselamatan obat, dan regresi fungsi",
        "LULUS untuk build lokal · UAT live menunggu deploy frontend terbaru",
    )
    doc.add_heading("Ringkasan hasil", level=1)
    add_summary_box(
        doc,
        "Kesimpulan",
        "Redesign UI/UX V2 telah diterapkan pada source code. Seluruh 1.073 tes regresi lulus, 76 tes terarah lulus, dan 5 skenario screenshot dengan 100 data lulus. Tampilan mempertahankan Tebus Resep pada layar 390 px, memakai katalog lazy, dan menampilkan penanda keselamatan tanpa bergantung pada warna saja.",
    )
    add_table(
        doc,
        ["Area", "Bukti", "Hasil"],
        [
            ["Regresi penuh", "flutter test test · 1.073 kasus", "PASS"],
            ["Tes redesign terarah", "POS, desain, gambar, shell · 76 kasus", "PASS"],
            ["Golden visual", "Desktop, mobile, kartu, status · 10 gambar", "PASS"],
            ["UAT 100 data", "1920, 1366, 768, 390, cart mobile · 5 skenario", "PASS"],
            ["Analisis statis", "0 error, 0 warning; 50 info non-blocking", "PASS"],
            ["Build target", "Windows Apotik release", "PASS"],
        ],
        [4.2, 8.4, 2.2],
    )
    doc.add_heading("Ruang lingkup dan dasar pengujian", level=2)
    doc.add_paragraph(
        "Pengujian membandingkan implementasi dengan paket Paket_Redesain_UI_UX_Apotik_Modern_2026. Fokusnya adalah layar Kasir Apotik, komponen kartu obat, gambar obat, breakpoint, navigasi mode, keranjang, aksesibilitas, dan perlindungan alur local-first. Pengujian ini tidak mengubah kontrak API atau aturan bisnis pembayaran."
    )
    add_bullets(doc, [
        "Data uji berisi 100 obat dengan variasi bentuk sediaan, harga, stok, dan penanda keselamatan.",
        "Lebar responsif diuji pada 560, 760, 960, 980, 1040, 1120, 1280, 1520, dan 1680 piksel melalui widget test.",
        "Bukti gambar diambil pada 1920×1080, 1366×768, 768×1024, dan 390×844 piksel.",
        "Mutasi sensitif tetap membutuhkan konfirmasi server; cache dipakai untuk data baca yang aman.",
    ])

    doc.add_page_break()
    doc.add_heading("Perubahan yang diterapkan", level=1)
    add_table(
        doc,
        ["ID", "Perubahan", "Hasil yang terlihat"],
        [
            ["UI-01", "Header fokus", "Dropdown duplikat dan FAB bantuan ganda dihilangkan; bantuan ringkas tetap ada."],
            ["UI-02", "Navigasi transaksi", "OTC, Resep Dokter, Racikan, Produksi Farmasi, dan Tebus Resep dapat dijangkau."],
            ["UI-03", "Filter cepat", "Semua stok, Tersedia, Stok menipis, bentuk sediaan, dan Perlu perhatian."],
            ["UI-04", "Kartu obat", "Gambar, identitas, harga, stok, kode, risiko, dan tombol tambah tersusun konsisten."],
            ["UI-05", "Urutan risiko", "High-alert, terkendali, golongan/Rx, LASA, cold-chain, lalu stok."],
            ["UI-06", "Gambar obat", "Bytes lokal, path lokal, fotoUrls pertama, gambarUrl, lalu fallback bentuk sediaan."],
            ["UI-07", "Katalog lazy", "Hanya baris yang terlihat dibangun; tinggi kartu mengikuti isi dan skala teks."],
            ["UI-08", "Keranjang responsif", "Panel tetap di desktop dan ringkasan bawah hanya muncul saat cart ponsel berisi."],
        ],
        [1.6, 4.8, 9.0],
    )
    add_figure(doc, ASSETS / "diagram-alur-pos-apotik.png", "Diagram 1  Alur POS dan batas integritas local-first", 6.65)

    doc.add_page_break()
    doc.add_heading("Matriks UST dan UAT", level=1)
    tests = [
        ["UST-01", "Memuat 100 obat", "Indikator menyatakan 100 dari 100", "PASS"],
        ["UST-02", "Mode transaksi", "Empat mode kasir aktif dan dapat dipilih", "PASS"],
        ["UST-03", "Tebus Resep", "Tampil pada lebar 390 px", "PASS"],
        ["UST-04", "Desktop 1920", "Katalog + panel keranjang tetap", "PASS"],
        ["UST-05", "Laptop 1366", "Katalog + panel keranjang tetap", "PASS"],
        ["UST-06", "Tablet 768", "Dua kolom, cart tidak memakan area katalog", "PASS"],
        ["UST-07", "Mobile 390", "Satu kolom, kontrol tetap dapat dijangkau", "PASS"],
        ["UST-08", "Cart mobile", "Ringkasan bawah muncul hanya setelah item dipilih", "PASS"],
        ["UST-09", "Filter tersedia", "Menyaring stok lebih dari nol", "PASS"],
        ["UST-10", "Filter stok menipis", "Menyaring item di bawah ambang", "PASS"],
        ["UST-11", "Gambar lokal", "Bytes/path lokal diprioritaskan", "PASS"],
        ["UST-12", "Gambar jaringan", "fotoUrls pertama mendahului gambarUrl", "PASS"],
        ["UST-13", "Fallback gambar", "Bentuk sediaan tetap memberi identitas visual", "PASS"],
        ["UST-14", "Badge risiko", "Ikon dan teks selalu tersedia", "PASS"],
        ["UST-15", "Stok habis", "Aksi tambah dinonaktifkan", "PASS"],
        ["UST-16", "Item terkunci", "Alasan ditampilkan dan tap ditolak", "PASS"],
        ["UST-17", "Pembayaran", "Gagal server mempertahankan cart", "PASS"],
        ["UST-18", "Idempotensi", "Kode sama dipakai saat retry gagal", "PASS"],
        ["UST-19", "Cache katalog", "Data cache diberi penanda usia", "PASS"],
        ["UST-20", "Regresi aplikasi", "1.073 tes tanpa kegagalan", "PASS"],
    ]
    add_table(doc, ["ID", "Skenario", "Kriteria penerimaan", "Hasil"], tests, [1.8, 4.4, 8.0, 1.7])

    doc.add_page_break()
    doc.add_heading("Bukti visual desktop", level=1)
    add_figure(doc, SHOTS / "01-kasir-modern-desktop-1920.png", "Gambar 1  Desktop 1920×1080 dengan 100 data dan panel keranjang tetap", 6.75)
    doc.add_paragraph(
        "Pada lebar desktop, mode transaksi berada di panel kiri, pencarian dan katalog berada di tengah, serta keranjang berada di kanan. Informasi barcode tidak mendominasi kartu; kode tetap terbaca dan barcode lengkap tersedia melalui tooltip/detail."
    )
    add_figure(doc, SHOTS / "02-kasir-modern-laptop-1366.png", "Gambar 2  Laptop 1366×768 dengan katalog tiga kolom", 6.75)

    doc.add_page_break()
    doc.add_heading("Bukti visual tablet dan ponsel", level=1)
    add_figure(doc, SHOTS / "03-kasir-modern-tablet-768.png", "Gambar 3  Tablet 768×1024 dengan dua kolom dan Tebus Resep terlihat", 5.45)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (image, caption) in enumerate([
        (SHOTS / "04-kasir-modern-mobile-390.png", "Katalog satu kolom"),
        (SHOTS / "05-kasir-modern-mobile-keranjang-390.png", "Cart melekat saat berisi"),
    ]):
        p = table.cell(0, i).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image), width=Inches(2.72))
        c = table.cell(0, i).add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in c.runs:
            set_run_font(run, size=8, bold=True)
    doc.add_paragraph(
        "Pada 390 piksel, Tebus Resep diletakkan pada baris tindakan tersendiri agar tidak hilang. Ringkasan keranjang tidak menutup katalog ketika kosong dan baru muncul setelah item dipilih."
    )

    doc.add_page_break()
    doc.add_heading("Matriks responsif", level=1)
    add_table(
        doc,
        ["Lebar isi", "Komposisi", "Katalog", "Keranjang"],
        [
            ["< 600 px", "Satu area", "1 kolom", "Ringkasan bawah saat berisi"],
            ["600–979 px", "Satu/dua area", "1–2 kolom", "Sheet atau ringkasan"],
            ["980–1519 px", "Dua area", "2–3 kolom", "Panel tetap 335–360 px"],
            ["≥ 1520 px", "Tiga area bila cukup", "3–4 kolom", "Panel tetap"],
        ],
        [3.0, 4.0, 4.0, 5.0],
    )
    doc.add_heading("Keselamatan dan aksesibilitas", level=2)
    add_bullets(doc, [
        "Setiap status memakai ikon dan teks; warna hanya memperkuat makna.",
        "Nama obat dan kekuatan berada sebelum harga agar identifikasi tidak bergantung pada foto.",
        "High-alert dan obat terkendali mendahului penanda komersial atau stok.",
        "Tombol tambah berukuran minimal 44 piksel dan memiliki label semantik lengkap.",
        "Kartu mengikuti tinggi konten sehingga badge banyak dan skala teks tidak dipotong.",
        "Foto adalah bantuan visual. Kasir tetap memverifikasi nama, kekuatan, bentuk, kode, dan batch.",
    ])
    add_figure(doc, ASSETS / "diagram-data-konseptual-apotik.png", "Diagram 2  Hubungan data konseptual POS Apotik", 6.65)

    doc.add_page_break()
    doc.add_heading("Kesiapan rilis dan UAT live", level=1)
    add_summary_box(
        doc,
        "Status rilis",
        "Source code, pengujian, screenshot 100 data, dan build Windows release siap. Frontend versi ini harus dideploy sebelum UAT live. Dokumen tidak menyatakan tampilan server sudah berubah sebelum deploy tersebut selesai.",
        PALE_PURPLE,
    )
    add_numbered(doc, [
        "Deploy frontend dari working tree/revisi yang memuat UI/UX V2. Backend tidak memerlukan WAR dari pekerjaan redesign ini.",
        "Buka Kasir Apotik pada 1920×1080 dan 1366×768; pastikan katalog dan panel keranjang tampil tanpa overflow.",
        "Ulangi pada layar kecil; pastikan tombol Tebus Resep terlihat sebelum menggulir dan cart bawah muncul setelah memilih obat.",
        "Lakukan smoke test OTC: pilih batch, bayar, periksa nota dan pengurangan stok.",
        "Lakukan smoke test Resep Dokter, Racikan, Produksi Farmasi, dan Tebus Resep tanpa mengubah aturan server.",
        "Catat nomor build dan waktu deploy pada tabel persetujuan."
    ])
    doc.add_heading("Kriteria rollback", level=2)
    add_bullets(doc, [
        "Mode transaksi atau Tebus Resep tidak dapat dijangkau pada resolusi yang didukung.",
        "Pembayaran atau produksi mengubah perilaku server dibanding build sebelumnya.",
        "Kartu obat menampilkan identitas, harga, stok, atau penanda risiko yang tidak sesuai respons API.",
        "Overflow menghalangi tombol utama atau keranjang pada 390, 768, 1366, atau 1920 piksel.",
    ])
    doc.add_heading("Persetujuan", level=2)
    add_table(
        doc,
        ["Peran", "Nama", "Tanggal", "Keputusan"],
        [
            ["Pelaksana UST/UAT", "", "", ""],
            ["Apoteker penanggung jawab", "", "", ""],
            ["Pemilik proses", "", "", ""],
            ["Penyetuju rilis", "", "", ""],
        ],
        [4.4, 4.2, 3.2, 4.2],
    )
    path = OUT / "Dokumen-UST-UAT-Redesign-Apotik-2026-09-08.docx"
    doc.save(path)
    return path


def build_manual() -> Path:
    doc = setup_document(
        "User Manual Kasir Apotik UI UX V2",
        "Panduan operasional tampilan terbaru Kasir Apotik",
    )
    cover(
        doc,
        "User Manual\nKasir Apotik UI/UX V2",
        "Panduan kasir dan apoteker untuk OTC, resep, racikan, produksi, dan tebus resep",
        "Panduan build UI/UX V2 · gunakan setelah frontend terbaru dideploy",
    )
    doc.add_heading("Tujuan dan pengguna", level=1)
    doc.add_paragraph(
        "Panduan ini membantu kasir dan apoteker menjalankan transaksi pada tampilan Kasir Apotik terbaru. Istilah pada layar ditulis sama dengan aplikasi agar langkah dapat diikuti tanpa menerjemahkan nama menu."
    )
    add_table(
        doc,
        ["Peran", "Tanggung jawab utama"],
        [
            ["Kasir", "Memilih toko, mencari item, mengelola keranjang, dan menerima pembayaran."],
            ["Apoteker", "Memverifikasi resep, racikan, produksi, batch, dan penanda keselamatan."],
            ["Supervisor", "Menangani penolakan yang memerlukan hak tambahan dan meninjau audit."],
        ],
        [4.2, 11.5],
    )
    doc.add_heading("Sebelum memulai", level=2)
    add_numbered(doc, [
        "Masuk dengan akun yang memiliki hak Kasir Apotik.",
        "Pilih toko atau apotek yang benar pada bagian atas aplikasi.",
        "Tekan Sinkronkan bila aplikasi menunjukkan data lokal belum diperbarui.",
        "Pastikan sesi kas terbuka dan metode pembayaran tersedia dari server.",
        "Periksa koneksi sebelum pembayaran, produksi, atau perubahan status batch."
    ])
    add_summary_box(
        doc,
        "Aturan penting",
        "Foto obat hanya membantu pencarian visual. Selalu cocokkan nama, kekuatan, bentuk sediaan, kode, batch, dan kedaluwarsa sebelum menyerahkan obat.",
        PALE_RED,
    )

    doc.add_page_break()
    doc.add_heading("Mengenali layar Kasir Apotik", level=1)
    add_figure(doc, ASSETS / "manual-01-mode-dan-tebus.png", "Gambar 1  Mode transaksi dan tombol Tebus Resep", 6.7)
    add_bullets(doc, [
        "OTC / Obat Bebas digunakan untuk penjualan obat jadi tanpa resep.",
        "Resep Dokter digunakan untuk memasukkan atau melayani resep dokter.",
        "Racikan digunakan untuk formula racikan yang sudah dapat dijual.",
        "Produksi Farmasi digunakan untuk menghasilkan barang jadi dari formula produksi.",
        "Tebus Resep membuka antrean resep yang menunggu penyerahan dan pembayaran.",
        "Detail membuka konteks transaksi pada lebar layar yang tidak menyediakan panel samping penuh."
    ])
    add_figure(doc, ASSETS / "manual-02-cari-dan-filter.png", "Gambar 2  Pencarian dan filter cepat", 6.7)
    doc.add_paragraph(
        "Ketik nama, kode, atau pindai barcode. Gunakan Tersedia untuk menyembunyikan stok kosong, Stok menipis untuk prioritas pengadaan, dan Perlu perhatian untuk melihat item dengan penanda risiko."
    )

    doc.add_page_break()
    doc.add_heading("Membaca kartu obat", level=1)
    add_figure(doc, ASSETS / "manual-03-kartu-obat.png", "Gambar 3  Informasi pada kartu obat", 6.75)
    add_table(
        doc,
        ["Elemen", "Cara membaca"],
        [
            ["Identitas", "Nama generik/merek, kekuatan, dan bentuk sediaan."],
            ["Harga", "Harga jual aktif untuk toko yang dipilih."],
            ["Stok", "Jumlah tersedia dan satuan; stok nol menonaktifkan tombol tambah."],
            ["Kode", "Kode item ringkas; barcode lengkap tersedia pada tooltip atau detail."],
            ["Penanda", "High-alert, terkendali, keras/Rx, LASA, cold-chain, dan kondisi stok."],
            ["Tombol tambah", "Menambahkan obat atau membuka pemilihan batch FEFO."],
        ],
        [4.0, 11.8],
    )
    doc.add_heading("Urutan pemeriksaan keselamatan", level=2)
    add_numbered(doc, [
        "Cocokkan nama obat dan kekuatan.",
        "Periksa bentuk sediaan dan satuan.",
        "Baca penanda High-alert, Terkendali, Keras/Rx, LASA, dan Cold-chain.",
        "Periksa stok, nomor batch, serta tanggal kedaluwarsa.",
        "Tambahkan item hanya setelah identitas dan risiko sesuai."
    ])

    doc.add_page_break()
    doc.add_heading("Penjualan OTC atau obat bebas", level=1)
    add_numbered(doc, [
        "Pilih OTC / Obat Bebas.",
        "Cari obat menggunakan nama, kode, atau barcode.",
        "Periksa identitas, harga, stok, dan penanda keselamatan pada kartu.",
        "Tekan tombol tambah. Bila pilihan batch muncul, gunakan batch FEFO yang layak dan masukkan jumlah.",
        "Tinjau item, batch, kuantitas, dan subtotal di keranjang.",
        "Isi identitas pembeli atau dokter bila diwajibkan oleh jenis obat.",
        "Tekan Bayar, pilih metode yang disediakan server, lalu konfirmasi uang diterima dan kembalian.",
        "Setelah server menyatakan berhasil, serahkan obat dan nota. Pastikan keranjang kembali kosong."
    ])
    add_figure(doc, ASSETS / "manual-04-keranjang-desktop.png", "Gambar 4  Panel keranjang desktop tetap terlihat selama transaksi", 2.45)
    add_summary_box(
        doc,
        "Jika pembayaran ditolak",
        "Jangan membuat transaksi baru untuk barang yang sama. Baca pesan server, pertahankan keranjang, perbaiki penyebabnya, lalu ulangi. Aplikasi memakai kode idempoten yang sama agar percobaan ulang tidak menggandakan transaksi.",
        PALE_RED,
    )

    doc.add_page_break()
    doc.add_heading("Resep Dokter", level=1)
    add_numbered(doc, [
        "Pilih Resep Dokter.",
        "Pilih pasien dan isi sumber resep, nama dokter, nomor izin praktik bila tersedia, tanggal resep, serta fasilitas asal.",
        "Catat diagnosis atau indikasi sesuai data resep. Jangan membuat diagnosis baru bila tidak tertulis.",
        "Tambahkan obat jadi dan aturan pakai. Pastikan dosis, frekuensi, rute, durasi, serta jumlah sesuai resep.",
        "Periksa alergi, duplikasi terapi, interaksi, kontraindikasi, dan penanda risiko yang tersedia.",
        "Simpan resep. Resep yang belum diserahkan masuk ke antrean Tebus Resep.",
        "Lanjutkan pembayaran setelah telaah dan dispensing selesai."
    ])
    add_summary_box(
        doc,
        "Data resep minimum",
        "Pasien, dokter, fasilitas asal, tanggal resep, diagnosis/indikasi, item, kekuatan, bentuk, aturan pakai, jumlah, catatan klinis, serta status telaah dan dispensing.",
    )

    doc.add_heading("Racikan", level=1)
    add_numbered(doc, [
        "Pilih Racikan lalu cari formula yang sudah aktif dan dapat dijual.",
        "Buka detail formula dan periksa komposisi, kekuatan, bentuk akhir, jumlah bungkus, serta aturan pakai.",
        "Masukkan jumlah yang diminta. Sistem menghitung kebutuhan bahan sesuai formula.",
        "Periksa ketersediaan batch bahan dan penanda keselamatan setiap komponen.",
        "Tambahkan racikan ke keranjang, lakukan verifikasi apoteker, lalu bayar.",
        "Cetak atau tulis etiket sesuai hasil yang sudah dikonfirmasi server."
    ])

    doc.add_page_break()
    doc.add_heading("Produksi Farmasi", level=1)
    add_numbered(doc, [
        "Pilih Produksi Farmasi lalu cari formula produksi aktif.",
        "Periksa barang jadi, target jumlah, komposisi, satuan, dan kebutuhan bahan.",
        "Pilih atau konfirmasi batch bahan FEFO yang layak.",
        "Isi hasil produksi, batch baru, tanggal produksi, kedaluwarsa, serta catatan QC bila diminta.",
        "Konfirmasi proses. Server mengurangi stok bahan dan menambah batch barang jadi secara atomik.",
        "Setelah produksi berhasil, pilih barang jadi untuk dijual melalui keranjang."
    ])
    add_summary_box(
        doc,
        "Batas integritas",
        "Produksi tidak boleh dianggap berhasil dari cache. Bila koneksi gagal, periksa status transaksi pada server sebelum mengulang agar stok bahan dan barang jadi tidak berubah dua kali.",
        PALE_RED,
    )
    doc.add_heading("Tebus Resep", level=1)
    add_numbered(doc, [
        "Tekan Tebus Resep. Tombol ini tetap terlihat pada layar kecil.",
        "Cari dan pilih resep berstatus menunggu.",
        "Cocokkan pasien, dokter, fasilitas asal, diagnosis, tanggal, serta daftar obat/racikan.",
        "Periksa hasil telaah, dispensing, batch, dan aturan pakai.",
        "Masukkan seluruh item obat jadi dan racikan dalam satu keranjang.",
        "Bayar satu kali. Server menyelesaikan transaksi, pengurangan stok, dan status resep secara atomik.",
        "Serahkan obat setelah status berhasil dan edukasi pasien selesai."
    ])

    doc.add_page_break()
    doc.add_heading("Menggunakan layar tablet dan ponsel", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (image, caption) in enumerate([
        (ASSETS / "manual-05-mobile.png", "Mode, Tebus Resep, dan pencarian"),
        (ASSETS / "manual-06-mobile-cart.png", "Ringkasan keranjang saat berisi"),
    ]):
        p = table.cell(0, i).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image), width=Inches(2.75))
        c = table.cell(0, i).add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in c.runs:
            set_run_font(run, size=8, bold=True)
    add_bullets(doc, [
        "Geser baris mode atau filter secara horizontal bila pilihan paling kanan belum terlihat.",
        "Tebus Resep berada pada baris tindakan tersendiri dan dapat dipilih tanpa membuka sidebar.",
        "Ketuk ringkasan Keranjang di bawah untuk membuka rincian transaksi.",
        "Panel bawah hanya tampil saat item sudah dipilih sehingga tidak menutup katalog kosong.",
        "Pada tablet, katalog menggunakan dua kolom. Pada ponsel, katalog menggunakan satu kolom."
    ])

    doc.add_page_break()
    doc.add_heading("Alur sistem dan data", level=1)
    add_figure(doc, ASSETS / "diagram-alur-pos-apotik.png", "Diagram 1  Alur operasional dari pencarian sampai transaksi", 6.65)
    add_figure(doc, ASSETS / "diagram-data-konseptual-apotik.png", "Diagram 2  Hubungan data konseptual", 6.65)
    doc.add_paragraph(
        "Diagram data bersifat konseptual untuk membantu pengguna memahami hubungan item, batch, formula, resep, transaksi, dan rincian transaksi. Nama tabel serta relasi fisik mengikuti skema backend yang dikelola tim teknis."
    )

    doc.add_page_break()
    doc.add_heading("Penanganan masalah", level=1)
    add_table(
        doc,
        ["Gejala", "Tindakan"],
        [
            ["Item tidak ditemukan", "Hapus filter, cari dengan kode, lalu sinkronkan katalog bila data lokal usang."],
            ["Tebus Resep tidak terlihat", "Pastikan memakai UI/UX V2; pada ponsel tombol berada di baris kedua di atas pencarian."],
            ["Batch tidak dapat dipilih", "Periksa status batch, stok, kedaluwarsa, dan koneksi server."],
            ["Tombol tambah padam", "Item stok kosong atau terkunci. Baca alasan pada kartu."],
            ["Pembayaran gagal", "Pertahankan cart, baca pesan server, perbaiki penyebab, lalu ulangi transaksi yang sama."],
            ["Produksi gagal", "Jangan mengulang sebelum memeriksa apakah server sudah mencatat produksi."],
            ["Data tampak lama", "Periksa penanda cache dan waktu sinkron terakhir, lalu tekan Sinkronkan."],
            ["Tampilan terpotong", "Catat resolusi dan skala tampilan; gunakan 100% bila memungkinkan lalu laporkan screenshot."],
        ],
        [5.0, 10.8],
    )
    doc.add_heading("Informasi untuk pelaporan", level=2)
    add_bullets(doc, [
        "Nama pengguna, toko, waktu, mode transaksi, dan nomor/kode transaksi.",
        "Resolusi layar dan skala tampilan Windows.",
        "Langkah yang dilakukan sebelum masalah muncul.",
        "Pesan yang terlihat dan screenshot tanpa data pribadi yang tidak perlu.",
        "Status koneksi serta waktu sinkron terakhir."
    ])
    doc.add_heading("Riwayat dokumen", level=2)
    add_table(
        doc,
        ["Versi", "Tanggal", "Perubahan"],
        [["UI/UX V2", "8 September 2026", "Redesign responsif, kartu obat, gambar, filter, cart mobile, dan UAT 100 data."]],
        [3.0, 4.0, 8.8],
    )
    path = OUT / "User-Manual-Kasir-Apotik-UIUX-V2-2026-09-08.docx"
    doc.save(path)
    return path


def main() -> None:
    required = [
        SHOTS / "01-kasir-modern-desktop-1920.png",
        SHOTS / "02-kasir-modern-laptop-1366.png",
        SHOTS / "03-kasir-modern-tablet-768.png",
        SHOTS / "04-kasir-modern-mobile-390.png",
        SHOTS / "05-kasir-modern-mobile-keranjang-390.png",
    ]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError("Screenshot UAT belum tersedia: " + ", ".join(missing))
    build_diagrams()
    build_annotated_assets()
    uat = build_uat()
    manual = build_manual()
    print(uat)
    print(manual)


if __name__ == "__main__":
    main()
