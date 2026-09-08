from __future__ import annotations

from pathlib import Path
from textwrap import wrap
from math import atan2, cos, sin, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "apotik-uiux-v2" / "uat-2026-09-08"
SHOTS = OUT / "screenshots-live-simplified"
ASSETS = OUT / "assets-live-simplified"
DOCX = OUT / "User-Manual-Apotik-Sederhana-Live-2026-09-08.docx"

GREEN = "154F3B"
MINT = "DDF4E8"
NAVY = "172033"
BLUE = "2563EB"
MUTED = "64748B"
LIGHT = "F3F6FA"
WHITE = "FFFFFF"


def font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def draw_centered_text(draw, box, text, used_font, fill):
    x1, y1, x2, y2 = box
    lines = []
    for paragraph in text.split("\n"):
        lines.extend(wrap(paragraph, width=max(12, int((x2 - x1) / (used_font.size * 0.62)))))
    heights = [draw.textbbox((0, 0), line, font=used_font)[3] for line in lines]
    total = sum(heights) + max(0, len(lines) - 1) * 8
    y = y1 + ((y2 - y1) - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=used_font)
        draw.text((x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2, y), line, font=used_font, fill=fill)
        y += height + 8


def rounded_box(draw, box, fill, outline, text, *, text_fill=NAVY, size=28):
    draw.rounded_rectangle(box, radius=22, fill=f"#{fill}", outline=f"#{outline}", width=3)
    draw_centered_text(draw, box, text, font(size, True), f"#{text_fill}")


def arrow(draw, start, end, color=GREEN):
    draw.line([start, end], fill=f"#{color}", width=8)
    angle = atan2(end[1] - start[1], end[0] - start[0])
    length = 22
    spread = pi / 7
    x, y = end
    draw.polygon(
        [
            (x, y),
            (x - length * cos(angle - spread), y - length * sin(angle - spread)),
            (x - length * cos(angle + spread), y - length * sin(angle + spread)),
        ],
        fill=f"#{color}",
    )


def build_flow_diagram(path: Path):
    image = Image.new("RGB", (1800, 1000), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    draw.text((70, 42), "Alur End-to-End Apotik Sederhana", font=font(48, True), fill=f"#{NAVY}")
    draw.text((70, 105), "Satu sumber produk, satu jalur penerimaan, dan transaksi yang dapat diaudit", font=font(25), fill=f"#{MUTED}")

    top = [
        ((70, 205, 330, 340), "Setup Produk\n& Lokasi"),
        ((400, 205, 660, 340), "PR"),
        ((730, 205, 990, 340), "PO"),
        ((1060, 205, 1320, 340), "BAST\nBatch & ED"),
        ((1390, 205, 1650, 340), "Tagihan &\nBayar Vendor"),
    ]
    for idx, (box, label) in enumerate(top):
        rounded_box(draw, box, MINT if idx in (0, 3) else WHITE, GREEN, label)
        if idx < len(top) - 1:
            arrow(draw, (box[2] + 10, 272), (top[idx + 1][0][0] - 10, 272))

    arrow(draw, (1190, 350), (1190, 470))
    rounded_box(draw, (910, 485, 1470, 625), "E8F0FE", BLUE, "Persediaan Layak Jual\nFEFO • Batch • Lokasi", text_fill=NAVY)
    branches = [
        ((70, 700, 365, 840), "OTC /\nObat Bebas"),
        ((430, 700, 725, 840), "Resep Dokter\n& Tebus Resep"),
        ((790, 700, 1085, 840), "Racikan\nPasien"),
        ((1150, 700, 1445, 840), "Produksi\nFarmasi"),
    ]
    stock_center_x = (910 + 1470) // 2
    stock_bottom_y = 625
    branch_y = 680
    draw.line([(stock_center_x, stock_bottom_y), (stock_center_x, branch_y)], fill=f"#{GREEN}", width=8)
    draw.line([(217, branch_y), (1297, branch_y)], fill=f"#{GREEN}", width=8)
    for box, label in branches:
        rounded_box(draw, box, WHITE, GREEN, label, size=26)
        center_x = (box[0] + box[2]) // 2
        arrow(draw, (center_x, branch_y), (center_x, box[1] - 8))
    rounded_box(draw, (1510, 700, 1730, 840), MINT, GREEN, "Laporan\nApotik", size=25)
    arrow(draw, (1455, 770), (1500, 770))
    draw.text((70, 910), "Kontrol: akses peran • idempotensi • audit • rekonsiliasi • persetujuan", font=font(27, True), fill=f"#{GREEN}")
    image.save(path)


def build_erd(path: Path):
    image = Image.new("RGB", (1800, 1000), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    draw.text((70, 42), "ERD Konseptual Apotik", font=font(48, True), fill=f"#{NAVY}")
    draw.text((70, 105), "Hubungan data utama yang digunakan dalam alur pengguna", font=font(25), fill=f"#{MUTED}")

    entities = {
        "Produk Obat": (90, 205, 370, 335),
        "Lokasi Simpan": (90, 455, 370, 585),
        "Batch / ED": (470, 205, 750, 335),
        "Formula Produksi": (470, 655, 750, 785),
        "Resep Dokter": (850, 205, 1130, 335),
        "Transaksi Jual": (850, 455, 1130, 585),
        "Formula Racikan": (850, 655, 1130, 785),
        "PR / PO": (1360, 205, 1640, 335),
        "BAST": (1360, 455, 1640, 585),
        "Tagihan / Bayar": (1360, 655, 1640, 785),
    }
    for label, box in entities.items():
        rounded_box(draw, box, WHITE if label not in ("Batch / ED", "Transaksi Jual", "BAST") else MINT, GREEN, label, size=24)

    def relation(points, label, label_at):
        draw.line(points, fill="#94A3B8", width=5, joint="curve")
        lx, ly = label_at
        draw.rounded_rectangle((lx - 72, ly - 20, lx + 72, ly + 20), 10, fill="#F7FAFC")
        draw_centered_text(draw, (lx - 72, ly - 20, lx + 72, ly + 20), label, font(17, True), f"#{MUTED}")

    relation([(370, 270), (470, 270)], "1:n", (420, 245))
    relation([(370, 520), (420, 520), (420, 310), (470, 310)], "1:n", (420, 420))
    relation([(610, 655), (610, 335)], "menghasilkan", (610, 500))
    relation([(750, 270), (800, 270), (800, 520), (850, 520)], "n:n", (800, 405))
    relation([(990, 335), (990, 455)], "0..1:1", (900, 395))
    relation([(990, 655), (990, 585)], "0..n", (1070, 620))
    relation([(1500, 335), (1500, 455)], "1:n", (1580, 395))
    relation([(1500, 585), (1500, 655)], "1:n", (1580, 620))
    relation([(1360, 520), (1240, 520), (1240, 385), (710, 385), (710, 335)], "menambah batch", (1160, 360))
    draw.text((70, 905), "Catatan: ERD ini konseptual untuk pengguna; nama tabel fisik mengikuti model backend AIS/SIRS.", font=font(24), fill=f"#{MUTED}")
    image.save(path)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Halaman ")
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_sep, fld_end])


def configure_document(doc: Document):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    for style_name, size, color in [
        ("Title", 30, NAVY),
        ("Heading 1", 22, NAVY),
        ("Heading 2", 16, GREEN),
        ("Heading 3", 12, GREEN),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "APOTIK • USER MANUAL BERBASIS UAT LIVE • 8 SEPTEMBER 2026"
    header.style = styles["Caption"]
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    add_page_field(section.footer.paragraphs[0])


def add_title(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    run = p.add_run("USER MANUAL\nAPOTIK SEDERHANA")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(34)
    run.font.color.rgb = RGBColor.from_string(GREEN)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Setup Produk • Persediaan • Pengadaan • Kasir • Resep • Racikan • Produksi • Laporan")
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = False
    box.columns[0].width = Inches(7.6)
    cell = box.cell(0, 0)
    shade(cell, MINT)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(
        "Berdasarkan screenshot aktual aplikasi desktop dan data live\n"
        "Server demo.ecampus.id/ecampus • Frontend daba55d • Backend r88240"
    )
    cr.bold = True
    cr.font.name = "Arial"
    cr.font.size = Pt(13)
    cr.font.color.rgb = RGBColor.from_string(GREEN)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(28)
    p3.add_run("Versi dokumen: 1.0 • Status UAT pascadeploy: PASS").bold = True
    doc.add_page_break()


def add_intro(doc: Document):
    doc.add_heading("Cara menggunakan manual ini", level=1)
    p = doc.add_paragraph(
        "Manual ini mengikuti urutan kerja harian apotek. Setiap prosedur menggunakan screenshot dari aplikasi yang telah dideploy. "
        "Nomor dan nama pada gambar adalah data sample/UAT; jangan dipakai sebagai keputusan klinis untuk pasien nyata."
    )
    p.paragraph_format.space_after = Pt(8)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Peran", "Tugas utama", "Menu yang digunakan"]
    for idx, label in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell, GREEN)
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
    set_repeat_table_header(table.rows[0])
    roles = [
        ("Admin / master data", "Menyiapkan produk dan referensi", "Setup Produk Obat; Batch & Kedaluwarsa"),
        ("Petugas pengadaan", "Melakukan kulakan yang terlacak", "Pengadaan Obat: PR, PO, BAST, Tagihan, Pembayaran"),
        ("Kasir / apoteker", "Menjual dan menyerahkan obat", "Kasir Apotik; Tebus Resep Dokter; Racikan"),
        ("Produksi / QC", "Membuat stok barang jadi", "Produksi Farmasi"),
        ("Supervisor", "Memantau stok dan penjualan", "Dashboard; Stok Opname; Retur; Laporan Apotik"),
    ]
    for role in roles:
        row = table.add_row()
        for idx, value in enumerate(role):
            row.cells[idx].text = value
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Prasyarat singkat")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(GREEN)
    for item in [
        "Masuk dengan akun yang memiliki hak Apotik dan pilih toko aktif.",
        "Tekan Sinkronkan bila ada indikator perubahan data lokal.",
        "Pastikan produk, satuan, jenis item medis, metode pembayaran, dan supplier sudah tersedia.",
        "Untuk transaksi resep, verifikasi identitas pasien dan informasi dokter sebelum menyimpan.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_page_break()


def add_diagram_page(doc: Document, title: str, description: str, image_path: Path):
    doc.add_heading(title, level=1)
    doc.add_paragraph(description)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(9.9))
    doc.add_page_break()


def add_screen_page(doc: Document, number: str, title: str, purpose: str, steps: list[str], screenshot: str, note: str | None = None):
    doc.add_heading(f"{number}. {title}", level=1)
    doc.add_paragraph(purpose)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Lebar ini menjaga screenshot dan langkah kerja tetap berada di halaman
    # yang sama pada A4 landscape, tanpa mengorbankan keterbacaan UI.
    p.add_run().add_picture(str(SHOTS / screenshot), width=Inches(7.75))
    caption = doc.add_paragraph(f"Gambar {number} — {title}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(0.8)
    table.columns[1].width = Inches(9.4)
    for idx, step in enumerate(steps, 1):
        row = table.rows[0] if idx == 1 else table.add_row()
        row.cells[0].text = str(idx)
        row.cells[1].text = step
        shade(row.cells[0], MINT)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if note:
        p = doc.add_paragraph()
        r = p.add_run(f"Catatan: {note}")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(GREEN)
    doc.add_page_break()


def add_final_checklists(doc: Document):
    doc.add_heading("Checklist operasional harian", level=1)
    groups = [
        ("Sebelum buka", ["Sinkronisasi selesai", "Toko dan kasir benar", "Printer/barcode siap", "Batch kedaluwarsa ditinjau"]),
        ("Saat pelayanan", ["Produk dan batch sesuai", "FEFO dipatuhi", "Resep dan identitas diverifikasi", "Total serta metode bayar dibacakan kembali"]),
        ("Saat tutup", ["Transaksi tertunda diperiksa", "Retur/selisih stok dicatat", "Rekonsiliasi kas dilakukan", "Laporan penjualan ditinjau"]),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (title, _) in enumerate(groups):
        cell = table.rows[0].cells[idx]
        shade(cell, GREEN)
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
    row = table.add_row()
    for idx, (_, items) in enumerate(groups):
        row.cells[idx].text = "\n".join(f"☐ {item}" for item in items)
    doc.add_paragraph()
    doc.add_heading("Penanganan kendala", level=2)
    issues = [
        ("Produk tidak tampil di kasir", "Pastikan produk dibuat melalui Setup Produk Obat, aktif, memiliki harga jual, dan stok/batch layak."),
        ("Tombol Bayar tidak aktif", "Periksa keranjang, alokasi batch, data wajib obat terkendali, serta koneksi server."),
        ("Resep tidak dapat diserahkan", "Lengkapi telaah apoteker, pemeriksaan kedua/konseling sesuai kewajiban, dan pastikan stok tersedia."),
        ("Barang pembelian belum menambah stok", "Pastikan penerimaan dilakukan melalui BAST dan dokumen disetujui; jangan memakai jalur PBF lama."),
        ("Data belum berubah", "Tekan Sinkronkan atau muat ulang. Jangan menekan konfirmasi berulang; catat pesan dan kode transaksi."),
    ]
    for issue, handling in issues:
        p = doc.add_paragraph()
        p.add_run(issue + ": ").bold = True
        p.add_run(handling)
    doc.add_page_break()


def add_uat_summary(doc: Document):
    doc.add_heading("Ringkasan UAT pascadeploy", level=1)
    doc.add_paragraph(
        "Pada 8 September 2026, harness Windows memuat aplikasi varian Apotik terhadap server live, menghitung data unik lintas halaman, "
        "membuka layar dan dialog yang didokumentasikan, lalu selesai dengan hasil All tests passed."
    )
    rows = [
        ("Produk obat", "100 / 11.000"),
        ("Racikan", "100 / 500 sampel"),
        ("Produksi", "100 / 500 sampel"),
        ("Resep menunggu", "100 / 3.897"),
        ("Batch", "100"),
        ("Laporan penjualan", "200"),
        ("PR / PO / BAST / Tagihan / Pembayaran", "105 per tahap"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, label in enumerate(["Cakupan", "Jumlah terverifikasi", "Status"]):
        cell = table.rows[0].cells[idx]
        shade(cell, GREEN)
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
    for area, count in rows:
        row = table.add_row()
        row.cells[0].text = area
        row.cells[1].text = count
        row.cells[2].text = "PASS"
        shade(row.cells[2], MINT)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Batas pengujian ini")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(GREEN)
    doc.add_paragraph(
        "Run pascadeploy ini tidak menekan konfirmasi pembayaran akhir dan tidak melakukan mutasi stok. Tujuannya adalah regresi live read-only dan bukti visual. "
        "UAT transaksi mutatif 100 data per alur telah dilakukan pada rangkaian UAT sebelumnya dan tercatat dalam handover proyek.",
        style=None,
    )
    doc.add_paragraph()
    doc.add_paragraph("Dokumen ini selesai dan siap digunakan untuk pelatihan/UAT pengguna.").runs[0].bold = True


def build():
    ASSETS.mkdir(parents=True, exist_ok=True)
    flow = ASSETS / "flow-end-to-end-apotik-sederhana.png"
    erd = ASSETS / "erd-konseptual-apotik-sederhana.png"
    build_flow_diagram(flow)
    build_erd(erd)

    required = [
        "00-dashboard-apotik.png",
        "01-setup-produk-obat-daftar.png",
        "02-setup-produk-obat-form.png",
        "03-batch-kedaluwarsa.png",
        "04-stok-opname.png",
        "05-retur-obat.png",
        "06-pengadaan-lima-tahap.png",
        "07b-daftar-pengadaan-pr.png",
        "08b-daftar-pengadaan-po.png",
        "09b-daftar-pengadaan-bast.png",
        "10b-daftar-tagihan-vendor.png",
        "11b-daftar-pembayaran-vendor.png",
        "12-kasir-otc-dan-resep.png",
        "12b-dialog-pembayaran-otc.png",
        "13-form-resep-baru-lengkap.png",
        "14-racikan-pasien.png",
        "15-produksi-farmasi.png",
        "16b-detail-tebus-resep-dokter.png",
        "17-monitoring-penjualan.png",
    ]
    missing = [name for name in required if not (SHOTS / name).exists()]
    if missing:
        raise FileNotFoundError(f"Screenshot belum tersedia: {missing}")

    doc = Document()
    configure_document(doc)
    add_title(doc)
    add_intro(doc)
    add_diagram_page(doc, "Alur proses end-to-end", "Ikuti alur utama dari kiri ke kanan. BAST adalah satu-satunya jalur penerimaan pembelian yang menambah persediaan.", flow)
    add_diagram_page(doc, "ERD konseptual", "Diagram ini membantu pengguna memahami keterkaitan produk, lokasi, batch, resep, transaksi, dan dokumen pengadaan.", erd)

    screens = [
        ("1", "Dashboard Apotik", "Gunakan dashboard sebagai daftar prioritas kerja saat mulai shift.", ["Pastikan toko aktif di bagian atas sudah benar.", "Tinjau resep menunggu, batch mendekati kedaluwarsa, dan stok habis.", "Buka item pada bagian Perlu tindakan sesuai prioritas."], "00-dashboard-apotik.png", None),
        ("2", "Daftar Setup Produk Obat", "Produk obat dibuat dari menu khusus agar langsung dikenali katalog Kasir Apotik.", ["Buka Setup Produk Obat.", "Cari dahulu kode/barcode/nama untuk mencegah duplikasi.", "Tekan Tambah Produk Obat bila belum tersedia."], "01-setup-produk-obat-daftar.png", "Jangan membuat obat dari master produk generik bila ingin langsung digunakan dalam alur farmasi."),
        ("3", "Form Tambah Produk Obat", "Lengkapi identitas, harga, atribut farmasi, dan lokasi fisik dalam satu form.", ["Isi kode, barcode, nama, satuan, jenis item medis, dan zat aktif.", "Isi harga beli, harga jual, stok minimum, golongan, kekuatan, serta bentuk sediaan.", "Tandai LASA, high-alert, cold-chain, dan retur sesuai kebijakan.", "Isi gudang/ruang sampai bin/posisi serta rentang suhu, lalu Simpan Produk."], "02-setup-produk-obat-form.png", "Stok awal tidak diinput pada form ini; stok masuk melalui BAST atau Stok Opname."),
        ("4", "Batch & Kedaluwarsa", "Pantau batch kedaluwarsa dan segera kedaluwarsa untuk menjalankan FEFO.", ["Pilih rentang hari kedaluwarsa.", "Pisahkan batch Kedaluwarsa dari batch Segera.", "Tahan/retur/musnahkan sesuai SOP; jangan menjual batch tidak layak."], "03-batch-kedaluwarsa.png", None),
        ("5", "Stok Opname Apotik", "Gunakan opname untuk mencatat hasil hitung fisik dan selisih persediaan.", ["Cari atau pindai obat.", "Pilih batch/lokasi yang dihitung.", "Isi jumlah fisik dan alasan selisih, lalu simpan sesuai kewenangan."], "04-stok-opname.png", None),
        ("6", "Retur Obat", "Catat barang yang keluar karena retur agar stok dan audit tetap konsisten.", ["Pilih obat dan batch yang diretur.", "Isi jumlah, alasan, dan pihak tujuan.", "Periksa ringkasan sebelum menyimpan."], "05-retur-obat.png", None),
        ("7", "Pusat Pengadaan Obat", "Satu halaman mengarahkan kulakan melalui lima tahap resmi.", ["Mulai dari PR untuk kebutuhan persediaan.", "Lanjutkan ke PO dan pilih supplier/PBF.", "Terima barang melalui BAST, lalu cocokkan tagihan dan pembayaran."], "06-pengadaan-lima-tahap.png", "Menu Penerimaan/PBF lama dihilangkan; penerimaan pembelian harus melalui BAST."),
        ("8", "Permintaan Pembelian (PR)", "PR merekam kebutuhan sebelum barang dipesan.", ["Tekan Buat PR.", "Isi kebutuhan, jumlah, keterangan, dan anggaran bila diwajibkan.", "Ajukan persetujuan dan pastikan status Disetujui sebelum membuat PO."], "07b-daftar-pengadaan-pr.png", None),
        ("9", "Pemesanan Pembelian (PO)", "PO mengikat permintaan dengan supplier, harga, dan termin.", ["Tekan Dari PR untuk menjaga keterlacakan.", "Pilih supplier/PBF, harga, jumlah, jadwal, serta termin.", "Setujui PO; pantau sisa kewajiban dan status Lunas/Disetujui."], "08b-daftar-pengadaan-po.png", None),
        ("10", "Penerimaan Barang (BAST)", "BAST adalah titik masuk stok dari pembelian.", ["Tekan Dari PO dan pilih pesanan yang diterima.", "Isi jumlah diterima, batch, tanggal kedaluwarsa, lokasi, dan hasil pemeriksaan.", "Setujui BAST; pastikan sumber PO dan nilai sesuai dokumen fisik."], "09b-daftar-pengadaan-bast.png", "BAST dapat mewakili penerimaan per termin; jangan menerima ulang dokumen yang sudah selesai."),
        ("11", "Terima Tagihan Vendor", "Tagihan barang dicocokkan dengan PO dan BAST yang sudah disetujui.", ["Buka tab Tagihan.", "Periksa BAST, PO, supplier, termin, nilai faktur, dan lampiran invoice.", "Terima tagihan setelah jumlah dan dokumen pendukung cocok."], "10b-daftar-tagihan-vendor.png", "Tagihan tanpa BAST hanya untuk tagihan rutin nonbarang sesuai konfigurasi."),
        ("12", "Pembayaran Vendor", "Bayar hanya tagihan yang sudah diterima dan disetujui.", ["Buka tab Pembayaran lalu tekan Bayar Vendor.", "Pilih tagihan, metode/cara transfer, tanggal, dan nilai.", "Setujui dan cocokkan status pelunasan dengan rekening koran."], "11b-daftar-pembayaran-vendor.png", None),
        ("13", "Kasir OTC dan Resep", "Kasir utama hanya menampilkan penjualan OTC/Obat Bebas dan Resep Dokter; racikan/produksi dipisah.", ["Pilih OTC/Obat Bebas atau Resep Dokter.", "Cari nama/kode atau pindai barcode.", "Tambah item, periksa batch FEFO dan jumlah pada keranjang.", "Isi identitas pembeli/dokter bila diwajibkan."], "12-kasir-otc-dan-resep.png", None),
        ("14", "Dialog Pembayaran OTC", "Dialog menampilkan total dan metode bayar sebelum transaksi dikirim.", ["Tekan Bayar setelah keranjang valid.", "Pilih Kredit Pelanggan, Online, Tunai, atau Bayar terpisah.", "Isi nomor referensi atau uang diterima sesuai metode.", "Periksa total/kembalian lalu tekan Bayar satu kali."], "12b-dialog-pembayaran-otc.png", "Jika hasil server belum pasti, jangan mengulangi pembayaran dengan transaksi baru; gunakan mekanisme retry/idempotensi."),
        ("15", "Membuat Resep Baru di Kasir", "Resep baru menyimpan asal, dokter, pasien, dan informasi klinis sebelum ditebus.", ["Tambahkan item obat dalam mode Resep Dokter lalu tekan Buat Resep Baru.", "Isi identitas pasien dan nomor rekam medis.", "Isi dokter/SIP, rumah sakit atau klinik, poli, dan tanggal resep.", "Isi ICD/diagnosis, penyakit/indikasi, serta catatan khusus; simpan dan kaitkan resep."], "13-form-resep-baru-lengkap.png", "Data klinis harus diverifikasi tenaga medis; data sample hanya untuk UAT."),
        ("16", "Racikan Pasien", "Racikan adalah pelayanan pasien berbasis formula resep, bukan proses menambah stok barang jadi.", ["Buka menu Racikan.", "Pilih formula yang sesuai resep pasien.", "Tentukan jumlah, telaah komposisi, siapkan, lalu jual/serahkan.", "Gunakan Tebus Resep bila racikan terkait antrean resep."], "14-racikan-pasien.png", None),
        ("17", "Produksi Farmasi", "Produksi mengonsumsi bahan untuk menghasilkan batch stok barang jadi melalui QC.", ["Buka menu Produksi Farmasi.", "Pilih formula/BOM dan tentukan jumlah produksi.", "Periksa bahan, konsumsi stok, hasil QC, nomor batch, kedaluwarsa, dan lokasi.", "Proses produksi setelah seluruh data siap."], "15-produksi-farmasi.png", "Produksi tidak sama dengan Racikan; hasilnya menjadi persediaan batch barang jadi."),
        ("18", "Tebus Resep Dokter", "Apoteker menelaah resep lengkap sebelum penyerahan.", ["Pilih resep berstatus Menunggu.", "Verifikasi pasien/RM, dokter/SIP, asal layanan, diagnosis, keluhan, anamnesis, alamat, dan kontak.", "Periksa alergi, LASA, interaksi/duplikasi/dosis, serta daftar periksa.", "Catat pemeriksaan kedua dan konseling sebelum obat diserahkan."], "16b-detail-tebus-resep-dokter.png", None),
        ("19", "Monitoring Penjualan", "Laporan Apotik merangkum nilai, kuantitas, golongan obat, serta rincian produk.", ["Pilih rentang tanggal dan tekan Muat.", "Periksa total penjualan dan total qty.", "Bandingkan Bebas, Psikotropika, dan Narkotika.", "Gunakan tab Obat Terkendali, Kedaluwarsa, dan Rekonsiliasi Kas untuk kontrol lanjutan."], "17-monitoring-penjualan.png", None),
    ]
    for args in screens:
        add_screen_page(doc, *args)
    add_final_checklists(doc)
    add_uat_summary(doc)
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
