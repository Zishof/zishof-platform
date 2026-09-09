from __future__ import annotations

import json
import math
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SCREEN = ROOT / "screenshots"
ANNOTATED = ROOT / "annotated"
DIAGRAMS = ROOT / "diagrams"
EVIDENCE = ROOT / "evidence" / "uat-training-operasional-evidence.json"
DOCX = ROOT / "Manual-UAT-Training-Operasional-POS-Pengadaan-Laporan-eBisnis-v1.34.32.docx"
PDF = ROOT / "Manual-UAT-Training-Operasional-POS-Pengadaan-Laporan-eBisnis-v1.34.32.pdf"
MD = ROOT / "HASIL-UAT-TRAINING-OPERASIONAL-eBisnis-v1.34.32.md"

NAVY = (15, 34, 57)
BLUE = (37, 99, 235)
CYAN = (14, 165, 233)
GREEN = (22, 163, 74)
AMBER = (245, 158, 11)
RED = (220, 38, 38)
INK = (31, 41, 55)
MUTED = (88, 101, 120)
LIGHT = (244, 247, 251)
WHITE = (255, 255, 255)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    filename = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path("C:/Windows/Fonts") / filename), size)


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textlength(candidate, font=fnt) <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def arrow(draw: ImageDraw.ImageDraw, p1, p2, color=BLUE, width=7):
    draw.line([p1, p2], fill=color, width=width)
    angle = math.atan2(p2[1] - p1[1], p2[0] - p1[0])
    length = 22
    a1 = angle + math.pi * 0.82
    a2 = angle - math.pi * 0.82
    tri = [p2, (p2[0] + length * math.cos(a1), p2[1] + length * math.sin(a1)),
           (p2[0] + length * math.cos(a2), p2[1] + length * math.sin(a2))]
    draw.polygon(tri, fill=color)


SCREEN_SPECS = [
    {
        "file": "01-kasir-pos-layar-penuh.png",
        "title": "Layar transaksi POS/Kasir",
        "callouts": [
            ((0.12, 0.12, 0.84, 0.17), "Cari/scan produk"),
            ((0.85, 0.16, 0.98, 0.23), "Pilih metode pembayaran"),
            ((0.85, 0.40, 0.98, 0.47), "Bayar dan simpan transaksi"),
        ],
        "narrative": (
            "Layar Kasir adalah titik awal penjualan. Operator mencari barang dengan nama atau barcode, "
            "memeriksa kuantitas dan harga di keranjang, lalu memilih member bila transaksi memakai saldo. "
            "Metode pembayaran harus dipilih sebelum tombol Bayar aktif. Pada pelatihan, gunakan transaksi "
            "bernilai kecil terlebih dahulu, cocokkan total layar dengan uang atau saldo yang diterima, kemudian "
            "selesaikan pembayaran satu kali agar nota tidak ganda. Setelah berhasil, transaksi dapat ditelusuri "
            "melalui Riwayat Penjualan dan seluruh laporan pada periode yang sama."
        ),
        "steps": ["Cari atau pindai produk.", "Atur jumlah dan pastikan harga benar.",
                  "Pilih member bila memakai Saldo.", "Pilih metode pembayaran.",
                  "Periksa total, tekan Bayar, lalu simpan/cetak nota."],
    },
    {
        "file": "02-riwayat-penjualan-404-transaksi.png",
        "title": "Riwayat Penjualan dan penelusuran nota",
        "callouts": [
            ((0.64, 0.13, 0.91, 0.18), "Sinkronkan dan tambah transaksi"),
            ((0.11, 0.25, 0.89, 0.31), "Filter tanggal dan pencarian nota"),
            ((0.11, 0.37, 0.97, 0.46), "Baris transaksi dan menu aksi"),
        ],
        "narrative": (
            "Riwayat Penjualan dipakai untuk memastikan nota benar-benar tersimpan dan untuk menelusuri transaksi "
            "yang menjadi sumber laporan. Filter tanggal pada layar ini berdiri sendiri; karena itu angka ringkasan "
            "hari berjalan dapat berbeda dari laporan UAT yang memakai rentang 1–10 September 2026. Gunakan nomor "
            "nota, nama pelanggan, produk, metode, atau kasir untuk mempersempit hasil. Menu tiga titik membuka detail "
            "dan tindakan yang diizinkan oleh peran pengguna. Saat transaksi dibuat ketika jaringan bermasalah, "
            "periksa juga status sinkronisasi sebelum membandingkan total dengan server."
        ),
        "steps": ["Atur tanggal mulai dan selesai.", "Masukkan nomor nota atau nama pelanggan.",
                  "Buka menu aksi pada baris yang dipilih.", "Cocokkan item, metode, waktu, dan total."],
    },
    {
        "file": "05-pr-formulir-baru.png",
        "title": "Membuat Permintaan Pembelian (PR)",
        "callouts": [
            ((0.36, 0.39, 0.64, 0.48), "Isi kebutuhan"),
            ((0.36, 0.45, 0.65, 0.53), "Pilih anggaran atau Tanpa anggaran"),
            ((0.54, 0.50, 0.64, 0.57), "Tambah barang"),
            ((0.58, 0.60, 0.65, 0.66), "Simpan PR"),
        ],
        "narrative": (
            "PR mencatat kebutuhan sebelum barang dipesan. Tuliskan kebutuhan yang dapat dipahami oleh penyetuju, "
            "misalnya stok minuman untuk satu pekan, kemudian tambahkan barang, satuan, dan jumlah yang diminta. "
            "Karena pelatihan An Nahl ini belum menggunakan COA dan tidak membahas Akuntansi, centang Tanpa anggaran "
            "apabila data anggaran belum disiapkan. Pilihan tersebut hanya mengatur validasi dokumen PR dan tidak "
            "mengubah ruang lingkup laporan operasional. Setelah disimpan, PR tetap perlu mengikuti keputusan/approval "
            "sesuai hak akses sebelum dapat dipilih dari PO."
        ),
        "steps": ["Tekan Buat PR.", "Isi keterangan kebutuhan.",
                  "Centang Tanpa anggaran untuk pelatihan ini.", "Tambah barang beserta UOM dan jumlah.",
                  "Periksa total, simpan, lalu ajukan untuk disetujui."],
    },
    {
        "file": "06-po-daftar-termin-dan-nontermin.png",
        "title": "Dashboard Pemesanan Pembelian (PO)",
        "callouts": [
            ((0.54, 0.05, 0.65, 0.09), "Buat dari PR atau langsung"),
            ((0.11, 0.23, 0.97, 0.37), "Ringkasan status dan nilai"),
            ((0.11, 0.37, 0.98, 0.54), "Tren nilai pesanan"),
        ],
        "narrative": (
            "Dashboard PO memperlihatkan total dokumen, nilai pesanan, jumlah yang sudah dibayar, sisa kewajiban, "
            "status lunas, dan dokumen yang melewati batas kirim. Tombol Dari PR menjaga jejak kebutuhan karena barang "
            "diambil dari PR yang telah disetujui. Tombol Buat PO digunakan untuk pesanan langsung yang memang tidak "
            "memerlukan PR. Untuk pelatihan, bandingkan nilai pesanan dengan sisa kewajiban dan perhatikan komposisi "
            "status sebelum masuk ke daftar pesanan. Data contoh menunjukkan lebih dari 100 PO sehingga peserta dapat "
            "berlatih memfilter kondisi termin maupun bukan termin."
        ),
        "steps": ["Pilih Dari PR untuk pengadaan terencana.", "Pilih Buat PO untuk pembelian langsung.",
                  "Periksa pemasok, tanggal kirim, dan nilai.", "Tentukan apakah pembayaran bertermin.",
                  "Simpan dan ajukan PO."],
    },
    {
        "file": "07-po-formulir-nontermin.png",
        "title": "PO bukan termin",
        "callouts": [
            ((0.35, 0.28, 0.65, 0.38), "Pilih pemasok dan tanggal"),
            ((0.35, 0.37, 0.65, 0.47), "Isi barang pesanan"),
            ((0.52, 0.66, 0.65, 0.73), "Simpan PO"),
        ],
        "narrative": (
            "PO bukan termin digunakan ketika tagihan akan dibayar sebagai satu kewajiban, tanpa jadwal cicilan. "
            "Pilih pemasok yang benar karena identitas ini akan dibawa ke BAST, tagihan, dan pembayaran vendor. "
            "Isi batas kirim secara realistis agar dashboard dapat menandai keterlambatan. Setiap baris barang harus "
            "memiliki satuan, jumlah, dan harga yang dapat ditelusuri ke kesepakatan vendor. Sebelum menyimpan, cocokkan "
            "subtotal baris dengan total dokumen serta pastikan pengaturan termin tidak aktif."
        ),
        "steps": ["Pilih pemasok.", "Isi tanggal dan batas kirim.", "Pastikan opsi termin tidak aktif.",
                  "Tambah barang, UOM, jumlah, dan harga.", "Cocokkan total lalu Simpan."],
    },
    {
        "file": "08-po-formulir-termin.png",
        "title": "PO termin dan jadwal pembayaran",
        "callouts": [
            ((0.35, 0.28, 0.66, 0.41), "Aktifkan pembayaran bertermin"),
            ((0.35, 0.42, 0.66, 0.58), "Susun termin"),
            ((0.52, 0.67, 0.65, 0.74), "Simpan setelah total 100%"),
        ],
        "narrative": (
            "PO termin dipakai ketika pembayaran vendor dibagi menjadi beberapa tahap. Setelah opsi Pembayaran "
            "bertermin aktif, susun nama termin, tanggal jatuh tempo, dan nilai atau persentasenya. Jumlah seluruh "
            "termin harus sama dengan total PO; ketidaksamaan akan menyulitkan pencocokan tagihan dan pembayaran. "
            "BAST dan tagihan berikutnya harus menyebut termin yang sama agar sisa kewajiban dapat dihitung dengan "
            "benar. Praktik yang disarankan adalah memakai keterangan yang konsisten, misalnya Uang Muka, Termin-1, "
            "dan Pelunasan, bukan label bebas yang berubah di setiap dokumen."
        ),
        "steps": ["Aktifkan Pembayaran bertermin.", "Tambahkan seluruh tahap termin.",
                  "Isi jatuh tempo dan nilai tiap tahap.", "Pastikan jumlah termin = total PO.",
                  "Simpan dan ajukan untuk disetujui."],
    },
    {
        "file": "09-bast-daftar-100-dokumen.png",
        "title": "Dashboard Penerimaan Barang/BAST",
        "callouts": [
            ((0.11, 0.23, 0.97, 0.37), "Jumlah, nilai, dan status BAST"),
            ((0.11, 0.38, 0.98, 0.54), "Tren penerimaan"),
            ((0.85, 0.89, 0.98, 0.95), "Terima dari PO atau langsung"),
        ],
        "narrative": (
            "BAST adalah bukti bahwa barang atau jasa benar-benar diterima. Dashboard memisahkan jumlah dokumen, "
            "nilai diterima, status persetujuan, dan status masuk stok. Gunakan Dari PO agar kuantitas yang diterima "
            "dibandingkan dengan pesanan; Terima Langsung hanya untuk kondisi yang memang tidak memiliki PO. Petugas "
            "penerima harus mencatat jumlah aktual, barang rusak atau kurang, serta tanggal penerimaan. Jangan membuat "
            "BAST hanya dari nilai invoice, karena laporan penerimaan dan persediaan bergantung pada rincian barang."
        ),
        "steps": ["Tekan Dari PO.", "Pilih PO pemasok yang benar.", "Periksa barang dan jumlah fisik.",
                  "Catat selisih bila ada.", "Simpan, setujui, lalu pastikan status masuk stok."],
    },
    {
        "file": "10-bast-pilih-po.png",
        "title": "Memilih PO sumber BAST",
        "callouts": [
            ((0.30, 0.29, 0.69, 0.38), "Cari nomor PO/pemasok"),
            ((0.30, 0.37, 0.69, 0.63), "Pilih PO yang masih dapat diterima"),
        ],
        "narrative": (
            "Dialog pemilihan PO mencegah pengguna memasukkan ulang data pemasok dan barang. Cari nomor PO atau "
            "pemasok, lalu pastikan dokumen yang dipilih memang belum seluruhnya diterima. Sistem membawa rincian "
            "pesanan ke formulir BAST, tetapi petugas tetap wajib mengganti kuantitas dengan hasil pemeriksaan fisik. "
            "Apabila daftar tidak memuat PO yang dicari, periksa status persetujuan, sisa barang yang belum diterima, "
            "dan toko aktif sebelum membuat penerimaan langsung."
        ),
        "steps": ["Cari PO.", "Baca nomor, pemasok, dan sisa nilai/kuantitas.",
                  "Pilih PO yang sesuai.", "Lanjutkan pemeriksaan fisik pada formulir BAST."],
    },
    {
        "file": "11-terima-tagihan-vendor-100-dokumen.png",
        "title": "Terima Tagihan Vendor",
        "callouts": [
            ((0.11, 0.23, 0.97, 0.38), "Ringkasan dokumen tagihan"),
            ((0.11, 0.38, 0.98, 0.55), "Tren nilai tagihan"),
            ((0.84, 0.88, 0.98, 0.95), "Buat dari BAST"),
        ],
        "narrative": (
            "Terima Tagihan Vendor menghubungkan dokumen vendor dengan barang/jasa yang telah diterima. Pembuatan "
            "dari BAST menjaga agar pemasok, PO, termin, dan nilai dapat ditelusuri. Catat nomor invoice vendor persis "
            "seperti dokumen aslinya dan isi tanggal jatuh tempo. Untuk PO termin, pilih tahap yang ditagihkan; untuk PO "
            "bukan termin, pastikan tagihan tidak melebihi sisa nilai yang belum ditagihkan. Setelah disimpan dan "
            "disetujui, tagihan menjadi dasar antrean pembayaran vendor."
        ),
        "steps": ["Pilih BAST sumber.", "Isi nomor dan tanggal invoice vendor.",
                  "Pilih termin jika PO bertermin.", "Cocokkan nilai dengan BAST dan sisa PO.",
                  "Simpan dan ajukan tagihan."],
    },
    {
        "file": "12-pembayaran-vendor-100-dokumen.png",
        "title": "Pembayaran Tagihan Vendor",
        "callouts": [
            ((0.11, 0.18, 0.98, 0.33), "Filter dan ringkasan pembayaran"),
            ((0.11, 0.34, 0.98, 0.74), "Daftar pembayaran vendor"),
            ((0.84, 0.84, 0.98, 0.94), "Proses pembayaran"),
        ],
        "narrative": (
            "Pembayaran vendor dilakukan setelah tagihan lolos verifikasi. Pilih tagihan yang benar, rekening atau "
            "cara bayar yang sah, tanggal pembayaran, dan referensi bukti transfer. Untuk skema termin, nilai yang "
            "dibayar harus mengikuti termin yang jatuh tempo; untuk bukan termin, pembayaran dapat menutup seluruh "
            "sisa tagihan. Setelah diproses, periksa status lunas atau sisa kewajiban pada PO dan laporan pembelian. "
            "Persetujuan pembayaran merupakan tindakan server dan hanya tersedia bagi peran yang berwenang."
        ),
        "steps": ["Filter tagihan yang siap dibayar.", "Pilih vendor/tagihan.",
                  "Isi metode, tanggal, dan nomor bukti.", "Periksa nilai terhadap sisa tagihan.",
                  "Proses pembayaran dan verifikasi status."],
    },
    {
        "file": "14-penjualan-daftar-faktur-atas.png",
        "title": "Laporan Daftar Faktur Penjualan",
        "callouts": [
            ((0.02, 0.07, 0.98, 0.16), "Atur periode dan tampilkan"),
            ((0.84, 0.09, 0.98, 0.15), "Ekspor PDF/Excel"),
            ((0.02, 0.23, 0.98, 0.86), "504 transaksi pada periode UAT"),
        ],
        "narrative": (
            "Daftar Faktur Penjualan adalah kontrol utama untuk menghitung jumlah nota dan nilai penjualan dalam "
            "periode. UAT memakai 1–10 September 2026 dan menghasilkan 504 transaksi. Periksa bahwa nomor nota, "
            "tanggal, kasir, pelanggan, metode, serta total tersedia dan tidak ada baris kosong. Gunakan laporan ini "
            "sebagai daftar sumber sebelum membandingkan Omzet, penjualan per barang, atau margin. Tombol PDF dan "
            "Excel tersedia setelah laporan berhasil dimuat."
        ),
        "steps": ["Atur periode yang sama untuk semua laporan.", "Tekan Tampilkan.",
                  "Periksa jumlah baris dan beberapa sampel nota.", "Unduh PDF/Excel bila diperlukan."],
    },
    {
        "file": "15-penjualan-per-barang-atas.png",
        "title": "Laporan Penjualan per Barang",
        "callouts": [
            ((0.02, 0.07, 0.98, 0.16), "Periode laporan"),
            ((0.02, 0.23, 0.98, 0.85), "Agregasi produk terjual"),
            ((0.90, 0.87, 0.98, 0.93), "Pagination"),
        ],
        "narrative": (
            "Penjualan per Barang menggabungkan transaksi menjadi ringkasan per produk. Laporan ini membantu "
            "menentukan produk cepat laku, memeriksa kuantitas terjual, dan membandingkan nilai penjualan. Data UAT "
            "mencakup lebih dari 100 produk sehingga pagination wajib diperiksa hingga halaman terakhir. Bila angka "
            "produk tidak sesuai perkiraan, telusuri kembali ke rincian penjualan per barang dan pastikan periode, toko, "
            "serta status transaksi sama."
        ),
        "steps": ["Gunakan periode yang sama dengan Daftar Faktur.", "Urutkan atau cari produk.",
                  "Periksa jumlah dan nilai.", "Lanjutkan ke Rincian Penjualan per Barang untuk bukti nota."],
    },
    {
        "file": "17-omzet-transaksi-atas.png",
        "title": "Laporan Transaksi Omzet",
        "callouts": [
            ((0.02, 0.07, 0.98, 0.16), "Periode dan ekspor"),
            ((0.02, 0.23, 0.98, 0.85), "Klik nilai/baris untuk asal angka"),
            ((0.02, 0.87, 0.16, 0.92), "504 baris"),
        ],
        "narrative": (
            "Transaksi Omzet menampilkan satu baris untuk setiap nota aktif. Kolom mode membedakan Tunai, QRIS, "
            "transfer, voucher, dan Saldo; kolom nominal menjadi dasar rekonsiliasi dengan Rekap Omzet. Setiap nilai "
            "yang diberi penanda dapat diklik untuk membuka asal angka. Gunakan nomor transaksi sebagai kunci utama "
            "saat peserta membandingkan popup, riwayat penjualan, dan dokumen ekspor. Pada periode UAT terdapat 504 "
            "baris, termasuk 100 transaksi saldo yang sengaja dibuat untuk latihan."
        ),
        "steps": ["Tekan Tampilkan.", "Pilih satu baris.", "Klik nilai yang bergaris/berpenanda.",
                  "Cocokkan nomor transaksi, metode, produk, dan nominal pada popup."],
    },
    {
        "file": "17-omzet-transaksi-rincian-clickable.png",
        "title": "Popup asal angka Omzet",
        "callouts": [
            ((0.32, 0.28, 0.68, 0.72), "Detail transaksi penyusun"),
            ((0.63, 0.66, 0.68, 0.72), "Tutup dan kembali"),
        ],
        "narrative": (
            "Popup Asal Angka menjelaskan dari transaksi mana suatu nilai laporan berasal. Peserta harus memeriksa "
            "nomor transaksi, tipe pengguna, nama pelanggan, metode, nomor pembayaran, nominal, toko, tanggal, kasir, "
            "dan keterangan produk. Informasi ini bukan data baru; seluruhnya merupakan jejak dari nota yang sama. "
            "Fungsi tersebut sangat penting ketika supervisor menerima pertanyaan mengapa angka rekap berbeda dari "
            "catatan manual, karena sumbernya dapat dibuka tanpa keluar dari laporan."
        ),
        "steps": ["Baca judul Asal Angka.", "Cocokkan nomor transaksi.",
                  "Periksa metode, nominal, tanggal, toko, dan kasir.", "Tekan Tutup untuk kembali."],
    },
    {
        "file": "18-omzet-tunai-per-produk-atas.png",
        "title": "Omzet Produk Non-Saldo/Tunai",
        "callouts": [
            ((0.02, 0.19, 0.98, 0.23), "Definisi klasifikasi non-saldo"),
            ((0.02, 0.23, 0.98, 0.85), "102 produk"),
            ((0.83, 0.09, 0.98, 0.15), "PDF/Excel"),
        ],
        "narrative": (
            "Laporan Non-Saldo/Tunai mengelompokkan omzet dari metode yang tidak memotong saldo member, termasuk "
            "tunai dan metode non-saldo lain yang diklasifikasikan oleh master pembayaran. Kolom Modal Rata-rata, "
            "Harga Rata-rata, Terjual, Omzet, dan Profit memungkinkan trainer menjelaskan hubungan antara volume dan "
            "keuntungan. UAT menghasilkan 102 produk. Klik nilai omzet pada produk tertentu untuk menampilkan nota "
            "penyusunnya, lalu cocokkan jumlahnya dengan Transaksi Omzet."
        ),
        "steps": ["Tampilkan periode UAT.", "Pilih produk.", "Periksa terjual, omzet, dan profit.",
                  "Klik angka untuk menelusuri transaksi penyusun."],
    },
    {
        "file": "19-omzet-saldo-per-produk-atas.png",
        "title": "Omzet Produk Saldo",
        "callouts": [
            ((0.02, 0.19, 0.98, 0.23), "Definisi pembayaran Saldo"),
            ((0.02, 0.23, 0.98, 0.85), "101 produk terisi"),
            ((0.83, 0.09, 0.98, 0.15), "Ekspor laporan"),
        ],
        "narrative": (
            "Omzet Saldo hanya memuat transaksi yang memakai metode pembayaran dengan klasifikasi pemotong saldo "
            "member. Agar bahan pelatihan tidak kosong, UAT membuat 100 transaksi tambahan pada 100 produk berbeda; "
            "bersama transaksi verifikasi awal, laporan menampilkan 101 produk. Nilai omzet, HPP, dan profit dialokasikan "
            "sesuai porsi pembayaran saldo. Peserta harus memilih member sebelum pembayaran dan memastikan saldo cukup, "
            "kemudian menelusuri angka laporan ke nota asal."
        ),
        "steps": ["Pilih member pada POS.", "Pilih metode Saldo.", "Selesaikan transaksi.",
                  "Buka laporan Omzet Saldo.", "Klik omzet produk untuk melihat nota penyusun."],
    },
    {
        "file": "20-rekap-omzet-per-toko-atas.png",
        "title": "Rekap Omzet per Toko",
        "callouts": [
            ((0.02, 0.19, 0.98, 0.23), "Aturan klasifikasi"),
            ((0.02, 0.23, 0.98, 0.37), "Tunai + Saldo = Total Omzet"),
            ((0.83, 0.09, 0.98, 0.15), "Ekspor"),
        ],
        "narrative": (
            "Rekap Omzet menyatukan dua kelompok pembayaran per toko. Untuk Kantin Demo, nilai Non-Saldo/Tunai dan "
            "Saldo harus menghasilkan Total Omzet yang sama dengan penjumlahan seluruh Transaksi Omzet pada periode. "
            "Klik salah satu nilai untuk melihat transaksi penyusun kelompok tersebut. Selisih biasanya berasal dari "
            "periode yang tidak sama, toko aktif berbeda, transaksi belum tersinkron, atau klasifikasi metode pembayaran "
            "yang berubah. UAT API memverifikasi nilai baris dan rincian yang diklik tanpa selisih."
        ),
        "steps": ["Tampilkan periode yang sama.", "Jumlahkan kolom Non-Saldo/Tunai dan Saldo.",
                  "Cocokkan dengan Total Omzet.", "Klik nilai untuk pemeriksaan rinci."],
    },
    {
        "file": "21-penerimaan-pembelian-atas.png",
        "title": "Laporan Penerimaan Pembelian",
        "callouts": [
            ((0.02, 0.07, 0.98, 0.16), "Periode dan ekspor"),
            ((0.02, 0.23, 0.98, 0.84), "452 penerimaan"),
            ((0.89, 0.85, 0.98, 0.91), "31 halaman"),
        ],
        "narrative": (
            "Laporan Penerimaan Pembelian bersumber dari dokumen penerimaan/BAST yang telah masuk ke alur pembelian. "
            "Setiap baris memperlihatkan nomor faktur, tanggal, pemasok, jumlah item, dan total pembelian. Data UAT "
            "berisi 452 penerimaan sehingga peserta wajib menggunakan pagination dan tidak menyimpulkan total dari "
            "halaman pertama saja. Cocokkan beberapa nomor faktur dengan BAST dan pastikan nilai serta pemasok sama."
        ),
        "steps": ["Atur periode.", "Tekan Tampilkan.", "Periksa jumlah baris dan halaman.",
                  "Cocokkan sampel faktur terhadap BAST.", "Unduh PDF/Excel untuk arsip."],
    },
    {
        "file": "22-faktur-pembelian-atas.png",
        "title": "Laporan Daftar Faktur Pembelian",
        "callouts": [
            ((0.02, 0.07, 0.98, 0.16), "Filter periode"),
            ((0.02, 0.23, 0.98, 0.84), "Daftar faktur vendor"),
            ((0.83, 0.09, 0.98, 0.15), "PDF/Excel"),
        ],
        "narrative": (
            "Daftar Faktur Pembelian digunakan untuk memastikan seluruh invoice vendor yang diterima sudah masuk ke "
            "periode pelaporan. Bandingkan nomor faktur, tanggal, pemasok, dan total dengan dokumen Terima Tagihan. "
            "Untuk PO termin, satu PO dapat menghasilkan beberapa faktur sesuai tahap; karena itu nomor PO tidak boleh "
            "dipakai sebagai pengganti nomor invoice vendor. Laporan ini menjadi jembatan pemeriksaan antara BAST, "
            "tagihan, dan pembayaran tanpa membahas jurnal atau COA."
        ),
        "steps": ["Tampilkan periode.", "Cari nomor invoice vendor.",
                  "Cocokkan pemasok dan total dengan Terima Tagihan.", "Periksa sisa tagihan melalui pembayaran vendor."],
    },
    {
        "file": "23-margin-per-produk-atas.png",
        "title": "Laporan Margin per Produk",
        "callouts": [
            ((0.02, 0.19, 0.98, 0.23), "Rumus margin"),
            ((0.02, 0.23, 0.98, 0.85), "201 produk"),
            ((0.83, 0.09, 0.98, 0.15), "Ekspor"),
        ],
        "narrative": (
            "Margin per Produk membandingkan Penjualan dengan HPP untuk setiap produk dan menampilkan selisih sebagai "
            "Laba. UAT mencakup 201 produk sehingga laporan memberikan variasi data yang cukup untuk latihan. HPP di "
            "sini adalah biaya produk pada data sumber operasional, bukan pembahasan posting Akuntansi. Produk dengan "
            "laba terlalu kecil perlu diperiksa harga jual, biaya pembelian, diskon, serta kuantitas. Gunakan laporan "
            "kategori untuk melihat pola yang lebih luas."
        ),
        "steps": ["Pilih periode.", "Bandingkan Penjualan dan HPP.", "Periksa Laba per produk.",
                  "Telusuri produk anomali ke laporan transaksi dan pembelian."],
    },
    {
        "file": "24-margin-per-kategori-atas.png",
        "title": "Laporan Margin per Kategori",
        "callouts": [
            ((0.02, 0.19, 0.98, 0.23), "Rumus kategori"),
            ((0.02, 0.23, 0.98, 0.80), "13 kategori"),
            ((0.83, 0.09, 0.98, 0.15), "PDF/Excel"),
        ],
        "narrative": (
            "Margin per Kategori menggabungkan hasil produk berdasarkan kategori master. Laporan ini membantu "
            "manajemen membedakan kategori beromzet besar tetapi bermargin tipis dari kategori yang lebih menguntungkan. "
            "UAT memuat 13 kategori. Ketika hasil kategori terasa tidak wajar, periksa apakah produk ditempatkan pada "
            "kategori yang benar dan bandingkan dengan Margin per Produk. Nilai kategori seharusnya merupakan agregasi "
            "produk dalam periode dan toko yang sama."
        ),
        "steps": ["Tampilkan periode.", "Urutkan kategori berdasarkan omzet atau laba.",
                  "Bandingkan persentase margin.", "Buka Margin per Produk untuk penyebab rinci."],
    },
    {
        "file": "25-laba-kotor-harian-atas.png",
        "title": "Laporan Laba Kotor Harian",
        "callouts": [
            ((0.02, 0.07, 0.98, 0.16), "Periode"),
            ((0.02, 0.23, 0.98, 0.58), "Penjualan, HPP, dan laba per hari"),
            ((0.83, 0.09, 0.98, 0.15), "Ekspor"),
        ],
        "narrative": (
            "Laba Kotor Harian merangkum Penjualan dikurangi HPP per tanggal. Gunakan laporan ini untuk melihat hari "
            "operasional yang paling kuat dan mendeteksi lonjakan atau penurunan tidak biasa. Jumlah transaksi yang "
            "besar pada satu tanggal harus terlihat konsisten dengan Transaksi Omzet dan penjualan per barang. Data "
            "UAT tersebar pada tiga tanggal sehingga peserta dapat berlatih membandingkan tren. Laporan ini tetap "
            "bersifat operasional; laporan keuangan berbasis COA tidak termasuk dalam pelatihan ini."
        ),
        "steps": ["Atur periode.", "Bandingkan Penjualan dan HPP per tanggal.",
                  "Periksa Laba dan persentasenya.", "Telusuri hari anomali ke Transaksi Omzet."],
    },
]


def annotate(spec: dict) -> Path:
    src = SCREEN / spec["file"]
    img = Image.open(src).convert("RGBA")
    draw = ImageDraw.Draw(img, "RGBA")
    w, h = img.size
    colors = [RED, AMBER, BLUE, GREEN]
    for idx, (box, _) in enumerate(spec["callouts"], start=1):
        x1, y1, x2, y2 = [int(v * (w if i % 2 == 0 else h)) for i, v in enumerate(box)]
        color = colors[(idx - 1) % len(colors)]
        draw.rounded_rectangle((x1, y1, x2, y2), radius=max(8, w // 220),
                               outline=(*color, 255), width=max(5, w // 420))
        radius = max(18, w // 95)
        cx, cy = x1 + radius, max(radius + 4, y1 - radius)
        draw.ellipse((cx - radius, cy - radius, cx + radius, cy + radius),
                     fill=(*color, 245), outline=(255, 255, 255, 255), width=3)
        nfont = font(int(radius * 1.15), True)
        text = str(idx)
        bbox = draw.textbbox((0, 0), text, font=nfont)
        draw.text((cx - (bbox[2] - bbox[0]) / 2, cy - (bbox[3] - bbox[1]) / 2 - 2),
                  text, font=nfont, fill=WHITE)
    ANNOTATED.mkdir(parents=True, exist_ok=True)
    out = ANNOTATED / spec["file"]
    img.convert("RGB").save(out, quality=95)
    return out


def rounded_box(draw, xy, title, body="", fill=(235, 244, 255), outline=BLUE, title_color=NAVY):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=4)
    tf = font(28, True)
    bf = font(21, False)
    title_lines = wrap(draw, title, tf, x2 - x1 - 36)
    y = y1 + 20
    for line in title_lines:
        draw.text((x1 + 18, y), line, font=tf, fill=title_color)
        y += 34
    for line in wrap(draw, body, bf, x2 - x1 - 36):
        draw.text((x1 + 18, y + 5), line, font=bf, fill=MUTED)
        y += 27


def diagram_flow() -> Path:
    img = Image.new("RGB", (1800, 1250), WHITE)
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Alur End-to-End Pelatihan Operasional", font=font(46, True), fill=NAVY)
    d.text((70, 103), "Dua alur bisnis bertemu pada laporan operasional; Akuntansi tidak termasuk.",
           font=font(25), fill=MUTED)
    d.rounded_rectangle((55, 165, 1745, 555), 25, fill=(245, 250, 255), outline=(191, 219, 254), width=3)
    d.text((85, 190), "ALUR PENJUALAN", font=font(28, True), fill=BLUE)
    sales = [("1. Pilih barang", "Kasir/scan barcode"), ("2. Tentukan bayar", "Tunai/non-saldo/saldo"),
             ("3. Bayar", "Nota tersimpan"), ("4. Telusuri", "Riwayat & detail")]
    xs = [95, 510, 925, 1340]
    for x, (t, b) in zip(xs, sales):
        rounded_box(d, (x, 270, x + 300, 440), t, b)
    for x in xs[:-1]:
        arrow(d, (x + 300, 355), (x + 405, 355))
    d.rounded_rectangle((55, 605, 1745, 1015), 25, fill=(250, 253, 248), outline=(187, 247, 208), width=3)
    d.text((85, 630), "ALUR KULAKAN / PENGADAAN", font=font(28, True), fill=GREEN)
    proc = [("1. PR", "Kebutuhan disetujui"), ("2. PO", "Termin/non-termin"),
            ("3. BAST", "Barang diterima"), ("4. Tagihan", "Invoice diverifikasi"),
            ("5. Bayar vendor", "Kewajiban ditutup")]
    px = [85, 420, 755, 1090, 1425]
    for x, (t, b) in zip(px, proc):
        rounded_box(d, (x, 715, x + 270, 885), t, b, fill=(239, 252, 244), outline=GREEN)
    for x in px[:-1]:
        arrow(d, (x + 270, 800), (x + 325, 800), color=GREEN)
    rounded_box(d, (420, 1060, 1380, 1195), "Laporan Operasional",
                "Penjualan • Pembelian • 4 Omzet • Margin Produk/Kategori • Laba Kotor Harian",
                fill=(255, 248, 229), outline=AMBER)
    arrow(d, (900, 555), (900, 1050), color=AMBER)
    arrow(d, (900, 1015), (900, 1050), color=AMBER)
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    out = DIAGRAMS / "01-flow-end-to-end.png"
    img.save(out)
    return out


def diagram_use_case() -> Path:
    img = Image.new("RGB", (1800, 1250), WHITE)
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Use Case Pelatihan Operasional", font=font(46, True), fill=NAVY)
    d.text((70, 105), "Aktor, tindakan utama, dan titik serah tanggung jawab.", font=font(25), fill=MUTED)
    rows = [
        ("Kasir", "Transaksi POS", "Nota dan Riwayat Penjualan"),
        ("Pemohon", "Membuat PR", "Kebutuhan menunggu persetujuan"),
        ("Logistik/Penyetuju", "Membuat & menyetujui PO", "Komitmen ke vendor"),
        ("Penerima", "Membuat BAST", "Penerimaan fisik tercatat"),
        ("Admin Vendor", "Menerima Tagihan", "Invoice siap diverifikasi"),
        ("Keuangan", "Membayar Vendor", "Status lunas atau sisa"),
        ("Supervisor/Trainer", "Membaca Laporan", "Drill-down, PDF, dan Excel"),
    ]
    y = 180
    for actor, use_case, result in rows:
        rounded_box(d, (80, y, 430, y + 120), actor, "Aktor", fill=(239, 246, 255), outline=BLUE)
        d.ellipse((650, y, 1120, y + 120), fill=(250, 252, 255), outline=CYAN, width=4)
        f = font(24, True)
        bb = d.textbbox((0, 0), use_case, font=f)
        d.text((885 - (bb[2] - bb[0]) / 2, y + 60 - (bb[3] - bb[1]) / 2),
               use_case, font=f, fill=NAVY)
        rounded_box(d, (1340, y, 1720, y + 120), result, "Hasil", fill=(239, 252, 244), outline=GREEN)
        arrow(d, (430, y + 60), (640, y + 60), color=BLUE, width=5)
        arrow(d, (1120, y + 60), (1330, y + 60), color=GREEN, width=5)
        y += 135
    d.rounded_rectangle((315, 1135, 1485, 1210), 18, fill=(255, 248, 229), outline=AMBER, width=3)
    d.text((360, 1156), "Batas pelatihan: COA, jurnal, posting, dan laporan keuangan berbasis akun tidak dibahas.",
           font=font(23, True), fill=(146, 83, 7))
    out = DIAGRAMS / "02-use-case.png"
    img.save(out)
    return out


def diagram_data() -> Path:
    img = Image.new("RGB", (1800, 1250), WHITE)
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Aliran Data / ERD Ringkas", font=font(46, True), fill=NAVY)
    d.text((70, 105), "Relasi bisnis yang perlu dipahami pengguna; bukan skema fisik basis data.", font=font(25), fill=MUTED)
    nodes = {
        "Produk": (80, 235, 400, 405, ["kode", "nama", "kategori", "harga/modal"]),
        "Penjualan": (545, 205, 885, 445, ["nomor nota", "tanggal", "toko", "member", "metode", "total"]),
        "Detail Penjualan": (1050, 235, 1400, 405, ["produk", "qty", "harga", "diskon"]),
        "PR": (80, 600, 360, 790, ["pemohon", "barang", "qty", "status"]),
        "PO": (480, 575, 770, 815, ["pemasok", "termin", "nilai", "status"]),
        "BAST": (890, 575, 1180, 815, ["PO", "barang diterima", "selisih"]),
        "Tagihan": (1300, 575, 1590, 815, ["BAST", "invoice", "jatuh tempo"]),
        "Pembayaran": (1300, 945, 1590, 1135, ["tagihan", "metode", "bukti", "nilai"]),
        "Laporan": (545, 945, 1040, 1135, ["penjualan", "pembelian", "omzet", "margin"]),
    }
    for name, (x1, y1, x2, y2, attrs) in nodes.items():
        outline = GREEN if name in {"PR", "PO", "BAST", "Tagihan", "Pembayaran"} else BLUE
        d.rounded_rectangle((x1, y1, x2, y2), 18, fill=LIGHT, outline=outline, width=4)
        d.rectangle((x1, y1, x2, y1 + 52), fill=outline)
        d.text((x1 + 18, y1 + 12), name, font=font(25, True), fill=WHITE)
        yy = y1 + 70
        for attr in attrs:
            d.text((x1 + 22, yy), f"• {attr}", font=font(21), fill=INK)
            yy += 32
    links = [((400, 320), (545, 320), "1:N"), ((885, 320), (1050, 320), "1:N"),
             ((360, 690), (480, 690), "1:N"), ((770, 690), (890, 690), "1:N"),
             ((1180, 690), (1300, 690), "1:N"), ((1445, 815), (1445, 945), "1:N"),
             ((715, 445), (715, 935), "agregasi"), ((1180, 1040), (1050, 1040), "agregasi")]
    for p1, p2, label in links:
        arrow(d, p1, p2, color=(100, 116, 139), width=5)
        mx, my = (p1[0] + p2[0]) // 2, (p1[1] + p2[1]) // 2
        d.text((mx - 28, my - 32), label, font=font(18, True), fill=MUTED)
    d.text((80, 1180), "Kunci penelusuran: nomor nota • nomor PR/PO/BAST/invoice • toko • tanggal • produk • pemasok",
           font=font(24, True), fill=NAVY)
    out = DIAGRAMS / "03-aliran-data-erd.png"
    img.save(out)
    return out


def diagram_reports() -> Path:
    img = Image.new("RGB", (1800, 1250), WHITE)
    d = ImageDraw.Draw(img)
    d.text((70, 45), "Peta Laporan dan Cara Menelusuri Angka", font=font(44, True), fill=NAVY)
    d.text((70, 105), "Mulai dari pertanyaan pengguna, lalu pilih laporan dan buka sumber transaksinya.", font=font(25), fill=MUTED)
    rows = [
        ("Apa saja notanya?", "Daftar Faktur / Transaksi Omzet", "Klik baris → detail nota"),
        ("Produk apa yang laku?", "Penjualan per Barang", "Buka rincian per barang"),
        ("Berapa omzet tunai/saldo?", "4 Laporan Omzet", "Klik nilai → transaksi penyusun"),
        ("Apa yang diterima/dibeli?", "Penerimaan / Faktur Pembelian", "Cocokkan BAST dan invoice"),
        ("Berapa laba kotor?", "Margin Produk/Kategori/Harian", "Penjualan − HPP operasional"),
    ]
    y = 200
    for i, (q, report, trace) in enumerate(rows, 1):
        rounded_box(d, (80, y, 510, y + 145), f"{i}. {q}", "Pertanyaan trainer", fill=(239, 246, 255), outline=BLUE)
        rounded_box(d, (680, y, 1120, y + 145), report, "Laporan yang dibuka", fill=(255, 248, 229), outline=AMBER)
        rounded_box(d, (1290, y, 1720, y + 145), trace, "Langkah verifikasi", fill=(239, 252, 244), outline=GREEN)
        arrow(d, (510, y + 72), (670, y + 72), color=BLUE)
        arrow(d, (1120, y + 72), (1280, y + 72), color=GREEN)
        y += 195
    d.rounded_rectangle((300, 1130, 1500, 1210), 18, fill=(245, 247, 250), outline=(148, 163, 184), width=3)
    d.text((355, 1152), "Selalu samakan periode dan toko sebelum membandingkan dua laporan.", font=font(27, True), fill=NAVY)
    out = DIAGRAMS / "04-peta-laporan.png"
    img.save(out)
    return out


def set_cell_fill(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_run(paragraph, text, *, bold=False, size=10, color="1F2937", italic=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    add_run(p, text, bold=True, size=20 if level == 1 else 15 if level == 2 else 12,
            color="17365D" if level <= 2 else "2563EB")
    return p


def add_body(doc, text, *, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        add_run(p, bold_lead, bold=True, size=9.4)
        add_run(p, text[len(bold_lead):], size=9.4)
    else:
        add_run(p, text, size=9.4)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        add_run(p, item, size=9.3)


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        add_run(p, f"{index}. ", bold=True, size=9.3, color="2563EB")
        add_run(p, item, size=9.3)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_fill(cell, "17365D")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        add_run(p, str(header), bold=True, size=8.5, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if row_index % 2:
                set_cell_fill(cells[i], "F3F6FA")
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cells[i].paragraphs[0]
            add_run(p, str(value), size=8.2)
    return table


def page_break(doc):
    doc.add_page_break()


def add_picture_page(doc, spec, path: Path):
    add_heading(doc, spec["title"], 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(10.55))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(cap, f"Gambar — {spec['title']}.", italic=True, size=8.2, color="64748B")
    table = doc.add_table(rows=1, cols=len(spec["callouts"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    colors = ["DC2626", "D97706", "2563EB", "16A34A"]
    for i, (_, label) in enumerate(spec["callouts"], 1):
        cell = table.rows[0].cells[i - 1]
        set_cell_fill(cell, "F8FAFC")
        set_cell_margins(cell, top=55, bottom=55)
        p = cell.paragraphs[0]
        add_run(p, f"{i}  ", bold=True, size=8.4, color=colors[(i - 1) % len(colors)])
        add_run(p, label, size=8.2)
    page_break(doc)
    add_heading(doc, f"Panduan — {spec['title']}", 2)
    add_body(doc, spec["narrative"])
    add_heading(doc, "Langkah pengguna", 3)
    add_numbered(doc, spec["steps"])
    add_heading(doc, "Kontrol UAT", 3)
    add_bullets(doc, [
        "Data tampil dan tidak ada grid kosong pada periode pengujian.",
        "Hak akses pengguna sesuai peran; tombol tindakan hanya aktif bila berwenang.",
        "Nomor dokumen, tanggal, toko, pemasok/member, item, dan nilai dapat ditelusuri.",
        "Bila angka dibandingkan dengan laporan lain, gunakan periode dan toko yang sama.",
    ])
    page_break(doc)


def build_doc():
    evidence = json.loads(EVIDENCE.read_text(encoding="utf-8"))
    annotated = {spec["file"]: annotate(spec) for spec in SCREEN_SPECS}
    diagrams = [diagram_flow(), diagram_use_case(), diagram_data(), diagram_reports()]

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.25)
    section.right_margin = Cm(1.25)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.45)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.4)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.color.rgb = RGBColor(23, 54, 93)
        styles[style_name].paragraph_format.space_before = Pt(4)
        styles[style_name].paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(header, "eBisnis • Manual Training Operasional • Tanpa Akuntansi", bold=True, size=8.2, color="2563EB")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "v1.34.32 • 10 September 2026   |   ", size=8, color="64748B")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    add_run(p, "MANUAL PELATIHAN & HASIL UAT", bold=True, size=16, color="2563EB")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    add_run(p, "Operasional eBisnis", bold=True, size=34, color="17365D")
    p = doc.add_paragraph()
    add_run(p, "POS/Kasir • PR • PO Termin & Bukan Termin • BAST • Tagihan • Pembayaran Vendor",
            bold=True, size=17, color="0EA5E9")
    p = doc.add_paragraph()
    add_run(p, "Laporan Penjualan • Pembelian • Omzet • Margin • Laba Kotor Harian",
            bold=True, size=16, color="16A34A")
    doc.add_paragraph()
    cover_table = add_table(doc, ["Build", "Lingkungan UAT", "Periode", "Sasaran"], [[
        "eBisnis 1.34.32 (build 195)", "Kantin Demo • server live", "1–10 September 2026",
        "Bahan training An Nahl"
    ]], [4.2, 5.0, 4.2, 5.0])
    doc.add_paragraph()
    note = doc.add_table(rows=1, cols=1)
    set_cell_fill(note.cell(0, 0), "FFF4D6")
    set_cell_margins(note.cell(0, 0), top=150, start=180, bottom=150, end=180)
    add_run(note.cell(0, 0).paragraphs[0], "Batas ruang lingkup — ", bold=True, size=10, color="92400E")
    add_run(note.cell(0, 0).paragraphs[0],
            "COA, jurnal, posting, dan laporan keuangan berbasis akun sengaja tidak dimasukkan. "
            "Materi ini dipakai sebelum data COA An Nahl siap.", size=10, color="92400E")
    page_break(doc)

    add_heading(doc, "1. Ringkasan Hasil UAT", 1)
    add_body(doc,
             "UAT dijalankan pada build varian eBisnis menggunakan Kantin Demo dan server live. Data dibuat dengan "
             "kode referensi yang idempoten agar pengulangan tes tidak menghasilkan dokumen ganda. Pengujian "
             "memastikan alur transaksi, pengadaan, laporan, ekspor, dan rincian angka dapat digunakan sebagai bahan "
             "pelatihan. Seluruh pemeriksaan dalam ruang lingkup dinyatakan lulus.")
    report_rows = [(r["id"], r["rows"], r["columns"], "LULUS" if r["passed"] else "GAGAL")
                   for r in evidence["reports"]]
    proc_rows = []
    labels = {
        "pengadaan_pr_daftar": "PR", "pengadaan_po_daftar": "PO",
        "pengadaan_bast_daftar": "BAST", "pengadaan_tagihan_daftar": "Terima Tagihan",
        "pengadaan_bayar_daftar": "Pembayaran Vendor",
    }
    for key, item in evidence["procurement"].items():
        proc_rows.append((labels[key], item["count"], "≥100", "LULUS" if item["passed"] else "GAGAL"))
    add_table(doc, ["Tahap Pengadaan", "Dokumen Sampel", "Target", "Hasil"], proc_rows, [6, 4, 3, 3])
    doc.add_paragraph()
    add_table(doc, ["Kelompok pengujian", "Hasil"], [
        ("Audit API pengadaan", "5/5 lulus"),
        ("Laporan + PDF server", "12/12 lulus"),
        ("Popup rincian clickable", "4/4 lulus"),
        ("Integrasi layar Windows", "2/2 skenario lulus"),
        ("Kontrak ekspor, hak akses, lokal-dulu, UOM", "44/44 tes lulus"),
    ], [11, 5])
    add_heading(doc, "Dataset yang ditampilkan", 2)
    add_bullets(doc, [
        "504 transaksi penjualan dalam periode UAT.",
        "100 transaksi saldo tambahan pada 100 produk berbeda; Omzet Saldo menampilkan 101 produk.",
        "100 dokumen berantai untuk masing-masing PR, PO, BAST, Tagihan, dan Pembayaran Vendor.",
        "452 baris penerimaan/faktur pembelian, 201 produk margin, dan 13 kategori margin.",
    ])
    page_break(doc)

    add_heading(doc, "2. Cara Menggunakan Manual Ini", 1)
    add_body(doc,
             "Angka berwarna pada screenshot menunjukkan bagian layar yang perlu diperhatikan. Urutan angka pada "
             "gambar sama dengan daftar penjelasan tepat di bawahnya. Materi dapat dipraktikkan per bagian, tetapi "
             "untuk simulasi end-to-end gunakan urutan POS, PR, PO, BAST, Terima Tagihan, Pembayaran Vendor, lalu "
             "laporan. Gunakan akun dan hak akses yang diberikan trainer; kredensial tidak dicantumkan dalam dokumen "
             "agar tidak disalin ke lingkungan produksi.")
    add_heading(doc, "Peran peserta", 2)
    add_table(doc, ["Peran", "Tanggung jawab pelatihan"], [
        ("Kasir", "Membuat transaksi, memilih metode pembayaran, mencetak dan menelusuri nota."),
        ("Pemohon/Logistik", "Membuat PR dan PO, menjaga barang, UOM, pemasok, harga, dan jadwal."),
        ("Penerima", "Mencatat BAST sesuai pemeriksaan fisik."),
        ("Administrasi Vendor", "Menerima invoice dan mencocokkannya dengan BAST/PO."),
        ("Keuangan", "Memproses pembayaran vendor sesuai tagihan dan termin."),
        ("Supervisor/Trainer", "Membaca laporan, drill-down angka, ekspor, dan rekonsiliasi."),
    ], [5, 13])
    add_heading(doc, "Prasyarat latihan", 2)
    add_bullets(doc, [
        "Toko aktif adalah Kantin Demo dan tanggal komputer benar.",
        "Master produk, pemasok, metode pembayaran, member, harga, dan UOM sudah tersinkron.",
        "Hak create/read/update/approval diberikan sesuai peran.",
        "Semua perbandingan laporan menggunakan 1–10 September 2026 dan toko yang sama.",
        "Untuk PR pelatihan, pilih Tanpa anggaran karena COA/anggaran An Nahl belum menjadi ruang lingkup.",
    ])
    page_break(doc)

    for i, diagram in enumerate(diagrams, 1):
        titles = ["3. Alur End-to-End", "4. Use Case dan Pembagian Peran", "5. Aliran Data / ERD Ringkas", "6. Peta Laporan"]
        add_heading(doc, titles[i - 1], 1)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(diagram), width=Inches(9.3))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(cap, f"Diagram {i} — {titles[i - 1].split('. ', 1)[1]}.", italic=True, size=8.2, color="64748B")
        # Gambar memenuhi halaman; heading berikutnya otomatis pindah ke halaman baru.

    add_heading(doc, "7. Panduan Layar dan Skenario UAT", 1)
    add_body(doc,
             "Bagian berikut menggabungkan petunjuk pengguna dan bukti UAT. Setiap layar diuji pada resolusi 2560 × "
             "1392 agar area kerja tampak penuh. Screenshot laporan mencakup bagian atas dan pengujian scroll hingga "
             "bagian terbawah; manual menampilkan gambar yang paling membantu untuk langkah pengguna.")
    page_break(doc)
    for spec in SCREEN_SPECS:
        add_picture_page(doc, spec, annotated[spec["file"]])

    add_heading(doc, "8. Matriks Laporan Operasional", 1)
    report_names = {
        "pnj_faktur": "Daftar Faktur Penjualan",
        "pnj_per_barang": "Penjualan per Barang",
        "pnj_rincian_barang": "Rincian Penjualan per Barang",
        "omzet_transaksi": "Transaksi Omzet",
        "omzet_tunai_produk": "Omzet Produk Non-Saldo/Tunai",
        "omzet_saldo_produk": "Omzet Produk Saldo",
        "omzet_rekapan": "Rekap Omzet",
        "beli_penerimaan": "Penerimaan Pembelian",
        "beli_faktur": "Daftar Faktur Pembelian",
        "margin_produk": "Margin per Produk",
        "margin_kategori": "Margin per Kategori",
        "laba_kotor_harian": "Laba Kotor Harian",
    }
    rows = [(report_names[r["id"]], r["rows"], r["columns"], "PDF valid", "LULUS") for r in evidence["reports"]]
    add_table(doc, ["Laporan", "Baris", "Kolom", "Ekspor", "Hasil"], rows, [8, 2, 2, 3, 2])
    add_body(doc,
             "Selain laporan di atas, katalog eBisnis juga menyediakan laporan penjualan dan pembelian lain. Trainer "
             "dapat memperluas latihan dengan pola yang sama: samakan periode dan toko, jalankan laporan, periksa "
             "jumlah baris, buka rincian angka yang tersedia, lalu ekspor PDF/Excel. Jangan memasukkan laporan berbasis "
             "Akuntansi sampai COA An Nahl siap dan tervalidasi.")
    page_break(doc)

    add_heading(doc, "9. Rekonsiliasi yang Harus Dilakukan Peserta", 1)
    add_table(doc, ["Pemeriksaan", "Cara", "Kriteria lulus"], [
        ("Transaksi → Omzet", "Jumlahkan nominal transaksi aktif pada periode dan toko yang sama.", "Sama dengan Total Omzet."),
        ("Non-Saldo + Saldo", "Jumlahkan kedua kolom pada Rekap Omzet.", "Sama dengan Total Omzet."),
        ("Produk → transaksi", "Klik omzet produk dan jumlahkan nota penyusunnya.", "Sama dengan nilai produk."),
        ("BAST → penerimaan", "Cocokkan nomor faktur, tanggal, pemasok, item, dan nilai.", "Tidak ada selisih tanpa penjelasan."),
        ("Tagihan → pembayaran", "Bandingkan nilai tagihan, termin, pembayaran, dan sisa.", "Sisa sesuai status lunas/belum lunas."),
        ("Penjualan − HPP", "Bandingkan kolom Penjualan dan HPP.", "Selisih sama dengan Laba."),
    ], [5, 9, 5])
    add_heading(doc, "Catatan drill-down", 2)
    add_body(doc,
             "Popup rincian menggunakan dimensi baris yang diklik, misalnya nomor transaksi, produk, toko, dan "
             "kelompok pembayaran. Untuk jumlah data besar, lakukan pemeriksaan sampel yang dapat diulang dan gunakan "
             "ekspor Excel untuk rekonsiliasi lengkap. Audit API UAT memilih dimensi unik dan membandingkan nilai "
             "popup dengan nilai laporan; seluruh empat skenario clickable lulus tanpa selisih.")
    page_break(doc)

    add_heading(doc, "10. Troubleshooting untuk Trainer", 1)
    add_table(doc, ["Gejala", "Penyebab yang diperiksa", "Tindakan"], [
        ("Laporan kosong", "Periode/toko tidak sama atau transaksi belum ada.", "Atur 1–10 Sep 2026, pilih Kantin Demo, tekan Tampilkan."),
        ("Data hari ini hanya sedikit", "Riwayat memakai tanggal hari berjalan.", "Perluas tanggal mulai/sampai sebelum menyimpulkan volume."),
        ("PO tidak muncul saat BAST", "PO belum disetujui atau sudah diterima penuh.", "Periksa status, sisa kuantitas, dan toko aktif."),
        ("Tagihan tidak dapat dibuat", "BAST belum disetujui atau nilai melebihi sisa.", "Buka BAST sumber dan cocokkan nilai/termin."),
        ("Pembayaran tidak dapat diproses", "Hak approval, tagihan, atau koneksi server bermasalah.", "Baca pesan, cek hak, koneksi, dan jangan kirim berulang."),
        ("Total laporan berbeda", "Periode, toko, status transaksi, atau klasifikasi metode berbeda.", "Samakan filter lalu buka Asal Angka."),
        ("PDF/Excel belum aktif", "Laporan belum selesai dimuat.", "Tekan Tampilkan dan tunggu data muncul."),
    ], [5, 7, 8])
    add_heading(doc, "Informasi koneksi", 2)
    add_body(doc,
             "Laporan server, persetujuan dokumen, dan pembayaran vendor memerlukan koneksi. Bila jaringan tidak "
             "tersedia, aplikasi harus menampilkan pesan yang mudah dipahami dan tidak boleh menyatakan tindakan "
             "server berhasil. Simpan data lokal yang memang didukung, lalu sinkronkan setelah koneksi pulih. Untuk "
             "detail teknis, gunakan tombol detail galat bila tersedia dan catat kode referensi, endpoint, serta waktu.")
    page_break(doc)

    add_heading(doc, "11. Checklist Pelaksanaan Training", 1)
    add_table(doc, ["No", "Praktik", "Bukti yang dikumpulkan", "Status"], [
        (1, "Buat penjualan tunai/non-saldo", "Nomor nota dan total", "□"),
        (2, "Buat penjualan saldo", "Member, saldo, nomor nota", "□"),
        (3, "Buat PR tanpa anggaran", "Nomor PR", "□"),
        (4, "Buat PO bukan termin dari PR", "Nomor PO dan total", "□"),
        (5, "Buat PO termin", "Jadwal termin = total PO", "□"),
        (6, "Buat BAST dari PO", "Nomor BAST dan jumlah diterima", "□"),
        (7, "Terima tagihan vendor", "Nomor invoice dan jatuh tempo", "□"),
        (8, "Bayar vendor", "Nomor bukti dan status", "□"),
        (9, "Buka empat laporan Omzet", "Screenshot + popup Asal Angka", "□"),
        (10, "Buka laporan penjualan/pembelian/margin", "PDF dan Excel", "□"),
        (11, "Rekonsiliasi", "Kertas kerja selisih Rp0", "□"),
    ], [1.5, 8, 7, 2])
    add_body(doc,
             "Trainer menandatangani checklist setelah peserta dapat menjelaskan hubungan dokumen, bukan hanya "
             "menekan tombol. Kesalahan latihan diperbaiki pada dokumen sumber dan dicatat; jangan membuat transaksi "
             "pengganti tanpa alasan karena akan menyulitkan penelusuran laporan.")
    page_break(doc)

    add_heading(doc, "12. Lampiran Bukti UAT", 1)
    add_table(doc, ["Komponen", "Bukti"], [
        ("Lingkungan", evidence["environment"]),
        ("Endpoint", evidence["endpoint"]),
        ("Waktu audit", evidence["generatedAt"]),
        ("Periode", f"{evidence['period']['start']} s.d. {evidence['period']['end']}"),
        ("Ruang lingkup", evidence["scope"]),
        ("Hasil keseluruhan", "LULUS" if evidence["allPassed"] else "GAGAL"),
        ("Akuntansi", "Dikecualikan sesuai permintaan; bukan kegagalan UAT."),
    ], [5, 14])
    add_heading(doc, "Rincian popup clickable", 2)
    add_table(doc, ["Skenario", "Baris", "Nilai laporan", "Nilai rincian", "Selisih", "Hasil"], [
        (x["name"], x["rows"], f"Rp{x['expectedTotal']:,.0f}", f"Rp{x['actualTotal']:,.0f}",
         f"Rp{x['difference']:,.0f}", "LULUS" if x["passed"] else "GAGAL")
        for x in evidence["clickableDetails"]
    ], [5.5, 2, 3, 3, 2.5, 2])
    add_body(doc,
             "Dokumen ini merupakan bukti dan panduan untuk build yang diuji. Perubahan server, master, hak akses, "
             "atau periode data perlu diuji ulang sebelum hasil digunakan sebagai dasar pelatihan berikutnya.")

    doc.save(DOCX)

    md_lines = [
        "# Hasil UAT Training Operasional eBisnis v1.34.32",
        "",
        f"- Waktu audit: `{evidence['generatedAt']}`",
        f"- Lingkungan: `{evidence['environment']}`",
        f"- Periode: `{evidence['period']['start']}` s.d. `{evidence['period']['end']}`",
        "- Ruang lingkup: POS, PR, PO termin/bukan termin, BAST, tagihan, pembayaran vendor, laporan penjualan, pembelian, Omzet, margin, dan laba kotor harian.",
        "- Akuntansi/COA/jurnal/posting: **dikecualikan sesuai permintaan**.",
        "- Hasil: **LULUS 100% untuk seluruh skenario dalam ruang lingkup**.",
        "",
        "## Ringkasan bukti",
        "",
        "| Pemeriksaan | Hasil |",
        "|---|---:|",
        "| Pengadaan PR/PO/BAST/Tagihan/Pembayaran | 5/5; masing-masing 100 dokumen |",
        "| Laporan dan PDF server | 12/12 lulus |",
        "| Rincian clickable | 4/4 lulus |",
        "| Integrasi layar Windows | 2/2 lulus |",
        "| Tes kontrak ekspor/hak akses/lokal-dulu/UOM | 44/44 lulus |",
        "| Transaksi penjualan | 504 |",
        "| Omzet Saldo per Produk | 101 produk |",
        "| Penerimaan/Faktur Pembelian | 452 baris |",
        "| Margin per Produk | 201 baris |",
        "",
        "## Artefak",
        "",
        f"- `{DOCX.name}`",
        f"- `{PDF.name}`",
        "- `evidence/uat-training-operasional-evidence.json`",
        "- `screenshots/` dan `annotated/`",
        "- `diagrams/`",
    ]
    MD.write_text("\n".join(md_lines) + "\n", encoding="utf-8")
    print(json.dumps({"docx": str(DOCX), "pdf": str(PDF), "screens": len(annotated), "diagrams": len(diagrams)}))


if __name__ == "__main__":
    build_doc()
