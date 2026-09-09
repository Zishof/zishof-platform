"""Regenerate page-friendly AB Chicken/RPA diagrams and replace DOCX embeds.

All canvases use a 3:2 ratio so they remain readable on a landscape Word/PDF
page. Connectors are explicitly routed through gutters and never cross text.
"""

from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parent
UAT = ROOT / "uat-pos-desktop-e2e-20260909"
RPA = ROOT / "integrasi-rpa-20260909"
W, H = 1800, 1200

NAVY = "#102A43"
BLUE = "#1F5D8F"
BLUE_FILL = "#E7F1F8"
GREEN = "#4D7F33"
GREEN_FILL = "#EAF5E5"
ORANGE = "#D56A12"
ORANGE_FILL = "#FFF0E3"
GRAY = "#5B6778"
GRAY_FILL = "#F3F5F7"
LINE = "#49657F"
WHITE = "#FFFFFF"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


F_TITLE = font(48, True)
F_SUBTITLE = font(25)
F_LANE = font(27, True)
F_CARD = font(27, True)
F_SMALL = font(21)
F_TINY = font(18)
F_HEAD = font(26, True)


def canvas(title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (W, H), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((55, 38), title, font=F_TITLE, fill=NAVY)
    draw.text((58, 102), subtitle, font=F_SUBTITLE, fill=GRAY)
    draw.line((55, 145, W - 55, 145), fill="#D7E2EC", width=3)
    return image, draw


def wrapped(draw: ImageDraw.ImageDraw, text: str, selected_font, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            trial = f"{current} {word}"
            if draw.textbbox((0, 0), trial, font=selected_font)[2] <= max_width:
                current = trial
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    selected_font=F_CARD,
    fill=NAVY,
    padding: int = 22,
) -> None:
    x1, y1, x2, y2 = box
    lines = wrapped(draw, text, selected_font, x2 - x1 - 2 * padding)
    heights = [draw.textbbox((0, 0), line or "Ag", font=selected_font)[3] for line in lines]
    line_gap = 7
    total = sum(heights) + max(0, len(lines) - 1) * line_gap
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=selected_font)
        width = bbox[2] - bbox[0]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=selected_font, fill=fill)
        y += height + line_gap


def card(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    fill=BLUE_FILL,
    outline=BLUE,
    radius: int = 20,
    selected_font=F_CARD,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=4)
    centered_text(draw, box, text, selected_font)


def lane(draw: ImageDraw.ImageDraw, y1: int, y2: int, title: str, alternate: bool = False) -> None:
    fill = "#F7FAFC" if alternate else "#EEF4F8"
    draw.rectangle((55, y1, W - 55, y2), fill=fill, outline="#CBD5E1", width=2)
    draw.text((72, y1 + 12), title, font=F_LANE, fill=NAVY)


def arrow(
    draw: ImageDraw.ImageDraw,
    points: list[tuple[int, int]],
    color=LINE,
    width: int = 5,
    label: str | None = None,
    label_at: tuple[int, int] | None = None,
    dashed: bool = False,
) -> None:
    for start, end in zip(points, points[1:]):
        if dashed:
            x1, y1 = start
            x2, y2 = end
            length = math.hypot(x2 - x1, y2 - y1)
            if length == 0:
                continue
            dx, dy = (x2 - x1) / length, (y2 - y1) / length
            pos = 0.0
            while pos < length:
                segment_end = min(pos + 14, length)
                draw.line(
                    (x1 + dx * pos, y1 + dy * pos, x1 + dx * segment_end, y1 + dy * segment_end),
                    fill=color,
                    width=width,
                )
                pos += 24
        else:
            draw.line((*start, *end), fill=color, width=width, joint="curve")
    (x1, y1), (x2, y2) = points[-2], points[-1]
    angle = math.atan2(y2 - y1, x2 - x1)
    size = 18
    left = (x2 - size * math.cos(angle - math.pi / 6), y2 - size * math.sin(angle - math.pi / 6))
    right = (x2 - size * math.cos(angle + math.pi / 6), y2 - size * math.sin(angle + math.pi / 6))
    draw.polygon([(x2, y2), left, right], fill=color)
    if label and label_at:
        bbox = draw.textbbox((0, 0), label, font=F_SMALL)
        x, y = label_at
        draw.rounded_rectangle(
            (x - 8, y - 4, x + bbox[2] + 8, y + bbox[3] + 5),
            radius=6,
            fill=WHITE,
        )
        draw.text((x, y), label, font=F_SMALL, fill=color)


def diamond(draw: ImageDraw.ImageDraw, cx: int, cy: int, w: int, h: int, text: str) -> None:
    points = [(cx, cy - h // 2), (cx + w // 2, cy), (cx, cy + h // 2), (cx - w // 2, cy)]
    draw.polygon(points, fill="#FFF7D6", outline=ORANGE)
    draw.line(points + [points[0]], fill=ORANGE, width=4, joint="curve")
    centered_text(draw, (cx - w // 2 + 18, cy - h // 2 + 8, cx + w // 2 - 18, cy + h // 2 - 8), text, F_HEAD)


def footer_note(draw: ImageDraw.ImageDraw, text: str) -> None:
    draw.rounded_rectangle((55, 1115, W - 55, 1170), radius=12, fill="#EDF4FA", outline="#C8D8E8", width=2)
    centered_text(draw, (70, 1120, W - 70, 1165), text, F_SMALL, GRAY)


def save(image: Image.Image, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, optimize=True)


def draw_uat_flow(path: Path) -> None:
    image, d = canvas(
        "Siklus Operasional POS Desktop AB Chicken",
        "Cabang stok tersedia dan stok kurang bertemu kembali sebelum Delivery Order.",
    )
    card(d, (80, 185, 390, 305), "1  Outlet buat\npesanan bahan")
    card(d, (485, 185, 795, 305), "2  Gudang periksa\nstok dan reservasi")
    diamond(d, 1010, 245, 290, 150, "Stok\ncukup?")
    card(d, (1270, 185, 1715, 305), "3A  Picking dan packing", GREEN_FILL, GREEN)
    arrow(d, [(390, 245), (485, 245)])
    arrow(d, [(795, 245), (865, 245)])
    arrow(d, [(1155, 245), (1270, 245)], GREEN, label="YA", label_at=(1180, 205))

    lane(d, 355, 785, "Cabang stok kurang — Pengadaan, penerimaan, dan persiapan")
    card(d, (880, 410, 1160, 520), "3B  PR", ORANGE_FILL, ORANGE)
    card(d, (1320, 410, 1660, 520), "4  PO vendor", ORANGE_FILL, ORANGE)
    card(d, (1320, 600, 1660, 710), "5  BAST + QC", ORANGE_FILL, ORANGE)
    card(d, (880, 600, 1160, 710), "6  Produksi / repack", GREEN_FILL, GREEN)
    card(d, (440, 600, 740, 710), "Tagihan + bayar vendor", GRAY_FILL, GRAY, selected_font=F_HEAD)
    arrow(d, [(1010, 320), (1010, 410)], ORANGE, label="TIDAK", label_at=(1025, 346))
    arrow(d, [(1160, 465), (1320, 465)], ORANGE)
    arrow(d, [(1490, 520), (1490, 600)], ORANGE)
    arrow(d, [(1320, 655), (1160, 655)], GREEN)
    arrow(d, [(1320, 685), (1225, 685), (1225, 750), (590, 750), (590, 710)], GRAY, dashed=True, label="AP/keuangan", label_at=(760, 719))

    lane(d, 815, 1085, "Pemenuhan outlet — Jalur operasi kembali menyatu")
    card(d, (1335, 885, 1690, 1000), "7  Delivery Order")
    card(d, (920, 885, 1275, 1000), "8  Kirim + pantau")
    card(d, (505, 885, 860, 1000), "9  BAST outlet")
    card(d, (90, 885, 445, 1000), "10  Produksi + POS", GREEN_FILL, GREEN)
    # Jalur stok cukup memakai gutter kanan; jangan menembus kotak PO/BAST.
    arrow(d, [(1490, 305), (1755, 305), (1755, 942), (1690, 942)], GREEN)
    arrow(d, [(1020, 710), (1020, 800), (1510, 800), (1510, 885)], GREEN)
    arrow(d, [(1335, 942), (1275, 942)])
    arrow(d, [(920, 942), (860, 942)])
    arrow(d, [(505, 942), (445, 942)])
    arrow(d, [(90, 942), (65, 942), (65, 245), (80, 245)], ORANGE, label="stok menipis → siklus baru", label_at=(80, 1040))
    footer_note(d, "Garis utuh = alur barang/status • garis putus = penyelesaian tagihan • seluruh transaksi dijalankan dari POS Desktop.")
    save(image, path)


def draw_uat_use_case(path: Path) -> None:
    image, d = canvas(
        "Use Case — Peran dan Tanggung Jawab",
        "Setiap aktor berada pada jalurnya sendiri; konektor berhenti di tepi use case dan tidak melintasi teks.",
    )
    actors = [
        ("Operator Outlet", "Buat pesanan bahan", "BAST, produksi, dan jual di POS"),
        ("Petugas Gudang Pusat", "Cek stok, picking, packing", "Buat DO dan pantau pengiriman"),
        ("Staf Pengadaan", "PR, PO, dan BAST vendor", "Tagihan, pembayaran, klaim"),
        ("Keuangan & Akuntansi", "Periksa sumber akun dan posting", "Jurnal, buku besar, laporan"),
    ]
    y = 180
    for index, (actor, first, second) in enumerate(actors):
        lane(d, y - 12, y + 200, f"{index + 1}. Jalur {actor}", index % 2 == 1)
        actor_box = (80, y + 40, 390, y + 155)
        first_box = (505, y + 18, 1040, y + 103)
        second_box = (1165, y + 18, 1710, y + 103)
        card(d, actor_box, actor, GRAY_FILL, GRAY)
        card(d, first_box, first)
        card(d, second_box, second, GREEN_FILL if index in (0, 3) else BLUE_FILL, GREEN if index in (0, 3) else BLUE)
        arrow(d, [(390, y + 80), (445, y + 80), (445, y + 60), (505, y + 60)])
        arrow(d, [(1040, y + 60), (1165, y + 60)])
        d.text((515, y + 120), "Tindakan awal", font=F_TINY, fill=GRAY)
        d.text((1175, y + 120), "Hasil / tindak lanjut", font=F_TINY, fill=GRAY)
        y += 225
    footer_note(d, "Admin Utama menggunakan Web hanya untuk registrasi, konfigurasi akses, monitoring, dan audit—bukan transaksi operasional.")
    save(image, path)


def entity(d, x: int, y: int, title: str, fields: str, fill=BLUE_FILL, outline=BLUE) -> None:
    box = (x, y, x + 300, y + 118)
    d.rounded_rectangle(box, radius=16, fill=fill, outline=outline, width=4)
    d.rectangle((x, y, x + 300, y + 42), fill=outline)
    centered_text(d, (x + 8, y + 2, x + 292, y + 40), title, F_HEAD, WHITE)
    centered_text(d, (x + 12, y + 45, x + 288, y + 113), fields, F_SMALL, NAVY)


def relationship(d, x1: int, y: int, x2: int, label: str = "1:N") -> None:
    arrow(d, [(x1, y), (x2, y)], LINE, width=4)
    d.text(((x1 + x2) // 2 - 18, y - 32), label, font=F_TINY, fill=GRAY)


def draw_uat_erd(path: Path) -> None:
    image, d = canvas(
        "ERD Ringkas dan Aliran Data End-to-End",
        "Empat lapisan memisahkan permintaan, pengadaan, pemenuhan, dan akuntansi agar relasi mudah ditelusuri.",
    )
    rows = [
        (175, "A. Permintaan Outlet", [("Outlet", "id, lokasi"), ("OutletOrder", "nomor, status"), ("OrderLine", "barang, qty"), ("Material", "sku, stok"), ("BOM", "produk, bahan, qty")]),
        (405, "B. Pengadaan Vendor", [("PurchaseRequest", "nomor, kebutuhan"), ("PurchaseOrder", "vendor, termin"), ("VendorReceipt", "BAST, diterima"), ("VendorInvoice", "tagihan, jatuh tempo"), ("VendorPayment", "bank, nominal")]),
        (635, "C. Produksi & Pemenuhan", [("ProductionOrder", "batch, hasil"), ("Shipment", "DO, tujuan"), ("OutletReceipt", "BAST, selisih"), ("PosSale", "faktur, total"), ("InventoryLedger", "lot, lokasi, mutasi")]),
        (865, "D. Akuntansi & Laporan", [("SourceDocument", "jenis, referensi"), ("JournalEntry", "tanggal, status"), ("JournalLine", "debit, kredit"), ("Account", "kode, kelompok"), ("FinancialReport", "periode, saldo")]),
    ]
    xs = [55, 410, 765, 1120, 1475]
    for row_index, (y, title, nodes) in enumerate(rows):
        d.text((55, y - 40), title, font=F_LANE, fill=[BLUE, ORANGE, GREEN, NAVY][row_index])
        for x, (name, fields) in zip(xs, nodes):
            fill, outline = (GREEN_FILL, GREEN) if row_index == 3 else ((ORANGE_FILL, ORANGE) if row_index == 1 else (BLUE_FILL, BLUE))
            entity(d, x, y, name, fields, fill, outline)
        for index in range(4):
            relationship(d, xs[index] + 300, y + 60, xs[index + 1], "1:N" if index != 3 else "N:1")
    footer_note(d, "Kunci telusur: nomor dokumen sumber → JournalEntry → JournalLine → Account → laporan periode; status POSTED menjadi gerbang.")
    save(image, path)


def draw_uat_accounts(path: Path) -> None:
    image, d = canvas(
        "Relasi Sumber Akun pada Setiap Posting",
        "Kolom proses, debit, dan kredit dibaca per baris—tidak ada garis yang melewati isi akun.",
    )
    headers = [(65, 390, "PROSES / SUMBER"), (430, 1050, "DEBIT"), (1090, 1735, "KREDIT")]
    for x1, x2, text in headers:
        d.rounded_rectangle((x1, 175, x2, 230), radius=10, fill=NAVY)
        centered_text(d, (x1, 175, x2, 230), text, F_HEAD, WHITE)
    rows = [
        ("BAST vendor", "114100 Persediaan Bahan Baku", "210100 GRNI / Barang Belum Ditagih"),
        ("Tagihan vendor", "210100 GRNI", "210200 Hutang Vendor"),
        ("Pembayaran vendor", "210200 Hutang Vendor", "111200 Bank Operasional"),
        ("Produksi / repack", "114200 Persediaan Barang Jadi", "114100 Bahan Baku + BOP"),
        ("Pengiriman antar lokasi", "114100 Persediaan Outlet", "114100 Persediaan Gudang"),
        ("Penjualan POS", "111200 Kas/Bank + 510100 HPP", "410100 Pendapatan + 114200 Persediaan"),
    ]
    y = 255
    for index, (process, debit, credit) in enumerate(rows):
        fill = "#FAFBFC" if index % 2 == 0 else "#F3F6F9"
        d.rounded_rectangle((65, y, 1735, y + 125), radius=14, fill=fill, outline="#D3DEE8", width=2)
        card(d, (80, y + 15, 375, y + 110), process, GRAY_FILL, GRAY, selected_font=F_HEAD)
        card(d, (445, y + 15, 1035, y + 110), debit, GREEN_FILL, GREEN, selected_font=F_HEAD)
        card(d, (1105, y + 15, 1720, y + 110), credit, ORANGE_FILL, ORANGE, selected_font=F_HEAD)
        arrow(d, [(375, y + 62), (445, y + 62)], GREEN, width=4)
        arrow(d, [(1035, y + 62), (1105, y + 62)], ORANGE, width=4)
        y += 137
    footer_note(d, "Sumber akun: Master Produk/BOM, metode pembayaran, lokasi, vendor, dan menu Sumber Akun Posting; ubah master lalu muat ulang pratinjau.")
    save(image, path)


def draw_rpa_context(path: Path) -> None:
    image, d = canvas(
        "Konteks Integrasi RPA dan AB Chicken",
        "Arus barang, dokumen internal, penjualan eksternal, dan pengawasan dipisahkan dengan jalur yang jelas.",
    )
    card(d, (80, 225, 390, 380), "Peternak\nAyam Hidup", GREEN_FILL, GREEN)
    card(d, (590, 205, 1110, 400), "Rumah Pemotongan Ayam\nQC • Produksi • Cold Chain", ORANGE_FILL, ORANGE)
    card(d, (1390, 205, 1720, 400), "Gudang Pusat\nPR • PO • BAST")
    arrow(d, [(390, 290), (590, 290)], BLUE, label="ayam hidup + dokumen asal", label_at=(400, 245))
    arrow(d, [(1390, 245), (1110, 245)], ORANGE, label="PO internal", label_at=(1170, 205))
    arrow(d, [(1110, 350), (1390, 350)], BLUE, label="daging + DO + suhu", label_at=(1130, 305))

    card(d, (1390, 690, 1720, 845), "Outlet AB Chicken\nProduksi • POS", GREEN_FILL, GREEN)
    card(d, (590, 710, 1110, 865), "Admin Utama Web\nKonfigurasi • Monitor • Audit", GRAY_FILL, GRAY)
    card(d, (80, 690, 390, 845), "Pelanggan\nRPA / Outlet", GRAY_FILL, GRAY)
    arrow(d, [(1555, 400), (1555, 690)], BLUE, label="replenishment + BAST", label_at=(1570, 515))
    # Jalur penjualan diputar di bawah kartu Admin agar tidak menutup teksnya.
    arrow(d, [(1390, 770), (1350, 770), (1350, 940), (430, 940), (430, 770), (390, 770)], GREEN, label="produk jadi / transaksi POS", label_at=(720, 895))
    arrow(d, [(850, 710), (850, 620), (850, 480)], GRAY, dashed=True, label="monitor & audit", label_at=(870, 570))
    arrow(d, [(850, 480), (850, 400)], GRAY, dashed=True)
    arrow(d, [(1110, 790), (1300, 790), (1300, 440), (1555, 440), (1555, 400)], GRAY, dashed=True)
    footer_note(d, "Operasional dilakukan dari POS Desktop; Web hanya untuk registrasi tenant, konfigurasi, monitoring, dan audit.")
    save(image, path)


def draw_rpa_flow(path: Path) -> None:
    image, d = canvas(
        "Flowchart Siklus PR sampai Penerimaan Daging",
        "Swimlane memperlihatkan pemilik aktivitas, bukti, dan titik serah antarunit.",
    )
    lane(d, 170, 360, "GUDANG PUSAT", False)
    lane(d, 360, 640, "RPA — PENERIMAAN & PRODUKSI", True)
    lane(d, 640, 850, "PENGIRIMAN / COLD CHAIN", False)
    lane(d, 850, 1090, "GUDANG PUSAT — PENERIMAAN", True)
    card(d, (230, 220, 500, 315), "1  Buat PR")
    card(d, (650, 220, 960, 315), "2  Kirim PO internal")
    arrow(d, [(500, 267), (650, 267)])
    card(d, (230, 430, 500, 525), "3  Terima pesanan")
    card(d, (650, 430, 960, 525), "4  Ayam datang + QC")
    card(d, (1110, 430, 1490, 525), "5  Potong + catat hasil lot", ORANGE_FILL, ORANGE)
    arrow(d, [(805, 315), (805, 430)], ORANGE)
    arrow(d, [(500, 477), (650, 477)])
    arrow(d, [(960, 477), (1110, 477)])
    card(d, (1110, 690, 1490, 785), "6  POS/transfer + DO")
    card(d, (650, 690, 960, 785), "7  Kirim + pantau suhu")
    arrow(d, [(1300, 525), (1300, 690)], ORANGE)
    arrow(d, [(1110, 737), (960, 737)])
    card(d, (1110, 920, 1490, 1015), "8  BAST penerimaan")
    card(d, (650, 920, 960, 1015), "9  Stok gudang masuk")
    card(d, (230, 920, 500, 1015), "10  Hitung kebutuhan")
    arrow(d, [(805, 785), (805, 875), (1300, 875), (1300, 920)])
    arrow(d, [(1110, 967), (960, 967)])
    arrow(d, [(650, 967), (500, 967)])
    # Loop berada di gutter kiri, di luar judul setiap swimlane.
    arrow(d, [(230, 967), (65, 967), (65, 267), (230, 267)], ORANGE, label="siklus berikutnya", label_at=(80, 790))
    footer_note(d, "Titik kontrol: QC ayam hidup → yield batch → suhu pengiriman → BAST/selisih → stok masuk → kebutuhan baru.")
    save(image, path)


def actor_symbol(d, cx: int, cy: int, label: str) -> None:
    d.ellipse((cx - 35, cy - 70, cx + 35, cy), outline=NAVY, width=5)
    d.line((cx, cy, cx, cy + 95), fill=NAVY, width=5)
    d.line((cx - 45, cy + 35, cx + 45, cy + 35), fill=NAVY, width=5)
    d.line((cx, cy + 95, cx - 45, cy + 160), fill=NAVY, width=5)
    d.line((cx, cy + 95, cx + 45, cy + 160), fill=NAVY, width=5)
    centered_text(d, (cx - 100, cy + 168, cx + 100, cy + 235), label, F_HEAD)


def draw_rpa_use_case(path: Path) -> None:
    image, d = canvas(
        "Use Case Operasional dan Pengawasan",
        "Empat kelompok peran; tiap garis berada di area kosong dan berhenti sebelum teks use case.",
    )
    groups = [
        (95, 195, "Perencana\nGudang", ["Buat PR + PO internal", "Pantau pemenuhan + backorder"]),
        (935, 195, "Operator\nRPA", ["Terima ayam hidup + QC", "Produksi lot + catat yield"]),
        (95, 665, "Pengemudi /\nPenerima", ["Kirim + pantau suhu", "BAST + klaim selisih"]),
        (935, 665, "Admin\nUtama", ["Konfigurasi + audit", "Rekonsiliasi + laporan"]),
    ]
    for x, y, actor, cases in groups:
        d.rounded_rectangle((x - 15, y - 20, x + 770, y + 390), radius=20, fill="#F8FAFC", outline="#D7E2EC", width=2)
        actor_symbol(d, x + 105, y + 105, actor)
        first = (x + 300, y + 35, x + 720, y + 145)
        second = (x + 300, y + 220, x + 720, y + 330)
        card(d, first, cases[0])
        card(d, second, cases[1], GREEN_FILL, GREEN)
        arrow(d, [(x + 150, y + 120), (x + 235, y + 120), (x + 235, y + 90), (x + 300, y + 90)], GRAY, width=3)
        arrow(d, [(x + 150, y + 140), (x + 250, y + 140), (x + 250, y + 275), (x + 300, y + 275)], GRAY, width=3)
    footer_note(d, "Pisahkan pembuat, pemeriksa, penerima, dan auditor melalui role; Admin Web tidak menggantikan operator transaksi Desktop.")
    save(image, path)


def draw_rpa_erd(path: Path) -> None:
    image, d = canvas(
        "ERD Ringkas dan Rantai Ketertelusuran RPA",
        "Relasi dikelompokkan per tahap agar asal ternak, hasil produksi, pengiriman, dan jurnal dapat ditelusuri.",
    )
    rows = [
        (190, "A. ASAL & PENERIMAAN", [("BusinessUnit", "id, jenis, lokasi"), ("InternalOrder", "nomor, pemesan"), ("LiveBirdReceipt", "peternak, flock, qty"), ("QualityCheck", "parameter, hasil")]),
        (440, "B. PRODUKSI", [("SlaughterBatch", "lot, input, waktu"), ("OutputLot", "sku, kg, grade"), ("ByProduct", "jenis, berat"), ("InventoryLedger", "lot, lokasi, mutasi")]),
        (690, "C. DISTRIBUSI", [("DeliveryOrder", "nomor, tujuan, suhu"), ("OutletReceipt", "BAST, diterima"), ("ExceptionClaim", "selisih, sebab"), ("PosSale", "faktur, pelanggan")]),
        (940, "D. AKUNTANSI", [("SourceDocument", "referensi, status"), ("JournalEntry", "tanggal, sumber"), ("JournalLine", "akun, debit, kredit"), ("AccountReport", "kelompok, saldo")]),
    ]
    xs = [80, 510, 940, 1370]
    for row_index, (y, title, nodes) in enumerate(rows):
        d.text((80, y - 38), title, font=F_LANE, fill=[BLUE, ORANGE, GREEN, NAVY][row_index])
        for x, (name, fields) in zip(xs, nodes):
            fill, outline = ((ORANGE_FILL, ORANGE) if row_index == 1 else ((GREEN_FILL, GREEN) if row_index == 3 else (BLUE_FILL, BLUE)))
            entity(d, x, y, name, fields, fill, outline)
        for index in range(3):
            relationship(d, xs[index] + 300, y + 60, xs[index + 1], "1:N")
    footer_note(d, "Kunci audit: peternak/flock → penerimaan → QC → batch potong → output lot → DO → BAST → stok → jurnal.")
    save(image, path)


def draw_rpa_accounts(path: Path) -> None:
    image, d = canvas(
        "Aliran Akuntansi RPA dan Gudang Pusat",
        "Debit dan kredit disajikan dalam kartu vertikal yang proporsional; sumber master dicatat pada kolom terakhir.",
    )
    headers = [(55, 380, "PROSES"), (410, 920, "DEBIT"), (950, 1460, "KREDIT"), (1490, 1745, "SUMBER")]
    for x1, x2, text in headers:
        d.rounded_rectangle((x1, 175, x2, 230), radius=10, fill=NAVY)
        centered_text(d, (x1, 175, x2, 230), text, F_HEAD, WHITE)
    rows = [
        ("Terima ayam hidup", "Persediaan Ayam Hidup", "GRNI / Belum Ditagih", "BAST + barang"),
        ("Terima tagihan peternak", "GRNI", "Hutang Peternak", "Vendor + termin"),
        ("Produksi pemotongan", "Persediaan Daging + Samping", "Ayam Hidup + BOP/Upah", "BOM + batch"),
        ("Transfer satu entitas", "Persediaan Gudang Pusat", "Persediaan RPA", "Lokasi + lot"),
        ("Penjualan eksternal RPA", "Kas/Piutang + HPP", "Pendapatan + Persediaan", "POS + produk"),
    ]
    y = 260
    for index, values in enumerate(rows):
        d.rounded_rectangle((55, y, 1745, y + 145), radius=14, fill="#FAFBFC" if index % 2 == 0 else "#F2F6F9", outline="#D3DEE8", width=2)
        boxes = [(70, 370), (425, 905), (965, 1445), (1505, 1730)]
        fills = [(ORANGE_FILL, ORANGE), (BLUE_FILL, BLUE), (GREEN_FILL, GREEN), (GRAY_FILL, GRAY)]
        for (x1, x2), (fill, outline), value in zip(boxes, fills, values):
            card(d, (x1, y + 18, x2, y + 127), value, fill, outline, selected_font=F_HEAD if x2 - x1 > 300 else F_SMALL)
        y += 160
    footer_note(d, "Satu badan hukum: transfer antarlokasi tidak membentuk pendapatan eksternal • beda badan hukum: gunakan AP/AR intercompany + eliminasi.")
    save(image, path)


def resize_embedded_diagrams(docx_path: Path, media_names: set[str], width_inches: float) -> None:
    document = Document(docx_path)
    for shape in document.inline_shapes:
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        rel_id = blip.embed
        part = document.part.related_parts.get(rel_id)
        if part and Path(str(part.partname)).name in media_names:
            # Preserve the 3:2 source ratio.  UAT is landscape, while the RPA
            # guide is portrait, so each document receives a page-safe width.
            shape.width = Inches(width_inches)
            shape.height = Inches(width_inches * H / W)
    document.save(docx_path)


def replace_embedded_media(docx_path: Path, replacements: dict[str, Path], width_inches: float) -> None:
    """Replace selected /word/media members, then normalize their shape ratio."""
    import os
    import zipfile

    temporary = docx_path.with_suffix(".diagram-update.tmp")
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            name = Path(item.filename).name
            if item.filename.startswith("word/media/") and name in replacements:
                data = replacements[name].read_bytes()
            target.writestr(item, data)
    os.replace(temporary, docx_path)
    resize_embedded_diagrams(docx_path, set(replacements), width_inches)


def replace_text_preserving_runs(docx_path: Path, replacements: dict[str, str]) -> None:
    """Refresh release-summary text without discarding existing paragraph styles."""
    document = Document(docx_path)

    def replace_paragraph(paragraph) -> None:
        for old, new in replacements.items():
            if old not in paragraph.text:
                continue
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    break
            else:
                # The target can be split across runs. Keep the formatting of
                # the first run and clear the remainder of this paragraph.
                if paragraph.runs:
                    paragraph.runs[0].text = paragraph.text.replace(old, new)
                    for run in paragraph.runs[1:]:
                        run.text = ""

    for paragraph in document.paragraphs:
        replace_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_paragraph(paragraph)
    document.save(docx_path)


def main() -> None:
    uat_diagrams = UAT / "diagrams"
    rpa_diagrams = RPA / "diagrams"
    draw_uat_flow(uat_diagrams / "01-flow-siklus-operasional.png")
    draw_uat_use_case(uat_diagrams / "02-use-case-peran.png")
    draw_uat_erd(uat_diagrams / "03-erd-aliran-data.png")
    draw_uat_accounts(uat_diagrams / "04-relasi-akun-posting.png")

    draw_rpa_context(rpa_diagrams / "01-konteks-integrasi.png")
    draw_rpa_flow(rpa_diagrams / "02-flowchart-swimlane.png")
    draw_rpa_use_case(rpa_diagrams / "03-use-case.png")
    draw_rpa_erd(rpa_diagrams / "04-erd-aliran-data.png")
    draw_rpa_accounts(rpa_diagrams / "05-aliran-akuntansi.png")

    replace_embedded_media(
        UAT / "User-Manual-dan-UAT-E2E-AB-Chicken-POS-Desktop-2026-09-09.docx",
        {
            "image1.png": uat_diagrams / "02-use-case-peran.png",
            "image2.png": uat_diagrams / "01-flow-siklus-operasional.png",
            "image3.png": uat_diagrams / "03-erd-aliran-data.png",
            "image4.png": uat_diagrams / "04-relasi-akun-posting.png",
        },
        8.0,
    )
    replace_embedded_media(
        RPA / "Panduan-Integrasi-Rumah-Pemotongan-Ayam-dengan-AB-Chicken-2026-09-09.docx",
        {
            "image2.png": rpa_diagrams / "01-konteks-integrasi.png",
            "image3.png": rpa_diagrams / "03-use-case.png",
            "image4.png": rpa_diagrams / "02-flowchart-swimlane.png",
            "image5.png": rpa_diagrams / "04-erd-aliran-data.png",
            "image6.png": rpa_diagrams / "05-aliran-akuntansi.png",
        },
        6.2,
    )
    replace_text_preserving_runs(
        UAT / "User-Manual-dan-UAT-E2E-AB-Chicken-POS-Desktop-2026-09-09.docx",
        {
            "POS Desktop Windows variant abchicken, versi 1.34.27 build 190":
                "POS Desktop Windows variant abchicken, versi kandidat 1.34.30 build 193",
            "33 dari 33 pengujian": "33/33 kontrak awal + 23/23 regresi rilis",
            "19 dari 19 pemeriksaan": "19/19 kontrol integritas + 13/13 layar kandidat",
            "6 jurnal baru, ID 301 sampai 306": "10/10 proses; jurnal regresi 307–312",
            "594 baris jurnal dan selisih neraca nol": "6/6 laporan; 608 baris jurnal; neraca seimbang",
        },
    )


if __name__ == "__main__":
    main()
